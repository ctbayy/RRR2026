import os
from docx import Document
from docxtpl import DocxTemplate


def tao_template_mau_sach():
    """BƯỚC 1: Tự động tạo một Template Word sạch 100%, không dính mã XML rác"""
    doc = Document()

    # Các thẻ {{ }} này sẽ được docxtpl thay thế cực mượt
    doc.add_paragraph("PHIẾU NHẬN XÉT HỒ SƠ KHAI THUẾ").runs[0].bold = True
    doc.add_paragraph("Tên Công ty: {{ ten_cong_ty }}")
    doc.add_paragraph("Mã số thuế: {{ mst_cong_ty }}")
    doc.add_paragraph("\n1. Bảng đối chiếu Doanh thu Hóa đơn và Tờ khai:")

    # Tạo cái bảng khung xương (chỉ có dòng tiêu đề)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    headers = ['STT', 'Kỳ', 'Hóa đơn đầu ra', 'Tờ khai GTGT', 'Chênh lệch']
    for i, text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True

    doc.add_paragraph("\n2. Nhận xét Rủi ro Báo cáo Tài chính:")
    doc.add_paragraph("{{ nhan_xet_bctc }}")

    doc.save("Template_Test.docx")
    print("1. Đã tạo xong Form mẫu: Template_Test.docx")


def dien_du_lieu_thuc_chien():
    """BƯỚC 2: Bơm dữ liệu thật vào Form mẫu"""
    # ---------------------------------------------------------
    # A. DỮ LIỆU ĐẦU VÀO (Giả lập output từ Engine của mày)
    # ---------------------------------------------------------
    context_text = {
        "ten_cong_ty": "CÔNG TY TNHH QUẢNG CÁO VÀ TỔ CHỨC SỰ KIỆN THỜI NAY",
        "mst_cong_ty": "4300308931",
        "nhan_xet_bctc": (
            "+ Tài khoản 111- Tiền mặt có số dư Nợ cuối kỳ 13.591.322.024 đồng, nhưng phát sinh chi phí lãi vay.\n"
            "+ Tài khoản 131- Phải thu khách hàng lệch 1.350.000 đồng so với tổng TK 511 và TK 3331."
        )
    }

    du_lieu_bang_gtgt = [
        {"stt": 1, "ky": "Quý 1/2025", "hd": "273.926.729", "tk": "353.293.172", "chenh": "79.366.443"},
        {"stt": 2, "ky": "Quý 2/2025", "hd": "239.885.710", "tk": "239.973.710", "chenh": "88.000"},
        {"stt": 3, "ky": "Quý 3/2025", "hd": "150.000.000", "tk": "150.000.000", "chenh": "0"}
    ]

    # ---------------------------------------------------------
    # B. TIẾN HÀNH RENDER VÀ VẼ BẢNG
    # ---------------------------------------------------------
    doc = DocxTemplate("Template_Test.docx")

    # 1. Bơm toàn bộ Text và Đoạn văn (Nhanh và không bao giờ lỗi)
    doc.render(context_text)

    # 2. Móc vào lõi Document để tự sinh thêm dòng cho Bảng (An toàn tuyệt đối)
    if len(doc.docx.tables) > 0:
        bang_mau = doc.docx.tables[0]  # Lấy cái bảng đầu tiên trong Word

        for row_data in du_lieu_bang_gtgt:
            dong_moi = bang_mau.add_row()
            dong_moi.cells[0].text = str(row_data["stt"])
            dong_moi.cells[1].text = str(row_data["ky"])
            dong_moi.cells[2].text = str(row_data["hd"])
            dong_moi.cells[3].text = str(row_data["tk"])
            dong_moi.cells[4].text = str(row_data["chenh"])

    doc.save("KetQua_Test.docx")
    print("2. Đã bơm số liệu và xuất Báo cáo thành công: KetQua_Test.docx")


if __name__ == "__main__":
    # Đảm bảo cài: pip install python-docx docxtpl
    tao_template_mau_sach()
    dien_du_lieu_thuc_chien()