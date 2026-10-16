# module_ho_tro/tao_template_excel.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def tao_file_mau_cdsps(duong_dan_luu):
    """Tạo file Excel mẫu chứa 2 sheet: Hướng dẫn và Dữ liệu"""
    wb = openpyxl.Workbook()

    # ==========================================
    # SHEET 1: HƯỚNG DẪN NHẬP LIỆU (Cực kỳ chi tiết)
    # ==========================================
    ws_hd = wb.active
    ws_hd.title = "HUONG_DAN"
    ws_hd.sheet_properties.tabColor = "10B981"  # Màu xanh lá

    huong_dan = [
        ["HƯỚNG DẪN NHẬP LIỆU BẢNG CÂN ĐỐI SỐ PHÁT SINH (THÔNG TƯ 200)"],
        [""],
        ["1. HỆ THỐNG HỖ TRỢ NẠP TRỰC TIẾP file Excel xuất từ MISA, FAST, Bravo, SmartPro, HTKK..."],
        ["2. Hoặc bạn có thể dán số liệu vào sheet 'DuLieu_CDSPS' của file mẫu này."],
        ["3. Hệ thống tự động nhận diện dòng tiêu đề, nhận diện năm tài chính và lọc bỏ dòng 'Tổng cộng'."],
        ["4. Các cột số tiền (Dư Nợ, Dư Có, Phát sinh) phải là định dạng SỐ."],
        ["5. Tự động kiểm tra 3 cặp cân đối kế toán (Đầu kỳ, Trong kỳ, Cuối kỳ) và phân cấp tài khoản Mẹ - Con."],
        ["6. Chống tính trùng lặp (Anti Double-Counting) giữa tài khoản tổng hợp và tài khoản chi tiết."]
    ]

    for row in huong_dan:
        ws_hd.append(row)

    ws_hd['A1'].font = Font(bold=True, size=14, color="D83B01")
    ws_hd.column_dimensions['A'].width = 80

    # ==========================================
    # SHEET 2: BẢNG DỮ LIỆU ĐỂ DÁN (FORM CHUẨN)
    # ==========================================
    ws_data = wb.create_sheet(title="DuLieu_CDSPS")
    ws_data.sheet_properties.tabColor = "0052D4"

    # Dòng 1: Tiêu đề gộp
    headers_r1 = ["Tài khoản", "", "Số Dư Đầu Kỳ", "", "Số Phát Sinh Trong Kỳ", "", "Số Dư Cuối Kỳ", ""]
    # Dòng 2: Tiêu đề cột chi tiết
    headers_r2 = ["Mã TK", "Tên Tài Khoản", "Nợ", "Có", "Nợ", "Có", "Nợ", "Có"]

    ws_data.append(headers_r1)
    ws_data.append(headers_r2)

    # Định dạng Header
    fill_header = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
    font_header = Font(color="FFFFFF", bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                         bottom=Side(style='thin'))

    for r in [1, 2]:
        for c in range(1, 9):
            cell = ws_data.cell(row=r, column=c)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_thin

    # Gộp ô cho đẹp
    ws_data.merge_cells("A1:B1")
    ws_data.merge_cells("C1:D1")
    ws_data.merge_cells("E1:F1")
    ws_data.merge_cells("G1:H1")

    ws_data.column_dimensions['A'].width = 12
    ws_data.column_dimensions['B'].width = 40
    for col in ['C', 'D', 'E', 'F', 'G', 'H']:
        ws_data.column_dimensions[col].width = 20

    wb.save(duong_dan_luu)