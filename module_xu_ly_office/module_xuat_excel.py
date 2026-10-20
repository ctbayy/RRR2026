# module_xuat_excel.py

import os
from PySide6.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QCheckBox,
                               QPushButton, QFileDialog, QMessageBox, QLabel, QGroupBox)
from PySide6.QtCore import Qt
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import qtawesome as qta  # Bổ sung thư viện Icon
from openpyxl.cell import MergedCell


class DialogXuatExcel(QDialog):
    def __init__(self, main_app, parent=None):
        super().__init__(parent)
        self.main_app = main_app
        self.setWindowTitle("Kết xuất dữ liệu ra Excel")
        self.setFixedSize(450, 510)  # Tăng chiều cao để chứa đủ 7 tùy chọn kết xuất

        # ĐỒNG BỘ TRIỆT ĐỂ SANG ICE BLUE FLUENT DESIGN
        self.setStyleSheet("""
            QDialog { 
                background-color: #EBF3FA; 
                font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, Arial, sans-serif; 
            }
            QLabel { color: #0F172A; }
            QGroupBox { 
                font-weight: bold; 
                color: #1E3A8A; 
                background-color: #F4F9FF; 
                border: 1px solid #C0D6E4; 
                border-radius: 8px; 
                margin-top: 15px; 
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                subcontrol-position: top left;
                padding: 0 5px;
                left: 10px;
            }
            QCheckBox { 
                font-size: 14px; 
                color: #334155;
                padding: 5px;
                font-weight: 500;
            }
            QCheckBox:hover {
                background-color: #E2EEF9;
                border-radius: 5px;
            }
            QPushButton#BtnXuat { 
                background-color: #005FB8; 
                color: white; 
                font-weight: bold; 
                padding: 12px 15px; 
                border-radius: 6px; 
                font-size: 13px;
                border: none;
            }
            QPushButton#BtnXuat:hover { background-color: #004C93; }
            
            QPushButton#BtnHuy { 
                background-color: #E2EEF9; 
                color: #334155; 
                font-weight: bold; 
                padding: 12px 15px; 
                border-radius: 6px; 
                font-size: 13px;
                border: 1px solid #C0D6E4;
            }
            QPushButton#BtnHuy:hover { background-color: #DDEBF7; }
        """)
        self.thiet_lap_ui()

    def thiet_lap_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(25, 20, 25, 20)
        layout.setSpacing(15)

        lbl_info = QLabel("Vui lòng chọn các nội dung muốn kết xuất:")
        lbl_info.setStyleSheet("font-size: 15px; font-weight: bold;")
        layout.addWidget(lbl_info)

        box_chon = QGroupBox("Danh mục dữ liệu")
        box_layout = QVBoxLayout(box_chon)
        box_layout.setSpacing(8)
        box_layout.setContentsMargins(15, 30, 15, 15)

        # Checkbox "Chọn tất cả" và dải phân cách
        self.chk_tat_ca = QCheckBox(" Chọn tất cả")
        self.chk_tat_ca.setStyleSheet("font-weight: bold; color: #005FB8; margin-bottom: 5px;")
        self.chk_tat_ca.setChecked(True)
        self.chk_tat_ca.toggled.connect(self.xu_ly_chon_tat_ca)
        box_layout.addWidget(self.chk_tat_ca)

        from PySide6.QtWidgets import QFrame
        line = QFrame()
        line.setFrameShape(QFrame.Shape.HLine)
        line.setFrameShadow(QFrame.Shadow.Sunken)
        line.setStyleSheet("background-color: #E2EEF9; margin-bottom: 5px;")
        box_layout.addWidget(line)

        # Các tùy chọn kết xuất
        self.chk_kq = QCheckBox(" Kết quả Phân tích & Cảnh báo rủi ro")
        self.chk_to_khai = QCheckBox(" Dữ liệu Tờ khai thuế (XML)")
        self.chk_bc_doanh_so = QCheckBox(" Báo cáo Đối chiếu Bán ra & Mua vào")
        self.chk_hd_ban = QCheckBox(" Dữ liệu Hóa đơn Bán ra")
        self.chk_hd_mua = QCheckBox(" Dữ liệu Hóa đơn Mua vào")
        self.chk_hd_rui_ro = QCheckBox(" Danh sách Hóa đơn Rủi ro")
        self.chk_hd_dac_biet = QCheckBox(" Danh sách Hóa đơn Trạng thái Đặc biệt")

        # Đồng bộ trạng thái: Nếu bỏ chọn 1 cái thì tự gỡ tick "Chọn tất cả"
        for chk in [
            self.chk_kq, self.chk_to_khai, self.chk_bc_doanh_so,
            self.chk_hd_ban, self.chk_hd_mua, self.chk_hd_rui_ro, self.chk_hd_dac_biet
        ]:
            chk.setChecked(True)
            chk.toggled.connect(self._cap_nhat_trang_thai_tat_ca)
            box_layout.addWidget(chk)

        layout.addWidget(box_chon)
        layout.addStretch()

        khung_nut = QHBoxLayout()
        khung_nut.setSpacing(15)

        btn_huy = QPushButton(" Hủy bỏ")
        btn_huy.setObjectName("BtnHuy")
        btn_huy.setIcon(qta.icon('fa5s.times', color='#334155'))
        btn_huy.clicked.connect(self.reject)

        btn_xuat = QPushButton(" Chọn nơi lưu Excel")
        btn_xuat.setObjectName("BtnXuat")
        btn_xuat.setIcon(qta.icon('fa5s.save', color='white'))
        btn_xuat.clicked.connect(self.thuc_hien_xuat)

        khung_nut.addWidget(btn_huy)
        khung_nut.addWidget(btn_xuat)
        layout.addLayout(khung_nut)

    def xu_ly_chon_tat_ca(self, checked):
        """Hàm bật/tắt đồng loạt các tùy chọn khi click vào nút Chọn Tất Cả"""
        for chk in [
            self.chk_kq, self.chk_to_khai, self.chk_bc_doanh_so,
            self.chk_hd_ban, self.chk_hd_mua, self.chk_hd_rui_ro, self.chk_hd_dac_biet
        ]:
            chk.blockSignals(True)
            chk.setChecked(checked)
            chk.blockSignals(False)

    def _cap_nhat_trang_thai_tat_ca(self):
        """Kiểm tra nếu người dùng bỏ tick 1 mục thì tự gỡ tick nút Tất cả, và ngược lại"""
        trang_thai = all([
            self.chk_kq.isChecked(), self.chk_to_khai.isChecked(),
            self.chk_bc_doanh_so.isChecked(), self.chk_hd_ban.isChecked(),
            self.chk_hd_mua.isChecked(), self.chk_hd_rui_ro.isChecked(),
            self.chk_hd_dac_biet.isChecked()
        ])
        self.chk_tat_ca.blockSignals(True)
        self.chk_tat_ca.setChecked(trang_thai)
        self.chk_tat_ca.blockSignals(False)

    def _dinh_dang_tieu_de(self, sheet, danh_sach_cot):
        """Tô màu xanh biển, chữ trắng, in đậm cho dòng tiêu đề"""
        fill_header = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
        font_header = Font(color="FFFFFF", bold=True)
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))

        sheet.append(danh_sach_cot)
        for col_idx in range(1, len(danh_sach_cot) + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_thin

    def _ghi_so_excel(self, sheet, row, col, gia_tri):
        """Hàm helper để tự động nhận dạng số và set format định dạng Excel"""
        cell = sheet.cell(row=row, column=col)
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))
        cell.border = border_thin

        try:
            str_val = str(gia_tri).replace(",", "")
            val_num = float(str_val)
            cell.value = val_num
            cell.number_format = '#,##0' if val_num.is_integer() else '#,##0.00'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        except Exception:
            cell.value = str(gia_tri)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _ghi_du_lieu_thong_minh(self, sheet, row, col, ten_cot, gia_tri):
        """Xử lý định dạng thông minh dựa trên tên cột (Text, Date, Number)"""
        from openpyxl.styles import Alignment, Border, Side
        import re
        from datetime import datetime

        cell = sheet.cell(row=row, column=col)
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))
        cell.border = border_thin

        if gia_tri is None:
            gia_tri = ""
        val_str = str(gia_tri).strip()
        ten_cot_lower = str(ten_cot).lower()

        # 1. NHÓM TEXT BẮT BUỘC (Giữ số 0 ở đầu)
        nhom_text = ["mst", "mã số thuế", "cmnd", "cccd", "hộ chiếu", "số giấy tờ", "điện thoại", "tài khoản"]
        if any(k in ten_cot_lower for k in nhom_text):
            cell.value = val_str
            cell.number_format = '@'
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            return

        # 2. NHÓM NGÀY THÁNG (Chuẩn hóa dd/mm/yyyy)
        if "ngày" in ten_cot_lower or "ngay" in ten_cot_lower:
            if re.match(r'^\d{4}-\d{2}-\d{2}', val_str):
                try:
                    date_obj = datetime.strptime(val_str[:10], '%Y-%m-%d')
                    cell.value = date_obj
                    cell.number_format = 'dd/mm/yyyy'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    return
                except Exception:
                    pass
            elif re.match(r'^\d{2}/\d{2}/\d{4}', val_str):
                cell.value = val_str
                cell.number_format = 'dd/mm/yyyy'
                cell.alignment = Alignment(horizontal="center", vertical="center")
                return

        # 3. NHÓM SỐ HỌC HOẶC MẶC ĐỊNH CHUNG
        try:
            val_num = float(val_str.replace(",", ""))
            cell.value = val_num
            cell.number_format = '#,##0' if val_num.is_integer() else '#,##0.00'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        except Exception:
            cell.value = val_str
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _can_chinh_kich_thuoc_cot(self, sheet):
        """Thuật toán Auto-fit chuẩn UX + Ép Font chuẩn Hành chính (Times New Roman)"""
        import openpyxl
        from openpyxl.styles import Font

        font_chuan = Font(name="Times New Roman", size=11)

        for col_idx in range(1, sheet.max_column + 1):
            max_length = 0
            col_letter = openpyxl.utils.get_column_letter(col_idx)

            for row_idx in range(1, sheet.max_row + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)

                if type(cell).__name__ == 'MergedCell':
                    continue

                if getattr(cell, 'font', None):
                    cell.font = Font(name="Times New Roman",
                                     size=11,
                                     bold=cell.font.bold,
                                     italic=cell.font.italic,
                                     color=cell.font.color)
                else:
                    cell.font = font_chuan

                if col_idx == 1 and getattr(cell, 'font', None) and cell.font.bold:
                    if cell.value and len(str(cell.value)) > 35:
                        continue

                try:
                    if cell.value is not None:
                        lines = str(cell.value).split('\n')
                        for line in lines:
                            if len(line) > max_length:
                                max_length = len(line)
                except Exception:
                    pass

            if max_length == 0:
                adjusted_width = 12
            else:
                adjusted_width = max_length + 2.0

            if adjusted_width < 5:
                adjusted_width = 5
            elif adjusted_width > 55:
                adjusted_width = 55

            sheet.column_dimensions[col_letter].width = adjusted_width

    def thuc_hien_xuat(self):
        nam_pt_goc = self.main_app.giao_dien.cb_nam.currentText()

        if not (self.chk_kq.isChecked() or self.chk_hd_ban.isChecked() or
                self.chk_hd_mua.isChecked() or self.chk_hd_rui_ro.isChecked() or
                self.chk_hd_dac_biet.isChecked() or self.chk_to_khai.isChecked() or
                self.chk_bc_doanh_so.isChecked()):
            QMessageBox.warning(self, "Lỗi", "Vui lòng chọn ít nhất 1 nội dung để xuất!")
            return

        # ===============================================================
        # PHÂN LUỒNG XUẤT EXCEL TỰ ĐỘNG
        # ===============================================================
        if nam_pt_goc != "Tất cả các năm":
            duong_dan_luu, _ = QFileDialog.getSaveFileName(
                self, "Lưu file Excel", f"BaoCao_PhanTich_{nam_pt_goc}.xlsx", "Excel Files (*.xlsx)"
            )
            if not duong_dan_luu:
                return
            danh_sach_nam_xuat = [nam_pt_goc]
            duong_dan_thu_muc = ""
        else:
            QMessageBox.information(
                self, "Xuất tự động hàng loạt",
                "Bạn đang phân tích 'Tất cả các năm'.\nHệ thống sẽ kết xuất tự động mỗi năm thành 1 file Excel riêng biệt.\n\nVui lòng chọn THƯ MỤC LƯU TRỮ!"
            )
            duong_dan_thu_muc = QFileDialog.getExistingDirectory(self, "Chọn thư mục lưu báo cáo")
            if not duong_dan_thu_muc:
                return
            if hasattr(self.main_app, 'kho_ket_qua_nam') and self.main_app.kho_ket_qua_nam:
                danh_sach_nam_xuat = list(self.main_app.kho_ket_qua_nam.keys())
            else:
                danh_sach_nam_xuat = [str(nam_pt_goc)]

        # ===============================================================
        # VÒNG LẶP ĐẺ FILE BÁO CÁO (CHẠY RIÊNG CHO TỪNG NĂM)
        # ===============================================================
        so_file = 0
        from openpyxl.comments import Comment

        for nam in danh_sach_nam_xuat:
            if hasattr(self.main_app, 'kho_ket_qua_nam') and nam in self.main_app.kho_ket_qua_nam:
                self.main_app.ket_qua_cuoi_cung = self.main_app.kho_ket_qua_nam[nam]["ket_qua_list"]
                du_lieu_br_nam = self.main_app.kho_ket_qua_nam[nam].get("du_lieu_ban_ra", [])
                du_lieu_mv_nam = self.main_app.kho_ket_qua_nam[nam].get("du_lieu_mua_vao", [])
            else:
                du_lieu_br_nam = []
                du_lieu_mv_nam = []

            if duong_dan_thu_muc:
                duong_dan_luu = os.path.join(duong_dan_thu_muc, f"BaoCao_PhanTich_Nam_{nam}.xlsx")

            try:
                wb = openpyxl.Workbook()
                wb.remove(wb.active)
                border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                                     top=Side(style='thin'), bottom=Side(style='thin'))

                # =========================================================
                # 1. XUẤT KẾT QUẢ PHÂN TÍCH
                # =========================================================
                if self.chk_kq.isChecked():
                    ds_kq = getattr(self.main_app, 'ket_qua_cuoi_cung', [])

                    def _dich_ten_chi_tieu(chuoi_goc):
                        try:
                            from module_doc_xml.tu_dien_chi_tieu import TU_DIEN_TO_KHAI
                        except Exception:
                            TU_DIEN_TO_KHAI = {}
                        if not chuoi_goc:
                            return ""
                        chuoi_str = str(chuoi_goc)
                        if "Chỉ tiêu" in chuoi_str or "=" in chuoi_str:
                            return chuoi_str
                        if "|" in chuoi_str:
                            ma_ct = chuoi_str.split("|")[-1].replace("]", "").strip()
                        else:
                            ma_ct = chuoi_str.replace("[", "").replace("]", "").strip()

                        for form, form_data in TU_DIEN_TO_KHAI.items():
                            if isinstance(form_data, dict):
                                for phan_he, phan_he_data in form_data.items():
                                    if isinstance(phan_he_data, dict) and ma_ct in phan_he_data:
                                        return f"[{ma_ct}] {phan_he_data[ma_ct]}"
                                    elif isinstance(phan_he_data, str) and phan_he == ma_ct:
                                        return f"[{ma_ct}] {phan_he_data}"
                        return f"Chỉ tiêu {ma_ct}"

                    ds_cheo, ds_bctc, ds_ngang = [], [], []
                    for r in ds_kq:
                        nhom = str(r.get('nhom', '')).lower()
                        if "chéo" in nhom:
                            ds_cheo.append(r)
                        elif "tài chính" in nhom or "bctc" in nhom:
                            ds_bctc.append(r)
                        elif "ngang" in nhom or "tỷ lệ" in nhom or "ty le" in nhom:
                            ds_ngang.append(r)
                        else:
                            ds_bctc.append(r)

                    def _ve_bang_7_cot(ten_sheet, data_list):
                        if not data_list:
                            return
                        ws = wb.create_sheet(ten_sheet)
                        cols = ["Nhóm Phân Tích", "Tên Phép Kiểm Tra (Logic)", "Số liệu Vế Trái (A)",
                                "Số liệu Vế Phải (B)", "Chênh lệch (A - B)", "Đánh giá", "Khuyến nghị / Thông điệp"]
                        self._dinh_dang_tieu_de(ws, cols)

                        ws.column_dimensions['A'].width = 15
                        ws.column_dimensions['B'].width = 35
                        ws.column_dimensions['C'].width = 25
                        ws.column_dimensions['D'].width = 25
                        ws.column_dimensions['E'].width = 20
                        ws.column_dimensions['F'].width = 18
                        ws.column_dimensions['G'].width = 40

                        align_wrap = Alignment(vertical="center", wrap_text=True)
                        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

                        for r_idx, r in enumerate(data_list, start=2):
                            is_rui_ro = r.get('is_rui_ro', False)
                            so_chenh_lech = r.get('ket_qua_so', 0)
                            tt_text = r.get('trang_thai', '')

                            mau_nen = PatternFill(
                                start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
                            ) if r_idx % 2 == 0 else PatternFill(
                                start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
                            )

                            ws.cell(row=r_idx, column=1, value=r.get('nhom', '')).alignment = align_wrap
                            ws.cell(row=r_idx, column=2, value=r.get('ten_phan_tich', '')).alignment = align_wrap

                            self._ghi_so_excel(ws, r_idx, 3, r.get('gia_tri_a', 0))
                            ws.cell(row=r_idx, column=3).comment = Comment(
                                _dich_ten_chi_tieu(r.get('ve_a_goc', '')), "Chi tiết"
                            )
                            ws.cell(row=r_idx, column=3).font = Font(bold=True)

                            if r.get('is_ab'):
                                self._ghi_so_excel(ws, r_idx, 4, r.get('gia_tri_b', 0))
                                ws.cell(row=r_idx, column=4).comment = Comment(
                                    _dich_ten_chi_tieu(r.get('ve_b_goc', '')), "Chi tiết"
                                )
                            else:
                                ws.cell(row=r_idx, column=4, value="---").alignment = align_center
                            ws.cell(row=r_idx, column=4).font = Font(bold=True)

                            self._ghi_so_excel(ws, r_idx, 5, so_chenh_lech)
                            if is_rui_ro:
                                ws.cell(row=r_idx, column=5).font = Font(color="DC2626", bold=True)
                            else:
                                ws.cell(row=r_idx, column=5).font = Font(bold=True)

                            c_tt = ws.cell(row=r_idx, column=6, value=tt_text)
                            c_tt.alignment = align_center
                            if is_rui_ro:
                                c_tt.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
                                c_tt.font = Font(color="FFFFFF", bold=True)
                            elif "Thiếu" in str(so_chenh_lech) or "Không đủ" in tt_text:
                                c_tt.fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
                                c_tt.font = Font(color="FFFFFF", bold=True)
                            else:
                                c_tt.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
                                c_tt.font = Font(color="FFFFFF", bold=True)

                            ws.cell(row=r_idx, column=7, value=r.get('thong_diep', '')).alignment = align_wrap

                            for c_idx in range(1, 8):
                                c = ws.cell(row=r_idx, column=c_idx)
                                if c_idx != 6:
                                    c.fill = mau_nen
                                c.border = border_thin

                    _ve_bang_7_cot("PT Báo cáo Tài chính", ds_bctc)
                    _ve_bang_7_cot("Phân tích Ngang", ds_ngang)

                    # VẼ ĐỐI CHIẾU CHÉO
                    if ds_cheo:
                        import re
                        ws_cheo = wb.create_sheet("Đối chiếu chéo Tờ khai")

                        def _lay_ma_form_excel(chuoi_goc):
                            if not chuoi_goc:
                                return "Dữ liệu"
                            match = re.search(
                                r'(01/GTGT|02/GTGT|03/TNDN|04/TNDN|05/QTT[-_]TNCN|05/KK[-_]TNCN|BCTC[-_]TT133|BCTC[-_]TT200|01/TAIN|02/TAIN|TAIN)',
                                str(chuoi_goc)
                            )
                            if match:
                                return match.group(1).replace('_', '-')
                            if "HD_" in str(chuoi_goc) or "bán ra" in str(chuoi_goc).lower():
                                return "Hóa đơn Bán ra"
                            if "HD_MUA_VAO" in str(chuoi_goc) or "mua vào" in str(chuoi_goc).lower():
                                return "Hóa đơn Mua vào"
                            return "Dữ liệu"

                        def format_header_cheo(ma_form):
                            if "hóa đơn" in ma_form.lower() or "sổ cái" in ma_form.lower() or "dữ liệu" in ma_form.lower():
                                return ma_form
                            return f"Tờ khai {ma_form}"

                        dict_nhom = {}
                        for r in ds_cheo:
                            form_a = _lay_ma_form_excel(r.get('ve_a_goc', ''))
                            form_b = _lay_ma_form_excel(r.get('ve_b_goc', ''))
                            ten_cap = f"{form_a} VÀ {form_b}" if r.get('is_ab') else f"{form_a}"
                            if ten_cap not in dict_nhom:
                                dict_nhom[ten_cap] = []
                            dict_nhom[ten_cap].append(r)

                        r_idx = 1
                        align_wrap = Alignment(vertical="center", wrap_text=True)
                        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

                        ws_cheo.column_dimensions['A'].width = 8
                        ws_cheo.column_dimensions['B'].width = 35
                        ws_cheo.column_dimensions['C'].width = 25
                        ws_cheo.column_dimensions['D'].width = 25
                        ws_cheo.column_dimensions['E'].width = 20
                        ws_cheo.column_dimensions['F'].width = 18
                        ws_cheo.column_dimensions['G'].width = 40

                        for ten_cap, ds_loi in dict_nhom.items():
                            tk_a = format_header_cheo(_lay_ma_form_excel(ds_loi[0].get('ve_a_goc', '')))
                            tk_b = format_header_cheo(_lay_ma_form_excel(ds_loi[0].get('ve_b_goc', ''))) if ds_loi[0].get('is_ab') else "Dữ liệu Vế Phải (B)"
                            cols = ["STT", "Nội dung đối chiếu", tk_a, tk_b, "Chênh lệch (A-B)", "Đánh giá",
                                    "Khuyến nghị / Thông điệp"]

                            fill_header = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
                            font_header = Font(color="0F172A", bold=True)
                            for c_i, c_name in enumerate(cols, start=1):
                                c = ws_cheo.cell(row=r_idx, column=c_i, value=c_name)
                                c.fill = fill_header
                                c.font = font_header
                                c.alignment = align_center
                                c.border = border_thin
                            r_idx += 1

                            for stt, r in enumerate(ds_loi, start=1):
                                is_rui_ro = r.get('is_rui_ro', False)
                                so_chenh_lech = r.get('ket_qua_so', 0)
                                tt_text = r.get('trang_thai', '')

                                mau_nen = PatternFill(
                                    start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
                                ) if stt % 2 != 0 else PatternFill(
                                    start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
                                )

                                ws_cheo.cell(row=r_idx, column=1, value=stt).alignment = align_center
                                ws_cheo.cell(row=r_idx, column=2, value=r.get('ten_phan_tich', '')).alignment = align_wrap

                                self._ghi_so_excel(ws_cheo, r_idx, 3, r.get('gia_tri_a', 0))
                                ws_cheo.cell(row=r_idx, column=3).comment = Comment(
                                    _dich_ten_chi_tieu(r.get('ve_a_goc', '')), "Chi tiết"
                                )
                                ws_cheo.cell(row=r_idx, column=3).font = Font(bold=True)

                                if r.get('is_ab'):
                                    self._ghi_so_excel(ws_cheo, r_idx, 4, r.get('gia_tri_b', 0))
                                    ws_cheo.cell(row=r_idx, column=4).comment = Comment(
                                        _dich_ten_chi_tieu(r.get('ve_b_goc', '')), "Chi tiết"
                                    )
                                else:
                                    ws_cheo.cell(row=r_idx, column=4, value="---").alignment = align_center
                                ws_cheo.cell(row=r_idx, column=4).font = Font(bold=True)

                                self._ghi_so_excel(ws_cheo, r_idx, 5, so_chenh_lech)
                                if is_rui_ro:
                                    ws_cheo.cell(row=r_idx, column=5).font = Font(color="DC2626", bold=True)
                                else:
                                    ws_cheo.cell(row=r_idx, column=5).font = Font(bold=True)

                                c_tt = ws_cheo.cell(row=r_idx, column=6, value=tt_text)
                                c_tt.alignment = align_center
                                if is_rui_ro:
                                    c_tt.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
                                    c_tt.font = Font(color="FFFFFF", bold=True)
                                elif "Thiếu" in str(so_chenh_lech) or "Không đủ" in tt_text:
                                    c_tt.fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
                                    c_tt.font = Font(color="FFFFFF", bold=True)
                                else:
                                    c_tt.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
                                    c_tt.font = Font(color="FFFFFF", bold=True)

                                ws_cheo.cell(row=r_idx, column=7, value=r.get('thong_diep', '')).alignment = align_wrap

                                for c_idx in range(1, 8):
                                    c = ws_cheo.cell(row=r_idx, column=c_idx)
                                    if c_idx != 6:
                                        c.fill = mau_nen
                                    c.border = border_thin
                                r_idx += 1

                # =========================================================
                # 1.5. XUẤT BÁO CÁO ĐỐI CHIẾU DOANH SỐ (MÓC TỪ RAM RA)
                # =========================================================
                if getattr(self, 'chk_bc_doanh_so', None) and self.chk_bc_doanh_so.isChecked():
                    ws_bc = wb.create_sheet("Báo cáo Đối chiếu HĐ")
                    dong_hien_tai = 1

                    def _ve_bang_bao_cao_excel(ws, start_row, tieu_de, data):
                        if not data:
                            return start_row
                        ws.cell(row=start_row, column=1, value=tieu_de).font = Font(bold=True, size=11, color="000000")
                        start_row += 1

                        fill_header = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
                        font_header = Font(bold=True)
                        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

                        headers_r1 = ["STT", "Kỳ", "Hóa đơn điện tử", "", "Hồ sơ khai thuế", "",
                                      "Chênh lệch (KK-Hóa đơn)", ""]
                        headers_r2 = ["", "", "Giá trị chưa thuế", "Thuế GTGT", "Giá trị chưa thuế", "Thuế GTGT",
                                      "Giá trị chưa thuế", "Thuế GTGT"]

                        for r_offset, h_list in enumerate([headers_r1, headers_r2]):
                            for i, text in enumerate(h_list, 1):
                                c = ws.cell(row=start_row + r_offset, column=i, value=text)
                                c.fill = fill_header
                                c.font = font_header
                                c.alignment = align_center
                                c.border = border_thin

                        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 1, end_column=1)
                        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row + 1, end_column=2)
                        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row, end_column=4)
                        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row, end_column=6)
                        ws.merge_cells(start_row=start_row, start_column=7, end_row=start_row, end_column=8)

                        row_idx = start_row + 2
                        tong_hd_truoc = tong_hd_thue = tong_tk_truoc = tong_tk_thue = 0

                        for stt, r_data in enumerate(data, 1):
                            hd_truoc, hd_thue = r_data.get('hd_truoc', 0), r_data.get('hd_thue', 0)
                            tk_truoc, tk_thue = r_data.get('tk_truoc', 0), r_data.get('tk_thue', 0)
                            tong_hd_truoc += hd_truoc
                            tong_hd_thue += hd_thue
                            tong_tk_truoc += tk_truoc
                            tong_tk_thue += tk_thue

                            ws.cell(row=row_idx, column=1, value=stt).alignment = align_center
                            ws.cell(row=row_idx, column=2, value=r_data.get('ky', '')).alignment = align_center
                            self._ghi_so_excel(ws, row_idx, 3, hd_truoc if hd_truoc else "-")
                            self._ghi_so_excel(ws, row_idx, 4, hd_thue if hd_thue else "-")
                            self._ghi_so_excel(ws, row_idx, 5, tk_truoc if tk_truoc else "-")
                            self._ghi_so_excel(ws, row_idx, 6, tk_thue if tk_thue else "-")
                            self._ghi_so_excel(ws, row_idx, 7, (tk_truoc - hd_truoc) if (tk_truoc - hd_truoc) else "-")
                            self._ghi_so_excel(ws, row_idx, 8, (tk_thue - hd_thue) if (tk_thue - hd_thue) else "-")

                            for c_fix in range(3, 9):
                                if ws.cell(row=row_idx, column=c_fix).value == "-":
                                    ws.cell(row=row_idx, column=c_fix).alignment = Alignment(
                                        horizontal="right", vertical="center"
                                    )

                            for c in range(1, 9):
                                ws.cell(row=row_idx, column=c).border = border_thin
                            row_idx += 1

                        ws.cell(row=row_idx, column=1, value="Tổng").alignment = align_center
                        ws.cell(row=row_idx, column=1).font = Font(bold=True)
                        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)

                        self._ghi_so_excel(ws, row_idx, 3, tong_hd_truoc if tong_hd_truoc else "-")
                        self._ghi_so_excel(ws, row_idx, 4, tong_hd_thue if tong_hd_thue else "-")
                        self._ghi_so_excel(ws, row_idx, 5, tong_tk_truoc if tong_tk_truoc else "-")
                        self._ghi_so_excel(ws, row_idx, 6, tong_tk_thue if tong_tk_thue else "-")
                        self._ghi_so_excel(
                            ws, row_idx, 7, (tong_tk_truoc - tong_hd_truoc) if (tong_tk_truoc - tong_hd_truoc) else "-"
                        )
                        self._ghi_so_excel(
                            ws, row_idx, 8, (tong_tk_thue - tong_hd_thue) if (tong_tk_thue - tong_hd_thue) else "-"
                        )

                        for c in range(1, 9):
                            ws.cell(row=row_idx, column=c).border = border_thin
                            ws.cell(row=row_idx, column=c).font = Font(bold=True)
                            if ws.cell(row=row_idx, column=c).value == "-":
                                ws.cell(row=row_idx, column=c).alignment = Alignment(
                                    horizontal="right", vertical="center"
                                )

                        return row_idx + 3

                    dong_hien_tai = _ve_bang_bao_cao_excel(
                        ws_bc, dong_hien_tai,
                        "1.1.2. Chênh lệch doanh số bán ra trên hồ sơ khai thuế và dữ liệu hóa đơn điện tử:",
                        du_lieu_br_nam
                    )
                    dong_hien_tai = _ve_bang_bao_cao_excel(
                        ws_bc, dong_hien_tai,
                        "1.1.3. Chênh lệch giá trị HHDV mua vào, thuế GTGT đầu vào được khấu trừ trên hồ sơ khai thuế và dữ liệu hóa đơn điện tử:",
                        du_lieu_mv_nam
                    )

                    if dong_hien_tai == 1:
                        wb.remove(ws_bc)
                    else:
                        ws_bc.column_dimensions['A'].width = 8
                        ws_bc.column_dimensions['B'].width = 12
                        for c in ['C', 'D', 'E', 'F', 'G', 'H']:
                            ws_bc.column_dimensions[c].width = 20

                # =========================================================
                # 2. XUẤT TỜ KHAI (XML) THEO NĂM
                # =========================================================
                if self.chk_to_khai.isChecked():
                    dict_to_khai = getattr(self.main_app, 'du_lieu_he_thong_to_khai', {})

                    def _sort_ky(ky_str):
                        try:
                            t_q, nam_k = ky_str.split('/')
                            if 'Q' in t_q.upper() or 'q' in t_q.lower():
                                thang = int(t_q.upper().replace('Q', '')) * 3 - 2
                            else:
                                thang = int(t_q)
                            return int(nam_k) * 100 + thang
                        except Exception:
                            return 999999

                    for ma_form, dict_ky in dict_to_khai.items():
                        ten_sheet = str(ma_form).replace("/", "_").replace("\\", "_")[:31]
                        ws_tk = wb.create_sheet(f"TKhai {ten_sheet}")

                        danh_sach_ky = sorted(list(dict_ky.keys()), key=_sort_ky)
                        danh_sach_ky = [ky for ky in danh_sach_ky if str(nam) in str(ky)]

                        if not danh_sach_ky:
                            wb.remove(ws_tk)
                            continue

                        ma_tran, ds_bang_ke, ds_bctc = {}, [], []
                        for ky in danh_sach_ky:
                            data = dict_ky[ky]
                            for ma_ct, val_dict in data.get("ChiTieuChinh", {}).items():
                                phan_he = "Chỉ tiêu chính"
                                if phan_he not in ma_tran:
                                    ma_tran[phan_he] = {}
                                if ma_ct not in ma_tran[phan_he]:
                                    ma_tran[phan_he][ma_ct] = {"DienGiai": val_dict.get("DienGiai", ""), "GiaTri": {}}
                                ma_tran[phan_he][ma_ct]["GiaTri"][ky] = val_dict.get("GiaTri", 0)

                            for ph in ["B01_CDKT", "B02_KQKD", "B03_LCTT", "B04_CDTK"]:
                                if ph in data and data[ph]:
                                    ds_bctc.append({"Ky": ky, "PhanHe": ph, "DuLieu": data[ph]})

                            for ten_pl, pl_data in data.get("PhuLuc", {}).items():
                                is_bang_ke = pl_data.get("IsBangKe", False)
                                du_lieu_pl = pl_data.get("DuLieu", {})
                                if not is_bang_ke:
                                    if ten_pl not in ma_tran:
                                        ma_tran[ten_pl] = {}
                                    for ma_ct, val_dict in du_lieu_pl.items():
                                        if ma_ct not in ma_tran[ten_pl]:
                                            ma_tran[ten_pl][ma_ct] = {"DienGiai": val_dict.get("DienGiai", ""), "GiaTri": {}}
                                        ma_tran[ten_pl][ma_ct]["GiaTri"][ky] = val_dict.get("GiaTri", 0)
                                else:
                                    ds_bang_ke.append({"Ky": ky, "TenPL": ten_pl, "DuLieu": du_lieu_pl})

                        fill_section = PatternFill(start_color="E2EBF5", end_color="E2EBF5", fill_type="solid")
                        font_bold = Font(bold=True)
                        r_idx = 1

                        if ma_tran:
                            cols_header = ["Phân Hệ / Phụ Lục", "Mã CT", "Diễn Giải"] + danh_sach_ky + ["TỔNG / CHỐT SỐ"]
                            fill_header = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
                            font_header = Font(color="FFFFFF", bold=True)
                            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

                            for c_i, c_name in enumerate(cols_header, start=1):
                                c = ws_tk.cell(row=r_idx, column=c_i, value=c_name)
                                c.fill = fill_header
                                c.font = font_header
                                c.alignment = align_center
                                c.border = border_thin
                            r_idx += 1

                            ws_tk.column_dimensions['A'].width = 18
                            ws_tk.column_dimensions['B'].width = 10
                            ws_tk.column_dimensions['C'].width = 45
                            for i in range(len(danh_sach_ky) + 1):
                                ws_tk.column_dimensions[openpyxl.utils.get_column_letter(4 + i)].width = 20

                            for phan_he, dict_ct in ma_tran.items():
                                ws_tk.cell(row=r_idx, column=1, value=phan_he).font = font_bold
                                for c in range(1, len(cols_header) + 1):
                                    ws_tk.cell(row=r_idx, column=c).fill = fill_section
                                    ws_tk.cell(row=r_idx, column=c).border = border_thin
                                r_idx += 1

                                for ma_ct, info in dict_ct.items():
                                    ws_tk.cell(row=r_idx, column=1, value="").border = border_thin
                                    ws_tk.cell(row=r_idx, column=2, value=ma_ct).border = border_thin
                                    ws_tk.cell(row=r_idx, column=3, value=info["DienGiai"]).border = border_thin

                                    mang_gia_tri_so = []
                                    for c_idx, ky in enumerate(danh_sach_ky):
                                        val = info["GiaTri"].get(ky, 0)
                                        self._ghi_so_excel(ws_tk, r_idx, 4 + c_idx, val)
                                        try:
                                            mang_gia_tri_so.append(float(str(val).replace(",", "")))
                                        except Exception:
                                            mang_gia_tri_so.append(0.0)

                                    tong_so = 0.0
                                    if mang_gia_tri_so:
                                        if str(ma_ct).strip().lower() == "ct22":
                                            tong_so = mang_gia_tri_so[0]
                                        elif str(ma_ct).strip().lower() == "ct43":
                                            tong_so = mang_gia_tri_so[-1]
                                        else:
                                            tong_so = sum(mang_gia_tri_so)

                                    cell_tong = ws_tk.cell(row=r_idx, column=4 + len(danh_sach_ky))
                                    cell_tong.value = tong_so
                                    cell_tong.number_format = '#,##0' if float(tong_so).is_integer() else '#,##0.00'
                                    cell_tong.border = border_thin
                                    cell_tong.font = font_bold
                                    r_idx += 1

                        if ds_bang_ke:
                            r_idx += 2
                            for bk in ds_bang_ke:
                                cell_title = ws_tk.cell(
                                    row=r_idx, column=1, value=f"--- {bk['TenPL']} (Kỳ: {bk['Ky']}) ---"
                                )
                                cell_title.font = Font(bold=True, color="0052D4")
                                r_idx += 1

                                du_lieu_dong = bk['DuLieu']
                                if not du_lieu_dong:
                                    continue

                                dong_dau = list(du_lieu_dong.values())[0]
                                danh_sach_den = ["Mã CT", "Chỉ tiêu ct09a_ma", "Chỉ tiêu ct12_ma", "Chỉ tiêu ct14_ma",
                                                 "Chỉ tiêu ct09_ma"]
                                cols_bk = [c for c in list(dong_dau.keys()) if str(c).strip() not in danh_sach_den]

                                for c_i, c_name in enumerate(cols_bk):
                                    c = ws_tk.cell(row=r_idx, column=1 + c_i, value=c_name)
                                    c.fill = fill_section
                                    c.font = font_bold
                                    c.border = border_thin
                                r_idx += 1

                                for ma_dong, row_data in du_lieu_dong.items():
                                    for c_i, c_name in enumerate(cols_bk):
                                        val = row_data.get(c_name, "")
                                        self._ghi_du_lieu_thong_minh(ws_tk, r_idx, 1 + c_i, c_name, val)
                                    r_idx += 1
                                r_idx += 1

                        if ds_bctc:
                            r_idx += 2
                            for bctc in ds_bctc:
                                cell_title = ws_tk.cell(
                                    row=r_idx, column=1,
                                    value=f"--- BÁO CÁO TÀI CHÍNH: {bctc['PhanHe']} (Kỳ: {bctc['Ky']}) ---"
                                )
                                cell_title.font = Font(bold=True, color="D83B01")
                                r_idx += 1

                                du_lieu_bctc = bctc['DuLieu']
                                if not du_lieu_bctc:
                                    continue

                                first_key = list(du_lieu_bctc.keys())[0]
                                first_item = du_lieu_bctc[first_key]
                                cot_gia_tri = [k for k in first_item.keys() if k != "DienGiai"]

                                cols_bctc = ["Mã Chỉ Tiêu", "Nội Dung Diễn Giải"] + cot_gia_tri
                                fill_header = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
                                font_header = Font(color="FFFFFF", bold=True)

                                for c_i, c_name in enumerate(cols_bctc):
                                    c = ws_tk.cell(row=r_idx, column=1 + c_i, value=c_name)
                                    c.fill = fill_header
                                    c.font = font_header
                                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                    c.border = border_thin
                                r_idx += 1

                                for ma_ct, row_data in du_lieu_bctc.items():
                                    c_ma = ws_tk.cell(row=r_idx, column=1, value=ma_ct)
                                    c_ma.alignment = Alignment(horizontal="center", vertical="center")
                                    c_ma.border = border_thin

                                    c_dg = ws_tk.cell(row=r_idx, column=2, value=row_data.get("DienGiai", ""))
                                    c_dg.border = border_thin
                                    c_dg.alignment = Alignment(wrap_text=True)

                                    for i, ten_cot in enumerate(cot_gia_tri):
                                        val = row_data.get(ten_cot, "0")
                                        self._ghi_so_excel(ws_tk, r_idx, 3 + i, val)
                                    r_idx += 1
                                r_idx += 1

                # =========================================================
                # 3. XUẤT HÓA ĐƠN BÁN RA / MUA VÀO / RỦI RO / TRẠNG THÁI ĐẶC BIỆT
                # =========================================================
                def xuat_bang_hoa_don(ten_sheet, du_lieu_hd, nam_loc):
                    import pandas as pd
                    if du_lieu_hd is None:
                        return
                    if isinstance(du_lieu_hd, pd.DataFrame) and du_lieu_hd.empty:
                        return
                    if isinstance(du_lieu_hd, dict) and not du_lieu_hd:
                        return

                    if isinstance(du_lieu_hd, pd.DataFrame):
                        if '_Nam' in du_lieu_hd.columns:
                            mask_nam = (
                                (du_lieu_hd['_Nam'].astype(str) == str(nam_loc))
                                | (du_lieu_hd['_Nam'].astype(str).isin(['', 'nan', 'NaT']))
                            )
                            du_lieu_hd = du_lieu_hd[mask_nam].copy()
                            if du_lieu_hd.empty:
                                return

                    ws_hd = wb.create_sheet(ten_sheet)
                    tu_khoa_so = [
                        "tiền", "thuế", "đơn giá", "số lượng", "giá trị", "tỷ giá",
                        "tien", "thue", "dongia", "soluong", "giatri", "tygia"
                    ]
                    tu_khoa_text_bat_buoc = [
                        "mst", "mãsốthuế", "masothue", "cccd", "cmnd",
                        "sốhđ", "sốhóađơn", "sohoadon", "kýhiệu", "kyhieu", "tàikhoản"
                    ]

                    fill_do = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    fill_vang = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
                    tu_khoa_tms_do = ["không hoạt động tại địa chỉ", "bỏ địa chỉ", "chờ xác minh", "mất tích"]
                    tu_khoa_tms_vang = ["ngừng hoạt động", "tạm nghỉ", "tạm ngừng", "chấm dứt"]
                    tu_khoa_hd_do = ["hủy", "xóa bỏ", "bị xóa bỏ", "bị hủy", "bị thay thế"]
                    tu_khoa_hd_vang = ["thay thế", "điều chỉnh", "bị điều chỉnh", "chiết khấu", "trả lại"]

                    if isinstance(du_lieu_hd, pd.DataFrame):
                        cols = [
                            col for col in du_lieu_hd.columns
                            if not str(col).startswith("_") and str(col) != "Row_Color"
                        ]
                        self._dinh_dang_tieu_de(ws_hd, cols)

                        idx_tms = cols.index('TrangThaiTMS') if 'TrangThaiTMS' in cols else -1
                        idx_rr = cols.index('Đối chiếu HĐ rủi ro') if 'Đối chiếu HĐ rủi ro' in cols else -1
                        cot_tt_indices = [
                            cols.index(c) for c in [
                                'TrangThaiHoaDon', 'Trạng thái hóa đơn', 'Tinhchat', 'Tính chất',
                                'PhanLoaiTrangThai', 'Nhóm Trạng Thái Đặc Biệt'
                            ] if c in cols
                        ]

                        for r_idx, row in enumerate(du_lieu_hd[cols].itertuples(index=False), start=2):
                            tt_tms = str(row[idx_tms]).lower() if idx_tms != -1 and pd.notna(row[idx_tms]) else ""
                            tt_rr = str(row[idx_rr]).lower() if idx_rr != -1 and pd.notna(row[idx_rr]) else ""
                            tt_hd_goc = " | ".join(
                                str(row[i]).lower() for i in cot_tt_indices if pd.notna(row[i])
                            ) if cot_tt_indices else ""

                            is_do = (
                                "thuộc ds hóa đơn rủi ro" in tt_rr
                                or any(k in tt_tms for k in tu_khoa_tms_do)
                                or any(k in tt_hd_goc for k in tu_khoa_hd_do)
                            )
                            is_vang = (not is_do) and (
                                any(k in tt_tms for k in tu_khoa_tms_vang)
                                or any(k in tt_hd_goc for k in tu_khoa_hd_vang)
                            )

                            for c_idx, val in enumerate(row, start=1):
                                col_name = cols[c_idx - 1]
                                col_lower = str(col_name).lower().replace(" ", "").replace("_", "")

                                is_id_col = any(k in col_lower for k in tu_khoa_text_bat_buoc)
                                is_num_col = (not is_id_col) and any(k in col_lower for k in tu_khoa_so)

                                if is_num_col:
                                    self._ghi_so_excel(ws_hd, r_idx, c_idx, val)
                                else:
                                    val_str = "" if pd.isna(val) else str(val).strip()
                                    if is_id_col and val_str.endswith(".0"):
                                        val_str = val_str[:-2]
                                    cell = ws_hd.cell(row=r_idx, column=c_idx, value=val_str)
                                    if is_id_col:
                                        cell.number_format = '@'
                                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                    else:
                                        cell.alignment = Alignment(vertical="center", wrap_text=True)
                                    cell.border = border_thin

                                if is_do:
                                    ws_hd.cell(row=r_idx, column=c_idx).fill = fill_do
                                elif is_vang:
                                    ws_hd.cell(row=r_idx, column=c_idx).fill = fill_vang
                    else:
                        dong_mau = list(du_lieu_hd.values())[0]
                        cols = list(dong_mau.keys())
                        self._dinh_dang_tieu_de(ws_hd, cols)

                        for r_idx, row_data in enumerate(du_lieu_hd.values(), start=2):
                            for c_idx, col_name in enumerate(cols, start=1):
                                val = row_data.get(col_name, "")
                                col_lower = str(col_name).lower().replace(" ", "").replace("_", "")
                                is_id_col = any(k in col_lower for k in tu_khoa_text_bat_buoc)
                                is_num_col = (not is_id_col) and any(k in col_lower for k in tu_khoa_so)

                                if is_num_col:
                                    self._ghi_so_excel(ws_hd, r_idx, c_idx, val)
                                else:
                                    cell = ws_hd.cell(row=r_idx, column=c_idx, value=str(val))
                                    if is_id_col:
                                        cell.number_format = '@'
                                    cell.border = border_thin

                if self.chk_hd_ban.isChecked():
                    xuat_bang_hoa_don("HĐ Bán Ra", self.main_app.dict_hd_ban_ra, nam)
                if self.chk_hd_mua.isChecked():
                    xuat_bang_hoa_don("HĐ Mua Vào", self.main_app.dict_hd_mua_vao, nam)

                # 3.1. XUẤT RIÊNG SHEET HÓA ĐƠN RỦI RO
                if getattr(self, 'chk_hd_rui_ro', None) and self.chk_hd_rui_ro.isChecked():
                    import pandas as pd
                    df_mv_goc = getattr(self.main_app, 'dict_hd_mua_vao', None)
                    if isinstance(df_mv_goc, pd.DataFrame) and not df_mv_goc.empty:
                        if 'Đối chiếu HĐ rủi ro' not in df_mv_goc.columns:
                            try:
                                from module_doc_hoadon.hoadon_ruiro import HoaDonRuiRoManager
                                rui_ro_mgr = getattr(self.main_app, 'ruiro_mgr', None) or HoaDonRuiRoManager()
                                df_mv_goc, _ = rui_ro_mgr.check_bang_ke_01_1(df_mv_goc)
                                self.main_app.dict_hd_mua_vao = df_mv_goc
                            except Exception:
                                pass
                        if 'Đối chiếu HĐ rủi ro' in df_mv_goc.columns:
                            df_hd_rr = df_mv_goc[df_mv_goc['Đối chiếu HĐ rủi ro'] == "Thuộc DS hóa đơn rủi ro"].copy()
                            if not df_hd_rr.empty:
                                xuat_bang_hoa_don("HĐ Rủi Ro", df_hd_rr, nam)

                # 3.2. XUẤT RIÊNG SHEET HÓA ĐƠN TRẠNG THÁI ĐẶC BIỆT
                if getattr(self, 'chk_hd_dac_biet', None) and self.chk_hd_dac_biet.isChecked():
                    import pandas as pd
                    from module_doc_hoadon.invoice_processor import InvoiceProcessor
                    ds_db = []
                    df_br_goc = getattr(self.main_app, 'dict_hd_ban_ra', None)
                    df_mv_goc = getattr(self.main_app, 'dict_hd_mua_vao', None)

                    if isinstance(df_br_goc, pd.DataFrame) and not df_br_goc.empty:
                        df_db_br = InvoiceProcessor.loc_hoa_don_trang_thai_dac_biet(df_br_goc, "Bán ra")
                        if not df_db_br.empty:
                            ds_db.append(df_db_br)
                    if isinstance(df_mv_goc, pd.DataFrame) and not df_mv_goc.empty:
                        df_db_mv = InvoiceProcessor.loc_hoa_don_trang_thai_dac_biet(df_mv_goc, "Mua vào")
                        if not df_db_mv.empty:
                            ds_db.append(df_db_mv)

                    if ds_db:
                        df_tong_db = pd.concat(ds_db, ignore_index=True)
                        xuat_bang_hoa_don("HĐ TT Đặc Biệt", df_tong_db, nam)

                # --- LƯU CHỐT TỪNG FILE EXCEL CỦA MỖI NĂM ---
                if wb.sheetnames:
                    for sheet_name in wb.sheetnames:
                        self._can_chinh_kich_thuoc_cot(wb[sheet_name])
                    wb.save(duong_dan_luu)
                    so_file += 1

            except Exception as e:
                QMessageBox.critical(self, "Lỗi", f"Có lỗi khi tạo file năm {nam}:\n{e}")

        # THÔNG BÁO HOÀN TẤT
        if duong_dan_thu_muc:
            QMessageBox.information(
                self, "Thành công", f"Đã kết xuất hoàn tất {so_file} file Excel vào thư mục:\n{duong_dan_thu_muc}"
            )
        else:
            QMessageBox.information(self, "Thành công", f"Đã kết xuất báo cáo thành công tại:\n{duong_dan_luu}")
        self.accept()