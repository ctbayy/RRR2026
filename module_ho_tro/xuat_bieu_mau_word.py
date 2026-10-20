# -*- coding: utf-8 -*-
"""
ENGINE KẾT XUẤT BIỂU MẪU WORD THANH KIỂM TRA THUẾ - HIGH-FIDELITY XML CLONING & RUN-LEVEL REPLACEMENT
Bảo lưu 100% cấu trúc, font chữ (Times New Roman), kích cỡ (14pt/13pt/11pt), khoảng cách dòng,
thụt đầu dòng và bảng biểu của các biểu mẫu hành chính trong Data/Mau:
1. Phiếu Nhận Xét Hồ Sơ Khai Thuế (PhieuNhanXet.docx)
2. Thông Báo Giải Trình, Bổ Sung Thông Tin Tài Liệu (ThongBaoGiaiTrinh.docx)
3. Tờ Trình Lãnh Đạo Phê Duyệt Lưu Hồ Sơ (To trinh phe duyet luu ho so_BETONG.docx)
"""

import os
import copy
import datetime
import unicodedata
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from module_ho_tro.dinh_dang_so import dinh_dang_so_vn
from module_ho_tro.thu_vien_mau_cau import trich_xuat_bang_gtgt, trich_xuat_bang_tncn, format_tien


THU_MUC_GOC_DU_AN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
THU_MUC_MAU_MAC_DINH = os.path.join(THU_MUC_GOC_DU_AN, "Data", "Mau")


class XuatBieuMauWord:
    def __init__(self, thu_muc_mau=None):
        if thu_muc_mau is None:
            thu_muc_mau = THU_MUC_MAU_MAC_DINH
        elif not os.path.isabs(thu_muc_mau):
            thu_muc_mau = os.path.join(THU_MUC_GOC_DU_AN, thu_muc_mau)

        self.thu_muc_mau = os.path.abspath(thu_muc_mau)
        self.mau_phieu_nhan_xet = os.path.join(self.thu_muc_mau, "PhieuNhanXet.docx")
        self.mau_thong_bao = os.path.join(self.thu_muc_mau, "ThongBaoGiaiTrinh.docx")
        self.mau_to_trinh = os.path.join(self.thu_muc_mau, "To trinh phe duyet luu ho so_BETONG.docx")

    # =========================================================================
    # CÁC HÀM XỬ LÝ RUN & XML BẢO LƯU 100% ĐỊNH DẠNG (FONT, SIZE, INDENT)
    # =========================================================================
    @staticmethod
    def _chuan_hoa_nfc(doc):
        """Chuẩn hóa toàn bộ text trong Document về Unicode NFC để tìm kiếm tiếng Việt chính xác 100%"""
        for p in doc.paragraphs:
            for r in p.runs:
                if r.text:
                    r.text = unicodedata.normalize('NFC', r.text)

        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            if r.text:
                                r.text = unicodedata.normalize('NFC', r.text)

    @staticmethod
    def _set_run_font_and_size(r, font_name="Times New Roman", size_pt=14, is_bold=None, is_italic=None):
        """Cài đặt tường minh font và size qua XML tag (tránh Word fallback sang Calibri)"""
        if is_bold is not None:
            r.bold = is_bold
        if is_italic is not None:
            r.italic = is_italic

        rPr = r._r.get_or_add_rPr()

        # w:rFonts
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)

        # w:sz và w:szCs (half-points: 14pt = 28 half-points)
        sz_val = str(int(size_pt * 2))
        sz = rPr.find(qn('w:sz'))
        if sz is None:
            sz = OxmlElement('w:sz')
            rPr.append(sz)
        sz.set(qn('w:val'), sz_val)

        szCs = rPr.find(qn('w:szCs'))
        if szCs is None:
            szCs = OxmlElement('w:szCs')
            rPr.append(szCs)
        szCs.set(qn('w:val'), sz_val)

    @classmethod
    def _replace_text_across_runs(cls, p, search_str, replace_str):
        """
        Thay thế từ khóa xuyên các Runs trong một Paragraph mà KHÔNG hủy cấu trúc Runs.
        Bảo lưu 100% font chữ, cỡ chữ, in đậm, in nghiêng và paragraph formatting.
        """
        if not p.runs or not search_str:
            return False

        search_norm = unicodedata.normalize('NFC', search_str)
        replace_norm = unicodedata.normalize('NFC', str(replace_str))

        if search_norm == replace_norm:
            return False

        # Đảm bảo các runs đã được normalize NFC
        for r in p.runs:
            if r.text:
                r.text = unicodedata.normalize('NFC', r.text)

        full_text = p.text
        occurrences = full_text.count(search_norm)
        if occurrences == 0:
            return False

        cur_search_pos = 0
        for _ in range(occurrences):
            run_ranges = []
            cur_pos = 0
            for r in p.runs:
                txt_len = len(r.text)
                run_ranges.append((cur_pos, cur_pos + txt_len, r))
                cur_pos += txt_len

            full_text = p.text
            start_idx = full_text.find(search_norm, cur_search_pos)
            if start_idx == -1:
                break
            end_idx = start_idx + len(search_norm)

            i_start, i_end = None, None
            for i, (r_start, r_end, r) in enumerate(run_ranges):
                if r_start <= start_idx < r_end:
                    i_start = i
                if r_start < end_idx <= r_end:
                    i_end = i
                    break
            if i_start is None:
                break
            if i_end is None:
                i_end = len(run_ranges) - 1

            r_start_info = run_ranges[i_start]
            r_end_info = run_ranges[i_end]

            if i_start == i_end:
                # Nằm trong cùng 1 run
                r = r_start_info[2]
                local_start = start_idx - r_start_info[0]
                local_end = end_idx - r_start_info[0]
                r.text = r.text[:local_start] + replace_norm + r.text[local_end:]
            else:
                # Nằm xuyên qua nhiều runs
                r_first = r_start_info[2]
                r_last = r_end_info[2]
                local_start = start_idx - r_start_info[0]
                local_end = end_idx - r_end_info[0]

                r_first.text = r_first.text[:local_start] + replace_norm
                for mid_idx in range(i_start + 1, i_end):
                    run_ranges[mid_idx][2].text = ''
                r_last.text = r_last.text[local_end:]

            cur_search_pos = start_idx + len(replace_norm)

        return True

    @classmethod
    def _thay_the_tu_dien(cls, doc, tu_dien):
        """Quét và thay thế từ điển trong tất cả Paragraphs và Tables của Document"""
        for p in doc.paragraphs:
            for k, v in tu_dien.items():
                cls._replace_text_across_runs(p, k, v)

        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for k, v in tu_dien.items():
                            cls._replace_text_across_runs(p, k, v)

    @classmethod
    def _clone_paragraph(cls, template_p, new_text, is_bold=None, is_italic=None, size_pt=14):
        """
        Nhân bản một Paragraph mẫu bằng deepcopy XML:
        Kế thừa 100% <w:pPr> (căn lề justify, thụt đầu dòng, khoảng cách dòng, dãn dòng 1.15)
        và thiết lập nội dung mới với font Times New Roman chuẩn.
        """
        import docx.text.paragraph
        new_elem = copy.deepcopy(template_p._p)
        new_p = docx.text.paragraph.Paragraph(new_elem, template_p._parent)

        if new_p.runs:
            r0 = new_p.runs[0]
            r0.text = unicodedata.normalize('NFC', str(new_text))
            cls._set_run_font_and_size(r0, font_name="Times New Roman", size_pt=size_pt, is_bold=is_bold, is_italic=is_italic)
            # Xóa các runs phụ còn sót lại trong bản copy
            for r in new_p.runs[1:]:
                new_p._p.remove(r._r)
        else:
            r0 = new_p.add_run(unicodedata.normalize('NFC', str(new_text)))
            cls._set_run_font_and_size(r0, font_name="Times New Roman", size_pt=size_pt, is_bold=is_bold, is_italic=is_italic)

        return new_p

    @classmethod
    def _set_cell_text_preserving_format(cls, cell, text, is_bold=None, align=None, size_pt=11):
        """Cập nhật text trong cell bảng word mà giữ nguyên định dạng viền, padding, font Times New Roman"""
        text_str = unicodedata.normalize('NFC', str(text))
        if cell.paragraphs:
            p = cell.paragraphs[0]
            if align is not None:
                p.alignment = align
            if p.runs:
                r0 = p.runs[0]
                r0.text = text_str
                cls._set_run_font_and_size(r0, font_name="Times New Roman", size_pt=size_pt, is_bold=is_bold)
                for r in p.runs[1:]:
                    p._p.remove(r._r)
            else:
                r0 = p.add_run(text_str)
                cls._set_run_font_and_size(r0, font_name="Times New Roman", size_pt=size_pt, is_bold=is_bold)
            # Xóa các paragraph thừa trong cell nếu có
            for extra_p in cell.paragraphs[1:]:
                cell._tc.remove(extra_p._p)
        else:
            p = cell.add_paragraph()
            if align is not None:
                p.alignment = align
            r0 = p.add_run(text_str)
            cls._set_run_font_and_size(r0, font_name="Times New Roman", size_pt=size_pt, is_bold=is_bold)

    # =========================================================================
    # CẬP NHẬT SỐ LIỆU THỰC TẾ VÀO CÁC BẢNG TRONG PHIẾU NHẬN XÉT
    # =========================================================================
    def _cap_nhat_bang_phieu_nhan_xet(self, doc, thong_tin_chung):
        """
        Cập nhật số liệu thực tế vào Bảng 1 (GTGT 12 chỉ tiêu), Bảng 2 (TNCN)
        và Bảng 3 (Đối chiếu HĐĐT bán ra vs 01/GTGT) theo đúng mẫu PhieuNhanXet.docx
        """
        du_lieu_to_khai = thong_tin_chung.get("du_lieu_to_khai", {})
        bang_br = thong_tin_chung.get("bang_hoa_don_br", [])

        # 1. Cập nhật Bảng 1: Số liệu Tờ khai thuế 01/GTGT
        gtgt_data = trich_xuat_bang_gtgt(du_lieu_to_khai)
        if len(doc.tables) > 1 and gtgt_data:
            t1 = doc.tables[1]
            years = sorted(list(gtgt_data.keys()))
            if years:
                y_main = years[-1]
                y_prev = years[-2] if len(years) > 1 else None

                # Header dòng 0
                if len(t1.columns) >= 3:
                    if y_prev:
                        self._set_cell_text_preserving_format(t1.rows[0].cells[2], f"Năm {y_prev}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                        if len(t1.columns) >= 4:
                            self._set_cell_text_preserving_format(t1.rows[0].cells[3], f"Năm {y_main}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                    else:
                        self._set_cell_text_preserving_format(t1.rows[0].cells[2], f"Năm {y_main}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

                # 12 chỉ tiêu dòng 1..12
                for i in range(1, 13):
                    if i < len(t1.rows):
                        if y_prev:
                            val_prev = gtgt_data[y_prev].get(i, 0.0)
                            val_prev_str = format_tien(val_prev) if val_prev != 0 else "-"
                            self._set_cell_text_preserving_format(t1.rows[i].cells[2], val_prev_str, align=WD_ALIGN_PARAGRAPH.RIGHT)

                            if len(t1.columns) >= 4:
                                val_main = gtgt_data[y_main].get(i, 0.0)
                                val_main_str = format_tien(val_main) if val_main != 0 else "-"
                                self._set_cell_text_preserving_format(t1.rows[i].cells[3], val_main_str, align=WD_ALIGN_PARAGRAPH.RIGHT)
                        else:
                            val_main = gtgt_data[y_main].get(i, 0.0)
                            val_main_str = format_tien(val_main) if val_main != 0 else "-"
                            self._set_cell_text_preserving_format(t1.rows[i].cells[2], val_main_str, align=WD_ALIGN_PARAGRAPH.RIGHT)

        # 2. Cập nhật Bảng 2: Quyết toán thuế TNCN (05/QTT-TNCN)
        tncn_data = trich_xuat_bang_tncn(du_lieu_to_khai)
        if len(doc.tables) > 2 and tncn_data:
            t2 = doc.tables[2]
            years_tncn = sorted(list(tncn_data.keys()))
            if years_tncn:
                y_tn = years_tncn[-1]
                if len(t2.rows[0].cells) >= 3:
                    self._set_cell_text_preserving_format(t2.rows[0].cells[2], f"05/QTT-TNCN năm {y_tn}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                val_tnct = tncn_data[y_tn].get('tnct', 0)
                val_thue = tncn_data[y_tn].get('thue', 0)
                if len(t2.rows) > 1 and len(t2.rows[1].cells) >= 3:
                    self._set_cell_text_preserving_format(t2.rows[1].cells[2], format_tien(val_tnct), align=WD_ALIGN_PARAGRAPH.RIGHT)
                if len(t2.rows) > 2 and len(t2.rows[2].cells) >= 3:
                    self._set_cell_text_preserving_format(t2.rows[2].cells[2], format_tien(val_thue), align=WD_ALIGN_PARAGRAPH.RIGHT)

        # 3. Cập nhật Bảng 3: Bảng đối chiếu HĐĐT bán ra và Tờ khai 01/GTGT
        if len(doc.tables) > 3 and bang_br:
            t3 = doc.tables[3]
            # Giữ lại Row 2 làm template row để clone
            template_tr = copy.deepcopy(t3.rows[2]._tr) if len(t3.rows) > 2 else None

            # Xóa các dòng mẫu cũ từ dòng index 2 đến dòng trước dòng Tổng cuối cùng
            while len(t3.rows) > 3:
                tr = t3.rows[2]._tr
                tr.getparent().remove(tr)

            tot_hd_t, tot_hd_th, tot_tk_t, tot_tk_th = 0.0, 0.0, 0.0, 0.0
            row_tong = t3.rows[-1]

            for idx, r_data in enumerate(bang_br):
                hd_t = float(r_data.get('hd_truoc', 0) or 0)
                hd_th = float(r_data.get('hd_thue', 0) or 0)
                tk_t = float(r_data.get('tk_truoc', 0) or 0)
                tk_th = float(r_data.get('tk_thue', 0) or 0)
                diff_t = tk_t - hd_t
                diff_th = tk_th - hd_th

                tot_hd_t += hd_t
                tot_hd_th += hd_th
                tot_tk_t += tk_t
                tot_tk_th += tk_th

                # Clone template_tr và chèn trước dòng Tổng
                if template_tr is not None:
                    new_tr = copy.deepcopy(template_tr)
                    row_tong._tr.addprevious(new_tr)
                    # Lấy row vừa chèn (dòng ngay trước row_tong)
                    cur_row = t3.rows[-2]
                else:
                    cur_row = t3.add_row()

                cells = cur_row.cells
                self._set_cell_text_preserving_format(cells[0], str(idx + 1), align=WD_ALIGN_PARAGRAPH.CENTER)
                self._set_cell_text_preserving_format(cells[1], str(r_data.get('ky', '')), align=WD_ALIGN_PARAGRAPH.CENTER)
                self._set_cell_text_preserving_format(cells[2], format_tien(hd_t), align=WD_ALIGN_PARAGRAPH.RIGHT)
                self._set_cell_text_preserving_format(cells[3], format_tien(hd_th), align=WD_ALIGN_PARAGRAPH.RIGHT)
                self._set_cell_text_preserving_format(cells[4], format_tien(tk_t), align=WD_ALIGN_PARAGRAPH.RIGHT)
                self._set_cell_text_preserving_format(cells[5], format_tien(tk_th), align=WD_ALIGN_PARAGRAPH.RIGHT)
                self._set_cell_text_preserving_format(cells[6], format_tien(diff_t) if abs(diff_t) > 10 else "-", align=WD_ALIGN_PARAGRAPH.RIGHT)
                self._set_cell_text_preserving_format(cells[7], format_tien(diff_th) if abs(diff_th) > 10 else "-", align=WD_ALIGN_PARAGRAPH.RIGHT)

            # Xóa dòng mẫu ban đầu còn dư
            if len(t3.rows) > len(bang_br) + 3:
                tr_first = t3.rows[2]._tr
                tr_first.getparent().remove(tr_first)

            # Cập nhật dòng Tổng cộng cuối cùng
            tot_diff_t = tot_tk_t - tot_hd_t
            tot_diff_th = tot_tk_th - tot_hd_th
            cells_tot = t3.rows[-1].cells
            self._set_cell_text_preserving_format(cells_tot[0], "Tổng", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            self._set_cell_text_preserving_format(cells_tot[1], "Tổng", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            self._set_cell_text_preserving_format(cells_tot[2], format_tien(tot_hd_t), is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self._set_cell_text_preserving_format(cells_tot[3], format_tien(tot_hd_th), is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self._set_cell_text_preserving_format(cells_tot[4], format_tien(tot_tk_t), is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self._set_cell_text_preserving_format(cells_tot[5], format_tien(tot_tk_th), is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self._set_cell_text_preserving_format(cells_tot[6], format_tien(tot_diff_t) if abs(tot_diff_t) > 10 else "-", is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self._set_cell_text_preserving_format(cells_tot[7], format_tien(tot_diff_th) if abs(tot_diff_th) > 10 else "-", is_bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

        # Xóa Table 4 thừa (mẫu 2026 trong template cũ) và đoạn paragraph P25
        if len(doc.tables) > 5:
            t4 = doc.tables[4]
            tbl = t4._tbl
            tbl.getparent().remove(tbl)

            # Xóa paragraph "+ Năm 2026:" nếu còn
            for p in doc.paragraphs:
                if "+ Năm 2026:" in p.text.strip():
                    p._p.getparent().remove(p._p)
                    break

    # =========================================================================
    # 1. XUẤT PHIẾU NHẬN XÉT HỒ SƠ KHAI THUẾ
    # =========================================================================
    def xuat_phieu_nhan_xet(self, file_dich, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so: bool = True):
        """
        Xuất file Phiếu nhận xét hồ sơ khai thuế.
        Đọc và thay thế trực tiếp vào Data/Mau/PhieuNhanXet.docx bằng công nghệ Run & XML Cloning
        bảo đảm 100% font chữ Times New Roman 14pt, lề, thụt dòng và bố cục văn bản chuẩn ngành.
        """
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        dia_chi = thong_tin_chung.get("dia_chi", "")
        nnkd = thong_tin_chung.get("nnkd", "Sản xuất kinh doanh")
        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        cong_chuc = thong_tin_chung.get("cong_chuc", "Phan Trung Hiếu")
        lanh_dao = thong_tin_chung.get("lanh_dao", "Nguyễn Thị Phương")
        chuc_vu = thong_tin_chung.get("chuc_vu", "PHÓ TRƯỞNG PHÒNG").upper()
        ky_thue = thong_tin_chung.get("ky_thue", "Năm 2025")
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")

        now = datetime.datetime.now()
        ngay_chu = f"ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"
        ngay_ngan = f"{now.day:02d}/{now.month:02d}/{now.year}"

        if not os.path.exists(self.mau_phieu_nhan_xet):
            return self._tao_phieu_nhan_xet_moi(file_dich, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so=la_luu_ho_so)

        doc = Document(self.mau_phieu_nhan_xet)
        self._chuan_hoa_nfc(doc)

        # 1. Thay thế metadata qua Run Replacement (giữ 100% định dạng font và size)
        tu_dien = {
            "THUẾ TỈNH QUẢNG NGÃI": cqt,
            "PHÒNG QLDN1": phong_doi,
            "TÊN CÔNG TYI": ten_nnt,
            "TÊN CÔNG TY": ten_nnt,
            "Tên CÔNG TY": ten_nnt,
            "MST công  ty": mst,
            "MST công ty": mst,
            "MST CÔNG TY": mst,
            "KCN Tịnh Phong, xã Thọ Phong, tỉnh Quảng Ngãi.": dia_chi,
            "Sản xuất khác chưa được phân vào đâu.": nnkd,
            "Kỳ thuế: Năm 2025, Quý 1-2/2026": f"Kỳ thuế: {ky_thue}",
            "06/10/2026": ngay_ngan,
            "Quảng Ngãi, ngày 06 tháng 10 năm 2026": f"{dia_danh}, {ngay_chu}",
            "Nguyễn Thị Phương": lanh_dao,
            "Phan Trung Hiếu": cong_chuc,
            "PHÓ TRƯỞNG PHÒNG PHÒNG QLDN SỐ 1": chuc_vu
        }
        if not la_luu_ho_so:
            tu_dien["Kính trình Lãnh đạo Thuế tỉnh phê duyệt kết quả rà soát, lưu hồ sơ và tiếp tục theo dõi, quản lý đơn vị./."] = \
                "Kính trình Lãnh đạo Thuế tỉnh xem xét, phê duyệt ban hành Thông báo yêu cầu Người nộp thuế giải trình, bổ sung thông tin, tài liệu theo quy định./."

        self._thay_the_tu_dien(doc, tu_dien)

        # 2. Cập nhật Bảng 1, Bảng 2 và Bảng 3 số liệu thực tế
        self._cap_nhat_bang_phieu_nhan_xet(doc, thong_tin_chung)

        # 3. Định vị khu vực II. NHẬN XÉT HỒ SƠ KHAI THUẾ
        p21_head = None
        p22_sample = None
        p_p27_tncn = None
        p_p28_danh_gia = None
        p_cam_ket = None

        for p in doc.paragraphs:
            txt = p.text.strip()
            if "2.1. Về Thuế GTGT" in txt and p21_head is None:
                p21_head = p
            elif "- Trên tờ khai thuế GTGT bổ sung" in txt and p22_sample is None:
                p22_sample = p
            elif "2.2. Về Thuế TNCN" in txt and p_p27_tncn is None:
                p_p27_tncn = p
            elif "Qua rà soát hồ sơ khai thuế của" in txt and "công chức chưa phát hiện dấu hiệu rủi ro" in txt and p_p28_danh_gia is None:
                p_p28_danh_gia = p
            elif "Công chức thực hiện rà soát theo số liệu" in txt and p_cam_ket is None:
                p_cam_ket = p

        # Phân loại nội dung theo nhóm chuẩn (BCTC, Đối chiếu chéo, Phân tích tỷ lệ gom vào BCTC & Chéo)
        muc_gtgt = [m for m in ds_noi_dung_chon if "gtgt" in str(m.get("nhom", "")).lower() or "hóa đơn" in str(m.get("nhom", "")).lower()]
        muc_tndn = [m for m in ds_noi_dung_chon if "tndn" in str(m.get("nhom", "")).lower()]
        muc_bctc = [m for m in ds_noi_dung_chon if any(k in str(m.get("nhom", "")).lower() for k in ["tài chính", "bctc", "chéo", "đối chiếu", "tỷ lệ"])]
        muc_tncn = [m for m in ds_noi_dung_chon if "tncn" in str(m.get("nhom", "")).lower()]
        muc_khac = [m for m in ds_noi_dung_chon if m not in muc_gtgt and m not in muc_tndn and m not in muc_bctc and m not in muc_tncn]

        if p22_sample is not None and p21_head is not None:
            # Chèn các phát hiện GTGT bằng cách clone từ p22_sample
            if muc_gtgt:
                for item in reversed(muc_gtgt):
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_clone = self._clone_paragraph(p22_sample, f"- {cau}", size_pt=14)
                    p22_sample._p.addnext(p_clone._p)
            else:
                p_clone = self._clone_paragraph(p22_sample, "- Qua đối chiếu hóa đơn và tờ khai thuế GTGT, các chỉ tiêu kê khai phù hợp, chưa phát hiện sai lệch lớn.", size_pt=14)
                p22_sample._p.addnext(p_clone._p)

            # Xóa paragraph mẫu P22
            p22_sample._p.getparent().remove(p22_sample._p)

        # Xóa đoạn P27 cũ ("2.2. Về Thuế TNCN...") để tái cơ cấu thứ tự chuẩn 2.1, 2.2, 2.3, 2.4
        if p_p27_tncn is not None:
            try:
                p_p27_tncn._p.getparent().remove(p_p27_tncn._p)
            except Exception:
                pass

        # Cập nhật đoạn đánh giá tổng kết P28 phù hợp với luồng Lưu hồ sơ hoặc Ban hành Thông báo
        if p_p28_danh_gia is not None:
            if la_luu_ho_so:
                txt_p28 = f"\tQua rà soát hồ sơ khai thuế của {ten_nnt} đã kê khai nộp cho cơ quan thuế tính đến thời điểm rà soát: thủ tục và thời hạn nộp hồ sơ đúng theo quy định. Tại thời điểm rà soát, công chức chưa phát hiện dấu hiệu rủi ro. Công chức tiếp tục theo dõi quản lý và tiến hành đề xuất các bước xử lý khi phát hiện có dấu hiệu rủi ro."
            else:
                txt_p28 = f"\tQua rà soát hồ sơ khai thuế của {ten_nnt} đã kê khai nộp cho cơ quan thuế tính đến thời điểm rà soát: phát hiện một số nội dung có sự chênh lệch, rủi ro cần làm rõ. Công chức đề xuất ban hành Thông báo yêu cầu Người nộp thuế giải trình, bổ sung thông tin, tài liệu theo quy định."
            p_p28_danh_gia.text = unicodedata.normalize('NFC', txt_p28)
            r_eval = p_p28_danh_gia.runs[0] if p_p28_danh_gia.runs else p_p28_danh_gia.add_run()
            self._set_run_font_and_size(r_eval, font_name="Times New Roman", size_pt=14)

        # Chèn lần lượt 2.2 (TNDN), 2.3 (BCTC & Đối chiếu chéo), 2.4 (TNCN) ngay trước đoạn tổng kết
        target_p = p_p28_danh_gia if p_p28_danh_gia is not None else p_cam_ket
        if target_p is not None and p21_head is not None:
            # 2.2. Về Thuế Thu nhập doanh nghiệp
            p_head_tndn = self._clone_paragraph(p21_head, "2.2. Về Thuế Thu nhập doanh nghiệp:", is_bold=True, size_pt=14)
            target_p._p.addprevious(p_head_tndn._p)
            if muc_tndn:
                for item in muc_tndn:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_sub = self._clone_paragraph(p21_head, f"- {cau}", is_bold=False, size_pt=14)
                    target_p._p.addprevious(p_sub._p)
            else:
                p_sub = self._clone_paragraph(p21_head, "- Qua rà soát hồ sơ quyết toán thuế TNDN, các chỉ tiêu kê khai phù hợp, chưa phát hiện sai lệch lớn.", is_bold=False, size_pt=14)
                target_p._p.addprevious(p_sub._p)

            # 2.3. Về Báo cáo Tài chính và Đối chiếu chéo dữ liệu
            p_head_bctc = self._clone_paragraph(p21_head, "2.3. Về Báo cáo Tài chính và Đối chiếu chéo dữ liệu:", is_bold=True, size_pt=14)
            target_p._p.addprevious(p_head_bctc._p)
            if muc_bctc:
                for item in muc_bctc:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_sub = self._clone_paragraph(p21_head, f"- {cau}", is_bold=False, size_pt=14)
                    target_p._p.addprevious(p_sub._p)
            else:
                p_sub = self._clone_paragraph(p21_head, "- Qua rà soát Báo cáo tài chính và đối chiếu chéo dữ liệu với tờ khai thuế, chưa phát hiện sai lệch lớn.", is_bold=False, size_pt=14)
                target_p._p.addprevious(p_sub._p)

            # 2.4. Về Thuế Thu nhập cá nhân
            p_head_tncn = self._clone_paragraph(p21_head, "2.4. Về Thuế Thu nhập cá nhân:", is_bold=True, size_pt=14)
            target_p._p.addprevious(p_head_tncn._p)
            if muc_tncn:
                for item in muc_tncn:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_sub = self._clone_paragraph(p21_head, f"- {cau}", is_bold=False, size_pt=14)
                    target_p._p.addprevious(p_sub._p)
            else:
                p_sub = self._clone_paragraph(p21_head, "- Qua rà soát hồ sơ quyết toán thuế TNCN (mẫu 05/QTT-TNCN), đơn vị đã kê khai và nộp thuế theo quy định, chưa phát hiện dấu hiệu rủi ro.", is_bold=False, size_pt=14)
                target_p._p.addprevious(p_sub._p)

            # 2.5. Về các nội dung rủi ro khác (nếu có)
            if muc_khac:
                p_head_khac = self._clone_paragraph(p21_head, "2.5. Về các nội dung rủi ro khác:", is_bold=True, size_pt=14)
                target_p._p.addprevious(p_head_khac._p)
                for item in muc_khac:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_sub = self._clone_paragraph(p21_head, f"- {cau}", is_bold=False, size_pt=14)
                    target_p._p.addprevious(p_sub._p)

        doc.save(file_dich)
        return True

    # =========================================================================
    # 2. XUẤT THÔNG BÁO GIẢI TRÌNH, BỔ SUNG TÀI LIỆU
    # =========================================================================
    def xuat_thong_bao_giai_trinh(self, file_dich, thong_tin_chung, ds_noi_dung_chon):
        """
        Xuất file Thông báo giải trình gửi Doanh nghiệp.
        Đọc và thay thế trực tiếp vào Data/Mau/ThongBaoGiaiTrinh.docx bằng công nghệ Run & XML Cloning
        bảo đảm 100% chuẩn văn bản hành chính quy định tại Nghị định 30/2020/NĐ-CP và mẫu ngành Thuế.
        """
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        dia_chi = thong_tin_chung.get("dia_chi", "")
        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        lanh_dao = thong_tin_chung.get("lanh_dao", "Trần Quốc Anh")
        chuc_vu = thong_tin_chung.get("chuc_vu", "TRƯỞNG PHÒNG").upper()
        lan_tb = thong_tin_chung.get("lan_tb", "Lần 1")
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")
        ngay_hen = thong_tin_chung.get("ngay_hen", (datetime.datetime.now() + datetime.timedelta(days=10)).strftime("%d/%m/%Y"))
        dien_thoai_lh = thong_tin_chung.get("dien_thoai", "0255 3824589")
        dia_chi_cqt = thong_tin_chung.get("dia_chi_cqt", "507 Quang Trung, phường Cẩm Thành, tỉnh Quảng Ngãi")

        now = datetime.datetime.now()
        ngay_thang_nam = f"{dia_danh}, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"

        if not os.path.exists(self.mau_thong_bao):
            return self._tao_thong_bao_moi(file_dich, thong_tin_chung, ds_noi_dung_chon)

        doc = Document(self.mau_thong_bao)
        self._chuan_hoa_nfc(doc)

        # 1. Thay thế metadata giữ 100% font và size
        tu_dien = {
            "CỤC THUẾ THUẾ TỈNH QUẢNG NGÃI": cqt,
            "THUẾ TỈNH QUẢNG NGÃI": cqt,
            "Thuế tỉnh Quảng Ngãi": cqt.title(),
            "Tên CÔNG TY": ten_nnt,
            "TÊN CÔNG TY": ten_nnt,
            "MST CÔNG TY": mst,
            "MST công ty": mst,
            "243/01 Nguyễn Công Phương, phường Nghĩa Lộ, tỉnh Quảng Ngãi.": dia_chi,
            "Trần Quốc Anh": lanh_dao,
            "TRƯỞNG PHÒNG QUẢN LÝ,  HỖ TRỢ DOANH NGHIỆP SỐ 1": chuc_vu,
            "PHÒNG QLDN1": phong_doi,
            "(Lần 1)": f"({lan_tb})",
            "ngày 22/10/2026": f"ngày {ngay_hen}",
            "22/10/2026": ngay_hen,
            "0255 3824589": dien_thoai_lh,
            "507 Quang Trung, phường Cẩm Thành, tỉnh Quảng Ngãi.": dia_chi_cqt,
            "507 Quang Trung, phường Cẩm Thành, tỉnh Quảng Ngãi": dia_chi_cqt,
            "Quảng Ngãi, ngày       tháng      năm": ngay_thang_nam
        }
        self._thay_the_tu_dien(doc, tu_dien)

        # Phân loại rủi ro
        muc_gtgt = [m for m in ds_noi_dung_chon if "gtgt" in str(m.get("nhom", "")).lower() or "hóa đơn" in str(m.get("nhom", "")).lower()]
        muc_tndn = [m for m in ds_noi_dung_chon if "tndn" in str(m.get("nhom", "")).lower()]
        muc_bctc = [m for m in ds_noi_dung_chon if any(k in str(m.get("nhom", "")).lower() for k in ["tài chính", "bctc", "chéo", "đối chiếu", "tỷ lệ"])]
        muc_tncn = [m for m in ds_noi_dung_chon if "tncn" in str(m.get("nhom", "")).lower()]
        muc_khac = [m for m in ds_noi_dung_chon if m not in muc_gtgt and m not in muc_tndn and m not in muc_bctc and m not in muc_tncn]

        # 2. Định vị 2 khu vực nội dung chính
        p_head_sec1 = None   # "1. Thuế GTGT: " (Paragraph 9)
        p_item_sec1 = None   # Paragraph 10
        p_de_nghi = None     # "Thuế tỉnh Quảng Ngãi đề nghị..." (Paragraph 16)
        p_head_sec2 = None   # "2. Về thuế TNDN:" (Paragraph 18)
        p_item_sec2 = None   # Paragraph 19
        p_han_nop = None     # Paragraph 21 ("Văn bản giải trình...")

        for p in doc.paragraphs:
            txt = p.text.strip()
            if "1. Thuế GTGT" in txt and p_head_sec1 is None:
                p_head_sec1 = p
            elif "- Từ năm 2022 đến năm 2025" in txt and p_item_sec1 is None:
                p_item_sec1 = p
            elif "đề nghị" in txt and "có văn bản để giải trình" in txt:
                p_de_nghi = p
            elif "2. Về thuế TNDN" in txt and p_head_sec2 is None:
                p_head_sec2 = p
            elif "Giải trình nội dung chi tiết Hoạt động" in txt and p_item_sec2 is None:
                p_item_sec2 = p
            elif "Văn bản giải trình, thông tin, tài liệu bổ sung" in txt:
                p_han_nop = p

        # 3. Điền Phần 1: Để làm rõ các nội dung
        if p_de_nghi is not None and p_head_sec1 is not None and p_item_sec1 is not None:
            # Chèn các nội dung mới ngay trước p_de_nghi
            stt = 1
            if muc_gtgt:
                h_gtgt = self._clone_paragraph(p_head_sec1, f"{stt}. Thuế Giá trị gia tăng và Hóa đơn điện tử:", is_bold=True, size_pt=14)
                p_de_nghi._p.addprevious(h_gtgt._p)
                stt += 1
                for item in muc_gtgt:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_c = self._clone_paragraph(p_item_sec1, f"- {cau}", is_bold=False, size_pt=14)
                    p_de_nghi._p.addprevious(p_c._p)

            if muc_tndn:
                h_tndn = self._clone_paragraph(p_head_sec1, f"{stt}. Thuế Thu nhập doanh nghiệp:", is_bold=True, size_pt=14)
                p_de_nghi._p.addprevious(h_tndn._p)
                stt += 1
                for item in muc_tndn:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_c = self._clone_paragraph(p_item_sec1, f"- {cau}", is_bold=False, size_pt=14)
                    p_de_nghi._p.addprevious(p_c._p)

            if muc_bctc:
                h_bctc = self._clone_paragraph(p_head_sec1, f"{stt}. Báo cáo Tài chính và Đối chiếu chéo dữ liệu:", is_bold=True, size_pt=14)
                p_de_nghi._p.addprevious(h_bctc._p)
                stt += 1
                for item in muc_bctc:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_c = self._clone_paragraph(p_item_sec1, f"- {cau}", is_bold=False, size_pt=14)
                    p_de_nghi._p.addprevious(p_c._p)

            if muc_tncn:
                h_tncn = self._clone_paragraph(p_head_sec1, f"{stt}. Thuế Thu nhập cá nhân:", is_bold=True, size_pt=14)
                p_de_nghi._p.addprevious(h_tncn._p)
                stt += 1
                for item in muc_tncn:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_c = self._clone_paragraph(p_item_sec1, f"- {cau}", is_bold=False, size_pt=14)
                    p_de_nghi._p.addprevious(p_c._p)

            if muc_khac:
                h_khac = self._clone_paragraph(p_head_sec1, f"{stt}. Các nội dung khác:", is_bold=True, size_pt=14)
                p_de_nghi._p.addprevious(h_khac._p)
                stt += 1
                for item in muc_khac:
                    cau = item.get('cau_van', '').lstrip('- ').lstrip('+ ')
                    p_c = self._clone_paragraph(p_item_sec1, f"- {cau}", is_bold=False, size_pt=14)
                    p_de_nghi._p.addprevious(p_c._p)

            # Xóa các đoạn mẫu cũ giữa "Để làm rõ các nội dung:" và p_de_nghi
            mau_cu_sec1 = ["1. Thuế GTGT:", "Từ năm 2022 đến năm 2025", "2. Thuế TNDN:", "Tại tờ khai Quyết toán thuế TNDN", "có chênh lệch giữa Doanh thu"]
            for p in list(doc.paragraphs):
                if any(m in p.text for m in mau_cu_sec1) and p != p_de_nghi:
                    try:
                        p._p.getparent().remove(p._p)
                    except Exception:
                        pass

        # 4. Điền Phần 2: Yêu cầu đơn vị có văn bản giải trình cụ thể
        # Tái định vị p_han_nop
        p_han_nop_moi = None
        for p in doc.paragraphs:
            if "Văn bản giải trình, thông tin, tài liệu bổ sung" in p.text:
                p_han_nop_moi = p
                break

        if p_han_nop_moi is not None and p_head_sec2 is not None and p_item_sec2 is not None:
            stt_yc = 1
            if muc_gtgt:
                h_yc = self._clone_paragraph(p_head_sec2, f"{stt_yc}. Về thuế Giá trị gia tăng và Hóa đơn:", is_bold=True, size_pt=14)
                p_han_nop_moi._p.addprevious(h_yc._p)
                stt_yc += 1
                for item in muc_gtgt:
                    yc = item.get("yeu_cau", "Giải trình chi tiết nguyên nhân chênh lệch và cung cấp hồ sơ minh chứng.")
                    p_yc = self._clone_paragraph(p_item_sec2, f"- {yc}", is_bold=False, size_pt=14)
                    p_han_nop_moi._p.addprevious(p_yc._p)

            if muc_tndn:
                h_yc = self._clone_paragraph(p_head_sec2, f"{stt_yc}. Về thuế Thu nhập doanh nghiệp:", is_bold=True, size_pt=14)
                p_han_nop_moi._p.addprevious(h_yc._p)
                stt_yc += 1
                for item in muc_tndn:
                    yc = item.get("yeu_cau", "Giải trình nguyên nhân chênh lệch và làm rõ các khoản thu nhập/chi phí.")
                    p_yc = self._clone_paragraph(p_item_sec2, f"- {yc}", is_bold=False, size_pt=14)
                    p_han_nop_moi._p.addprevious(p_yc._p)

            if muc_bctc:
                h_yc = self._clone_paragraph(p_head_sec2, f"{stt_yc}. Về Báo cáo Tài chính và Đối chiếu chéo dữ liệu:", is_bold=True, size_pt=14)
                p_han_nop_moi._p.addprevious(h_yc._p)
                stt_yc += 1
                for item in muc_bctc:
                    yc = item.get("yeu_cau", "Giải trình chi tiết số liệu Báo cáo tài chính, số dư công nợ và các khoản chênh lệch.")
                    p_yc = self._clone_paragraph(p_item_sec2, f"- {yc}", is_bold=False, size_pt=14)
                    p_han_nop_moi._p.addprevious(p_yc._p)

            if muc_tncn:
                h_yc = self._clone_paragraph(p_head_sec2, f"{stt_yc}. Về thuế Thu nhập cá nhân:", is_bold=True, size_pt=14)
                p_han_nop_moi._p.addprevious(h_yc._p)
                stt_yc += 1
                for item in muc_tncn:
                    yc = item.get("yeu_cau", "Giải trình chi tiết thu nhập chịu thuế và số thuế đã khấu trừ.")
                    p_yc = self._clone_paragraph(p_item_sec2, f"- {yc}", is_bold=False, size_pt=14)
                    p_han_nop_moi._p.addprevious(p_yc._p)

            if muc_khac:
                h_yc = self._clone_paragraph(p_head_sec2, f"{stt_yc}. Về các nội dung khác:", is_bold=True, size_pt=14)
                p_han_nop_moi._p.addprevious(h_yc._p)
                stt_yc += 1
                for item in muc_khac:
                    yc = item.get("yeu_cau", "Giải trình và cung cấp tài liệu giải trình bổ sung.")
                    p_yc = self._clone_paragraph(p_item_sec2, f"- {yc}", is_bold=False, size_pt=14)
                    p_han_nop_moi._p.addprevious(p_yc._p)

            # Xóa các đoạn mẫu cũ của Phần 2
            mau_cu_sec2 = ["1. Về thuế GTGT: Giải trình nội dung", "2. Về thuế TNDN:", "Giải trình nội dung chi tiết Hoạt động", "Giải trình nguyên nhân chênh lệch giữa Doanh thu"]
            for p in list(doc.paragraphs):
                if any(m in p.text for m in mau_cu_sec2) and p != p_han_nop_moi:
                    try:
                        p._p.getparent().remove(p._p)
                    except Exception:
                        pass

        doc.save(file_dich)
        return True

    # =========================================================================
    # 3. XUẤT TỜ TRÌNH LÃNH ĐẠO (PHÊ DUYỆT LƯU HỒ SƠ HOẶC BAN HÀNH THÔNG BÁO)
    # =========================================================================
    def xuat_to_trinh(self, file_dich, thong_tin_chung, ds_noi_dung_chon, ds_can_cu_chon=None, la_luu_ho_so: bool = True):
        """
        Xuất Tờ trình Lãnh đạo Cục/Chi cục phê duyệt kết quả kiểm tra hoặc ban hành Thông báo giải trình.
        Đọc và thay thế trực tiếp vào Data/Mau/To trinh phe duyet luu ho so_BETONG.docx bằng Run & XML Cloning
        bảo lưu 100% thể thức tờ trình hành chính và 3 cấp phê duyệt.
        - la_luu_ho_so = True: Tờ trình phê duyệt lưu hồ sơ
        - la_luu_ho_so = False: Tờ trình ban hành Thông báo yêu cầu giải trình, bổ sung thông tin
        """
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        cong_chuc = thong_tin_chung.get("cong_chuc", "Công chức đề xuất")
        lanh_dao = thong_tin_chung.get("lanh_dao", "Nguyễn Thị Phương")
        chuc_vu = thong_tin_chung.get("chuc_vu", "PHÓ TRƯỞNG PHÒNG").upper()
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")

        now = datetime.datetime.now()
        ngay_chu = f"ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"
        dia_danh_ngay = f"{dia_danh}, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"

        if not os.path.exists(self.mau_to_trinh):
            return self._tao_to_trinh_moi(file_dich, thong_tin_chung, ds_noi_dung_chon, ds_can_cu_chon, la_luu_ho_so=la_luu_ho_so)

        doc = Document(self.mau_to_trinh)
        self._chuan_hoa_nfc(doc)

        tu_dien = {
            "THUẾ TỈNH QUẢNG NGÃI": cqt,
            "PHÒNG QLDN1": phong_doi,
            "TÊN CÔNG TY": ten_nnt,
            "Tên CÔNG TY": ten_nnt,
            "MST công  ty": mst,
            "MST công ty": mst,
            "MST CÔNG TY": mst,
            "Tên  công chức": cong_chuc,
            "Tên công chức": cong_chuc,
            "Tên công chức": cong_chuc,
            "TÊN LÃNH ĐẠO": lanh_dao,
            "CHỨC DANH LÃNH ĐẠO": chuc_vu,
            "Quảng Ngãi, ngày 06 tháng 10 năm 2026": dia_danh_ngay,
            "Địa danh, ngày 06 tháng 10 năm 2026": dia_danh_ngay,
            "Địa danh, ngày 06 tháng 10 năm 2026": dia_danh_ngay,
            "ngày 06 tháng 10 năm 2026": ngay_chu,
            "Căn cứ Phiếu nhận xét hồ sơ thuế ngày tháng năm của TÊN CÔNG TY.": f"Căn cứ Phiếu nhận xét hồ sơ thuế ngày {now.day:02d}/{now.month:02d}/{now.year} của {ten_nnt}."
        }

        # Nếu là luồng Ban hành Thông báo giải trình: chuyển đổi tiêu đề và nội dung công việc
        if not la_luu_ho_so:
            tu_dien.update({
                "Về việc phê duyệt lưu hồ sơ kiểm tra tại cơ quan thuế đối với": "Về việc ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với",
                "Trình Lãnh đạo Thuế tỉnh phê duyệt lưu hồ sơ kiểm tra tại cơ quan thuế đối với": "Trình Lãnh đạo Thuế tỉnh phê duyệt ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với",
                "Kính trình Lãnh đạo Thuế tỉnh xem xét phê duyệt kết quả rà soát để lưu hồ sơ./.": "Kính trình Lãnh đạo Thuế tỉnh xem xét phê duyệt ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với đơn vị./.",
                "kính trình Lãnh đạo Thuế tỉnh phê duyệt./.": "kính trình Lãnh đạo Thuế tỉnh ký ban hành Thông báo./.",
                "kính trình Lãnh đạo Thuế tỉnh phê duyệt": "kính trình Lãnh đạo Thuế tỉnh ký ban hành Thông báo"
            })

        self._thay_the_tu_dien(doc, tu_dien)

        # Xóa các dòng căn cứ placeholder "Căn cứ 1............."
        for p in list(doc.paragraphs):
            txt = p.text.strip()
            if "Căn cứ 1............." in txt or "Căn cứ 2..............." in txt:
                p._p.getparent().remove(p._p)

        # Định vị Section 2 và chèn căn cứ bổ sung nếu có
        p_pns = None
        p_mau_cc = None
        for p in doc.paragraphs:
            if "Căn cứ Phiếu nhận xét hồ sơ thuế" in p.text:
                p_pns = p
            elif "Căn cứ mặc định 1" in p.text or "Căn cứ mặc định 1" in p.text:
                p_mau_cc = p

        if p_pns is not None and p_mau_cc is not None and ds_can_cu_chon:
            for cc in ds_can_cu_chon:
                p_new_cc = self._clone_paragraph(p_mau_cc, f"- Căn cứ {cc};", size_pt=14)
                p_pns._p.addprevious(p_new_cc._p)

        # Định vị Section 3: Đề xuất của công chức
        p_de_xuat_head = None
        p_sample_dx = None
        p_cam_ket_dx = None

        for p in doc.paragraphs:
            txt = p.text.strip()
            if "3. Đề xuất của công chức thuế giải quyết:" in txt:
                p_de_xuat_head = p
            elif "Qua rà soát hồ sơ khai thuế của" in txt and "chưa phát hiện dấu hiệu rủi ro" in txt:
                p_sample_dx = p
            elif "Công chức thực hiện rà soát theo số liệu đơn vị đã kê khai" in txt:
                p_cam_ket_dx = p

        if p_sample_dx is not None and p_cam_ket_dx is not None:
            if ds_noi_dung_chon:
                # Tạo đoạn dẫn đề xuất phù hợp với luồng
                if la_luu_ho_so:
                    txt_dan = f"Qua rà soát hồ sơ khai thuế của {ten_nnt} đã kê khai nộp cho cơ quan thuế tính đến thời điểm rà soát và kết quả đối chiếu, giải trình, Công chức ghi nhận các nội dung cụ thể như sau:"
                else:
                    txt_dan = f"Qua rà soát hồ sơ khai thuế của {ten_nnt} đã kê khai nộp cho cơ quan thuế tính đến thời điểm rà soát, Công chức phát hiện các nội dung có sự chênh lệch/sai lệch cụ thể như sau:"
                p_dan = self._clone_paragraph(p_sample_dx, txt_dan, size_pt=14)
                p_cam_ket_dx._p.addprevious(p_dan._p)

                # Chèn các bullet item rủi ro
                for item in ds_noi_dung_chon:
                    p_finding = self._clone_paragraph(p_sample_dx, f"+ {item.get('cau_van', '')}", size_pt=14)
                    p_cam_ket_dx._p.addprevious(p_finding._p)

                # Xóa paragraph mẫu
                p_sample_dx._p.getparent().remove(p_sample_dx._p)

        doc.save(file_dich)
        return True

    # =========================================================================
    # 4. XUẤT TRỌN BỘ HỒ SƠ COMBO (2 HOẶC 3 BIỂU MẪU ĐỒNG THỜI)
    # =========================================================================
    def xuat_combo_ho_so(self, thu_muc_dich, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so: bool = True, ds_can_cu_chon=None):
        """
        Xuất trọn bộ hồ sơ (Combo) theo đúng quy trình nghiệp vụ kiểm tra thuế:
        - Nếu la_luu_ho_so = True (Combo 2 biểu mẫu):
            1. Tờ trình Lãnh đạo phê duyệt lưu hồ sơ
            2. Phiếu nhận xét hồ sơ khai thuế (Mẫu 02/QTKT)
        - Nếu la_luu_ho_so = False (Combo 3 biểu mẫu):
            1. Tờ trình Lãnh đạo ban hành Thông báo giải trình
            2. Thông báo yêu cầu giải trình, bổ sung thông tin (Lần 1)
            3. Phiếu nhận xét hồ sơ khai thuế (Mẫu 02/QTKT)
        Trả về danh sách tuple: [(ten_file, duong_dan_file), ...]
        """
        os.makedirs(thu_muc_dich, exist_ok=True)
        ten_raw = thong_tin_chung.get("ten_nnt", "NNT").strip()
        ten_clean = ten_raw.replace(' ', '_').replace('CÔNG_TY_', '').replace('TNHH_', '')
        mst = thong_tin_chung.get("mst", "").strip()
        tag = f"{mst}_{ten_clean}" if mst else ten_clean

        ket_qua = []
        if la_luu_ho_so:
            f1_name = f"01_To_trinh_phe_duyet_luu_ho_so_{tag}.docx"
            f1_path = os.path.join(thu_muc_dich, f1_name)
            self.xuat_to_trinh(f1_path, thong_tin_chung, ds_noi_dung_chon, ds_can_cu_chon, la_luu_ho_so=True)
            ket_qua.append((f1_name, f1_path))

            f2_name = f"02_Phieu_nhan_xet_ho_so_{tag}.docx"
            f2_path = os.path.join(thu_muc_dich, f2_name)
            self.xuat_phieu_nhan_xet(f2_path, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so=True)
            ket_qua.append((f2_name, f2_path))
        else:
            f1_name = f"01_To_trinh_ban_hanh_thong_bao_{tag}.docx"
            f1_path = os.path.join(thu_muc_dich, f1_name)
            self.xuat_to_trinh(f1_path, thong_tin_chung, ds_noi_dung_chon, ds_can_cu_chon, la_luu_ho_so=False)
            ket_qua.append((f1_name, f1_path))

            f2_name = f"02_Thong_bao_giai_trinh_{tag}.docx"
            f2_path = os.path.join(thu_muc_dich, f2_name)
            self.xuat_thong_bao_giai_trinh(f2_path, thong_tin_chung, ds_noi_dung_chon)
            ket_qua.append((f2_name, f2_path))

            f3_name = f"03_Phieu_nhan_xet_ho_so_{tag}.docx"
            f3_path = os.path.join(thu_muc_dich, f3_name)
            self.xuat_phieu_nhan_xet(f3_path, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so=False)
            ket_qua.append((f3_name, f3_path))

        return ket_qua

    # =========================================================================
    # CÁC BẢNG KHÔNG VIỀN (BORDERLESS TABLES) THEO NGHỊ ĐỊNH 30/2020/NĐ-CP
    # =========================================================================
    @staticmethod
    def _tao_bang_khong_vien(doc, rows, cols, col_widths=None):
        """Tạo bảng 2 cột không viền (borderless) để đặt Quốc hiệu, Tiêu ngữ, Kính gửi và Chữ ký"""
        tbl = doc.add_table(rows=rows, cols=cols)
        tbl.autofit = False

        tblPr = tbl._tbl.tblPr
        tblBorders = tblPr.find(qn('w:tblBorders'))
        if tblBorders is None:
            tblBorders = OxmlElement('w:tblBorders')
            tblPr.append(tblBorders)

        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'none')
            tblBorders.append(border)

        if col_widths:
            for r in tbl.rows:
                for idx, w in enumerate(col_widths):
                    if idx < len(r.cells):
                        r.cells[idx].width = w
        return tbl

    # =========================================================================
    # FALLBACK ENGINE: NẾU THIẾU FILE TEMPLATE GỐC TRONG DATA/MAU
    # =========================================================================
    def _tao_phieu_nhan_xet_moi(self, file_dich, thong_tin_chung, ds_noi_dung_chon, la_luu_ho_so: bool = True):
        doc = Document()
        for s in doc.sections:
            s.top_margin, s.bottom_margin = Cm(2), Cm(2)
            s.left_margin, s.right_margin = Cm(3), Cm(2)

        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        dia_chi = thong_tin_chung.get("dia_chi", "")
        nnkd = thong_tin_chung.get("nnkd", "Sản xuất kinh doanh")
        ky_thue = thong_tin_chung.get("ky_thue", "Năm 2025")
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")
        cong_chuc = thong_tin_chung.get("cong_chuc", "Công chức kiểm tra")
        lanh_dao = thong_tin_chung.get("lanh_dao", "Nguyễn Thị Phương")
        chuc_vu = thong_tin_chung.get("chuc_vu", "PHÓ TRƯỞNG PHÒNG").upper()

        now = datetime.datetime.now()
        ngay_chu = f"ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"

        # 1. Table 0 không viền: Quốc hiệu tiêu ngữ và CQT
        tb0 = self._tao_bang_khong_vien(doc, 2, 2, [Cm(7.5), Cm(10.5)])
        self._set_cell_text_preserving_format(tb0.rows[0].cells[0], f"{cqt}\n{phong_doi}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb0.rows[0].cells[1], "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("PHIẾU NHẬN XÉT HỒ SƠ KHAI THUẾ\n")
        self._set_run_font_and_size(r_t, size_pt=15, is_bold=True)
        r_k = p_t.add_run(f"Kỳ thuế: {ky_thue}")
        self._set_run_font_and_size(r_k, size_pt=13, is_italic=True)

        doc.add_paragraph(f"- Người nộp thuế: {ten_nnt}.")
        doc.add_paragraph(f"- Mã số thuế: {mst}")
        doc.add_paragraph(f"- Địa chỉ: {dia_chi}.")
        doc.add_paragraph(f"- Ngành nghề kinh doanh chính: {nnkd}.")

        p_s2 = doc.add_paragraph("II. NHẬN XÉT HỒ SƠ KHAI THUẾ")
        self._set_run_font_and_size(p_s2.runs[0], size_pt=14, is_bold=True)
        for idx, item in enumerate(ds_noi_dung_chon):
            p_i = doc.add_paragraph(f"- {item.get('cau_van', '')}")
            self._set_run_font_and_size(p_i.runs[0], size_pt=14)

        p_s3 = doc.add_paragraph("III. KIẾN NGHỊ")
        self._set_run_font_and_size(p_s3.runs[0], size_pt=14, is_bold=True)
        if la_luu_ho_so:
            doc.add_paragraph("Kính trình Lãnh đạo Thuế tỉnh phê duyệt kết quả rà soát, lưu hồ sơ và tiếp tục theo dõi, quản lý đơn vị./.")
        else:
            doc.add_paragraph("Kính trình Lãnh đạo Thuế tỉnh xem xét, phê duyệt ban hành Thông báo yêu cầu Người nộp thuế giải trình, bổ sung thông tin, tài liệu theo quy định./.")

        p_d = doc.add_paragraph(f"{dia_danh}, {ngay_chu}")
        p_d.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Table chữ ký không viền 2 cột
        tb_sig = self._tao_bang_khong_vien(doc, 2, 2, [Cm(9), Cm(9)])
        self._set_cell_text_preserving_format(tb_sig.rows[0].cells[0], chuc_vu, is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb_sig.rows[0].cells[1], "CÔNG CHỨC", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb_sig.rows[1].cells[0], f"\n\n\n\n{lanh_dao}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb_sig.rows[1].cells[1], f"\n\n\n\n{cong_chuc}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        doc.save(file_dich)
        return True

    def _tao_thong_bao_moi(self, file_dich, thong_tin_chung, ds_noi_dung_chon):
        doc = Document()
        for s in doc.sections:
            s.top_margin, s.bottom_margin = Cm(2), Cm(2)
            s.left_margin, s.right_margin = Cm(3), Cm(2)

        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        dia_chi = thong_tin_chung.get("dia_chi", "")
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")
        ngay_hen = thong_tin_chung.get("ngay_hen", (datetime.datetime.now() + datetime.timedelta(days=10)).strftime("%d/%m/%Y"))
        dien_thoai_lh = thong_tin_chung.get("dien_thoai", "0255 3824589")
        dia_chi_cqt = thong_tin_chung.get("dia_chi_cqt", "507 Quang Trung, phường Cẩm Thành, tỉnh Quảng Ngãi")
        lanh_dao = thong_tin_chung.get("lanh_dao", "Trần Quốc Anh")
        chuc_vu = thong_tin_chung.get("chuc_vu", "TRƯỞNG PHÒNG").upper()

        now = datetime.datetime.now()
        ngay_thang_nam = f"{dia_danh}, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"

        # 1. Table 0 không viền: Tiêu ngữ và CQT
        tb0 = self._tao_bang_khong_vien(doc, 2, 2, [Cm(7.5), Cm(10.5)])
        self._set_cell_text_preserving_format(tb0.rows[0].cells[0], f"CỤC THUẾ {cqt}\nSố:           /TB-QNG", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb0.rows[0].cells[1], f"CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n{ngay_thang_nam}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("THÔNG BÁO\n(Lần 1)\nVề việc giải trình, bổ sung thông tin, tài liệu")
        self._set_run_font_and_size(r_t, size_pt=14, is_bold=True)

        # 2. Table Kính gửi không viền
        tb_kg = self._tao_bang_khong_vien(doc, 1, 2, [Cm(3), Cm(15)])
        self._set_cell_text_preserving_format(tb_kg.rows[0].cells[0], "Kính gửi:", is_bold=True, size_pt=14)
        self._set_cell_text_preserving_format(tb_kg.rows[0].cells[1], f"{ten_nnt};\nMã số thuế: {mst};\nĐịa chỉ nhận thông báo: {dia_chi}", size_pt=14)

        doc.add_paragraph("Căn cứ quy định của pháp luật thuế hiện hành về trách nhiệm của người nộp thuế trong việc cung cấp thông tin theo yêu cầu của cơ quan thuế;")
        p_lr = doc.add_paragraph("Để làm rõ các nội dung:")
        self._set_run_font_and_size(p_lr.runs[0], size_pt=14, is_bold=True)

        for idx, item in enumerate(ds_noi_dung_chon):
            doc.add_paragraph(f"- {item.get('cau_van', '')}")

        doc.add_paragraph(f"{cqt} đề nghị {ten_nnt} có văn bản để giải trình và bổ sung thông tin, tài liệu sau đây:")
        for idx, item in enumerate(ds_noi_dung_chon):
            doc.add_paragraph(f"- {item.get('yeu_cau', 'Giải trình nguyên nhân và cung cấp chứng từ minh chứng.')}")

        doc.add_paragraph(f"Văn bản giải trình, thông tin, tài liệu bổ sung gửi về {cqt} trước ngày {ngay_hen}. Nếu có vướng mắc, đề nghị người nộp thuế liên hệ với {phong_doi} theo số điện thoại: {dien_thoai_lh}, địa chỉ: {dia_chi_cqt}.")
        doc.add_paragraph(f"{cqt} thông báo để {ten_nnt} biết và thực hiện./.")

        # Table chữ ký không viền
        tb_sig = self._tao_bang_khong_vien(doc, 1, 2, [Cm(8), Cm(10)])
        self._set_cell_text_preserving_format(tb_sig.rows[0].cells[0], f"Nơi nhận:\n- Như trên;\n- Lưu: VT, {phong_doi}.", size_pt=12.5)
        self._set_cell_text_preserving_format(tb_sig.rows[0].cells[1], f"TUQ. TRƯỞNG THUẾ TỈNH\n{chuc_vu}\n\n\n\n\n{lanh_dao}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        doc.save(file_dich)
        return True

    def _tao_to_trinh_moi(self, file_dich, thong_tin_chung, ds_noi_dung_chon, ds_can_cu_chon=None, la_luu_ho_so: bool = True):
        doc = Document()
        for s in doc.sections:
            s.top_margin, s.bottom_margin = Cm(2), Cm(2)
            s.left_margin, s.right_margin = Cm(3), Cm(2)

        cqt = thong_tin_chung.get("cqt", "THUẾ TỈNH QUẢNG NGÃI").upper()
        phong_doi = thong_tin_chung.get("phong_doi", "PHÒNG QLDN1").upper()
        ten_nnt = thong_tin_chung.get("ten_nnt", "DOANH NGHIỆP").upper()
        mst = thong_tin_chung.get("mst", "")
        dia_danh = thong_tin_chung.get("dia_danh", "Quảng Ngãi")
        cong_chuc = thong_tin_chung.get("cong_chuc", "Công chức đề xuất")
        lanh_dao = thong_tin_chung.get("lanh_dao", "Nguyễn Thị Phương")
        chuc_vu = thong_tin_chung.get("chuc_vu", "PHÓ TRƯỞNG PHÒNG").upper()

        now = datetime.datetime.now()
        ngay_chu = f"ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"
        dia_danh_ngay = f"{dia_danh}, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"

        # 1. Table 0 không viền: Tiêu ngữ và CQT
        tb0 = self._tao_bang_khong_vien(doc, 1, 2, [Cm(7.5), Cm(10.5)])
        self._set_cell_text_preserving_format(tb0.rows[0].cells[0], f"{cqt}\n{phong_doi}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)
        self._set_cell_text_preserving_format(tb0.rows[0].cells[1], "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do- Hạnh phúc", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if la_luu_ho_so:
            r_t = p_t.add_run(f"TỜ TRÌNH LÃNH ĐẠO THUẾ TỈNH\nVề việc phê duyệt lưu hồ sơ kiểm tra tại cơ quan thuế đối với\n{ten_nnt}.")
            self._set_run_font_and_size(r_t, size_pt=14, is_bold=True)
            doc.add_paragraph(f"1. Nội dung công việc:\nTrình Lãnh đạo Thuế tỉnh phê duyệt lưu hồ sơ kiểm tra tại cơ quan thuế đối với {ten_nnt}; MST: {mst}.")
        else:
            r_t = p_t.add_run(f"TỜ TRÌNH LÃNH ĐẠO THUẾ TỈNH\nVề việc ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với\n{ten_nnt}.")
            self._set_run_font_and_size(r_t, size_pt=14, is_bold=True)
            doc.add_paragraph(f"1. Nội dung công việc:\nTrình Lãnh đạo Thuế tỉnh phê duyệt ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với {ten_nnt}; MST: {mst}.")

        doc.add_paragraph("2. Căn cứ pháp lý để giải quyết:\n- Căn cứ Luật Quản lý thuế;\n- Căn cứ Quyết định số 1161/QĐ-CT ngày 14/8/2026 của Cục Thuế về việc ban hành Quy trình Kiểm tra thuế.")

        doc.add_paragraph("3. Đề xuất của công chức thuế giải quyết:")
        for idx, item in enumerate(ds_noi_dung_chon):
            doc.add_paragraph(f"+ {item.get('cau_van', '')}")

        if la_luu_ho_so:
            doc.add_paragraph("Kính trình Lãnh đạo Thuế tỉnh xem xét phê duyệt kết quả rà soát để lưu hồ sơ./.")
        else:
            doc.add_paragraph("Kính trình Lãnh đạo Thuế tỉnh xem xét phê duyệt ban hành Thông báo yêu cầu giải trình, bổ sung thông tin đối với đơn vị./.")

        p_d = doc.add_paragraph(dia_danh_ngay)
        p_d.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Table chữ ký cán bộ đề xuất
        tb_sig1 = self._tao_bang_khong_vien(doc, 1, 2, [Cm(9), Cm(9)])
        self._set_cell_text_preserving_format(tb_sig1.rows[0].cells[1], f"Cán bộ đề xuất\n\n\n\n\n{cong_chuc}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        if la_luu_ho_so:
            doc.add_paragraph(f"* Ý kiến của lãnh đạo Phòng {phong_doi}: Thống nhất theo đề xuất của công chức, kính trình Lãnh đạo Thuế tỉnh phê duyệt./.")
        else:
            doc.add_paragraph(f"* Ý kiến của lãnh đạo Phòng {phong_doi}: Thống nhất theo đề xuất của công chức, kính trình Lãnh đạo Thuế tỉnh ký ban hành Thông báo./.")

        # Table chữ ký lãnh đạo phòng
        tb_sig2 = self._tao_bang_khong_vien(doc, 1, 2, [Cm(9), Cm(9)])
        self._set_cell_text_preserving_format(tb_sig2.rows[0].cells[1], f"{dia_danh_ngay}\n{chuc_vu}\n\n\n\n\n{lanh_dao}", is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13)

        doc.add_paragraph("* Ý kiến phê duyệt của Lãnh đạo Thuế tỉnh:\n. . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .\n. . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .")

        doc.save(file_dich)
        return True
