import os
import time
import re
from PySide6.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QPushButton,
                               QLabel, QLineEdit, QFileDialog, QListWidget, QListWidgetItem,
                               QProgressBar, QTextEdit, QComboBox, QCheckBox, QGroupBox, QMessageBox)
from PySide6.QtCore import Qt, QThread, Signal
import qtawesome as qta
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from module_doc_xml.quan_ly_doc_to_khai import QuanLyDocToKhaiXML
from module_doc_xml.tu_dien_chi_tieu import MAP_TEN_TO_KHAI
from module_ho_tro.quan_ly_cong_thuc import QuanLyCongThuc
from module_phan_tich.engine_tinh_toan import EnginePhanTich
from module_phan_tich.bao_cao_doanh_so import BaoCaoDoanhSoEngine
from module_doc_hoadon.invoice_processor import InvoiceProcessor
from module_doc_hoadon.xu_ly_hoa_don_hub import doc_file_hoa_don_v2
from module_doc_hoadon.hoadon_ruiro import HoaDonRuiRoManager


# =========================================================================
# LUỒNG CHẠY NGẦM (WORKER): AUTO BOT AGENT (TỰ ĐỘNG HÓA 100% NHƯ THỦ CÔNG)
# =========================================================================
class AutoBotWorker(QThread):
    log_msg = Signal(str)
    progress_tong = Signal(int, int)
    progress_ct = Signal(str)
    finished_batch = Signal(int, int)
    sync_main_ui = Signal(dict)  # Bắn toàn bộ dữ liệu đã phân tích về Giao diện chính

    def __init__(self, thu_muc_goc, thu_muc_xuat, ds_mst, cau_hinh, parent=None):
        super().__init__(parent)
        self.thu_muc_goc = thu_muc_goc
        self.thu_muc_xuat = thu_muc_xuat
        self.ds_mst = ds_mst
        self.cau_hinh = cau_hinh
        self.is_running = True

    def _nap_va_phan_loai_ho_so_mst(self, path_mst):
        """
        Tự động thu thập và phân loại toàn bộ file trong thư mục của MST:
        - Hỗ trợ cấu trúc chia thư mục con (TK, HDMV, HDBR, CDSPS/TT200).
        - Hỗ trợ cả cấu trúc để phẳng tất cả file XML/Excel/CSV ngay trong thư mục MST (như kéo thả thủ công).
        """
        du_lieu_to_khai = {}
        ds_file_xml = []
        du_lieu_bcdsps_tt200 = {}
        dfs_hdmv = []
        dfs_hdbr = []

        quan_ly_xml = QuanLyDocToKhaiXML()
        allowed_cqt = self.cau_hinh.get("allowed_cqt", ["*"])

        def _lay_so_lan(thong_tin_dict):
            val = thong_tin_dict.get("LanBoSung")
            if val is None or str(val).strip() == "":
                val = thong_tin_dict.get("SoLan", 0)
            try:
                return int(val)
            except (ValueError, TypeError):
                return 0

        def _nap_1_file_xml(fpath):
            try:
                kq = quan_ly_xml.xu_ly_file_xml(fpath)
                if not kq:
                    return False
                ttc = kq.get("ThongTinChung", {})
                ma_tk = ttc.get("MaTKhai", "N/A")
                ky = ttc.get("KyKeKhai", "N/A")
                ma_cqt_hs = str(ttc.get("maCQTNoiNop", "")).strip()

                if allowed_cqt and "*" not in allowed_cqt:
                    if ma_cqt_hs and ma_cqt_hs not in allowed_cqt:
                        self.log_msg.emit(
                            f"   ⚠️ Bỏ qua tờ khai [{ma_tk}] kỳ {ky}: Mã CQT {ma_cqt_hs} không thuộc phạm vi phân quyền."
                        )
                        return False

                ten_form = MAP_TEN_TO_KHAI.get(ma_tk, ma_tk)
                lan_bs_moi = _lay_so_lan(ttc)

                if ten_form not in du_lieu_to_khai:
                    du_lieu_to_khai[ten_form] = {}
                if ky in du_lieu_to_khai[ten_form]:
                    tk_cu = du_lieu_to_khai[ten_form][ky]
                    lan_bs_cu = _lay_so_lan(tk_cu.get("ThongTinChung", {}))
                    if lan_bs_moi > lan_bs_cu:
                        du_lieu_to_khai[ten_form][ky] = kq
                else:
                    du_lieu_to_khai[ten_form][ky] = kq

                ds_file_xml.append(fpath)
                return True
            except Exception as e:
                self.log_msg.emit(f"   ❌ Lỗi đọc XML {os.path.basename(fpath)}: {e}")
                return False

        def _nap_1_file_cdsps(fpath):
            try:
                from module_ho_tro.kiem_tra_excel import kiem_tra_file_upload
                from module_ho_tro.tu_dien_tt200 import TU_DIEN_TK_TT200
                is_valid, _, df_sach, nam_phat_hien = kiem_tra_file_upload(fpath, TU_DIEN_TK_TT200)
                if is_valid and df_sach is not None and not df_sach.empty:
                    nam_nhap = str(nam_phat_hien) if nam_phat_hien else ""
                    if not nam_nhap:
                        cac_nam_tk = [
                            int(str(k)[-4:])
                            for d_k in du_lieu_to_khai.values()
                            for k in d_k.keys()
                            if str(k)[-4:].isdigit()
                        ]
                        nam_nhap = str(max(cac_nam_tk)) if cac_nam_tk else str(datetime.now().year)
                    du_lieu_bcdsps_tt200[nam_nhap] = df_sach
                    self.log_msg.emit(
                        f"   ✅ Đã nạp Bảng CĐSPS TT200 năm {nam_nhap} ({len(df_sach)} TK) từ {os.path.basename(fpath)}."
                    )
                    return True
            except Exception as e:
                self.log_msg.emit(f"   ⚠️ Không thể nạp CĐSPS từ {os.path.basename(fpath)}: {e}")
            return False

        def _nap_1_file_hoa_don(fpath, is_mua_vao):
            try:
                df = doc_file_hoa_don_v2(fpath, is_mua_vao=is_mua_vao)
                if df is None or df.empty:
                    # Fallback đọc trực tiếp nếu file không khớp mẫu chuẩn
                    df = pd.read_excel(fpath)
                if df is not None and not df.empty:
                    if is_mua_vao:
                        dfs_hdmv.append(df)
                    else:
                        dfs_hdbr.append(df)
                    return True
            except Exception as e:
                self.log_msg.emit(f"   ❌ Lỗi đọc hóa đơn {os.path.basename(fpath)}: {e}")
            return False

        # Duyệt đệ quy toàn bộ thư mục của MST để không bỏ sót bất kỳ file nào
        so_xml, so_mv, so_br, so_cd = 0, 0, 0, 0
        for root_dir, _, files in os.walk(path_mst):
            rel_dir = os.path.relpath(root_dir, path_mst).upper()
            for fname in files:
                if fname.startswith("~$") or fname.startswith("TMS_AUTO_"):
                    continue
                fpath = os.path.join(root_dir, fname)
                ext = os.path.splitext(fname)[1].lower()
                fname_lower = fname.lower()

                if ext == ".xml":
                    if _nap_1_file_xml(fpath):
                        so_xml += 1
                elif ext in [".xlsx", ".xls", ".csv"]:
                    # 1. Nếu nằm trong thư mục CDSPS / TT200 hoặc tên file có từ khóa CĐSPS
                    if any(k in rel_dir for k in ["CDSPS", "CDPS", "TT200"]) or any(
                        k in fname_lower for k in ["cdsps", "cdps", "tt200", "cân đối"]
                    ):
                        if _nap_1_file_cdsps(fpath):
                            so_cd += 1
                    # 2. Nếu nằm trong thư mục HDMV hoặc tên file là Mua vào
                    elif "HDMV" in rel_dir or any(
                        k in fname_lower for k in ["mua_vao", "muavao", "mua vào", "_in", "input", "đầu vào", "mua"]
                    ):
                        if _nap_1_file_hoa_don(fpath, is_mua_vao=True):
                            so_mv += 1
                    # 3. Nếu nằm trong thư mục HDBR hoặc mặc định là Bán ra
                    elif "HDBR" in rel_dir or not ("TK" in rel_dir):
                        if _nap_1_file_hoa_don(fpath, is_mua_vao=False):
                            so_br += 1

        self.log_msg.emit(
            f"   📂 Kết quả nạp hồ sơ: {so_xml} XML | {so_mv} file HĐ Mua vào | {so_br} file HĐ Bán ra | {so_cd} file CĐSPS."
        )
        df_hdmv = pd.concat(dfs_hdmv, ignore_index=True) if dfs_hdmv else pd.DataFrame()
        df_hdbr = pd.concat(dfs_hdbr, ignore_index=True) if dfs_hdbr else pd.DataFrame()
        return du_lieu_to_khai, ds_file_xml, du_lieu_bcdsps_tt200, df_hdmv, df_hdbr

    def run(self):
        import pythoncom
        pythoncom.CoInitialize()
        thanh_cong, that_bai = 0, 0
        tong_so = len(self.ds_mst)

        self.log_msg.emit(f"🚀 KHỞI ĐỘNG AGENT PHÂN TÍCH TỰ ĐỘNG CHO {tong_so} HỒ SƠ DOANH NGHIỆP...")
        dict_tong = QuanLyCongThuc.tai_du_lieu_tong_hop()
        config_tms = self.cau_hinh.get("config_tms", {})
        danh_sach_rule_chon = self.cau_hinh.get("danh_sach_rule_chon", None)
        rui_ro_mgr = HoaDonRuiRoManager()

        last_company_state = None

        for idx, mst in enumerate(self.ds_mst, 1):
            if not self.is_running:
                self.log_msg.emit("🛑 ĐÃ DỪNG TIẾN TRÌNH AUTO BOT BỞI NGƯỜI DÙNG!")
                break

            self.progress_tong.emit(idx, tong_so)
            self.log_msg.emit(f"\n{'=' * 60}\n🏢 ĐANG XỬ LÝ HỒ SƠ: {mst} ({idx}/{tong_so})\n{'=' * 60}")

            try:
                path_mst = self.thu_muc_goc if mst == os.path.basename(self.thu_muc_goc) and not os.path.isdir(
                    os.path.join(self.thu_muc_goc, mst)
                ) else os.path.join(self.thu_muc_goc, mst)

                # ===============================================================
                # 1. NẠP TOÀN DIỆN XML, HÓA ĐƠN CHUẨN HÓA & BẢNG CĐSPS TT200
                # ===============================================================
                self.progress_ct.emit(f"[{mst}] Đang nạp và chuẩn hóa Tờ khai, Hóa đơn, CĐSPS...")
                du_lieu_to_khai, ds_file_xml, du_lieu_bcdsps_tt200, df_hdmv, df_hdbr = self._nap_va_phan_loai_ho_so_mst(
                    path_mst
                )

                # ===============================================================
                # 2. QUÉT RADAR HÓA ĐƠN RỦI RO & HÓA ĐƠN TRẠNG THÁI ĐẶC BIỆT
                # ===============================================================
                df_hd_rui_ro = pd.DataFrame()
                if not df_hdmv.empty:
                    self.progress_ct.emit(f"[{mst}] Đang rà soát Hóa đơn Rủi ro...")
                    df_hdmv, rr_warnings = rui_ro_mgr.check_bang_ke_01_1(df_hdmv)
                    for w in rr_warnings:
                        self.log_msg.emit(f"   {w}")
                    if 'Đối chiếu HĐ rủi ro' in df_hdmv.columns:
                        df_hd_rui_ro = df_hdmv[df_hdmv['Đối chiếu HĐ rủi ro'] == "Thuộc DS hóa đơn rủi ro"].copy()

                ds_dac_biet = []
                if not df_hdbr.empty:
                    df_db_br = InvoiceProcessor.loc_hoa_don_trang_thai_dac_biet(df_hdbr, "Bán ra")
                    if not df_db_br.empty:
                        ds_dac_biet.append(df_db_br)
                if not df_hdmv.empty:
                    df_db_mv = InvoiceProcessor.loc_hoa_don_trang_thai_dac_biet(df_hdmv, "Mua vào")
                    if not df_db_mv.empty:
                        ds_dac_biet.append(df_db_mv)
                df_hd_dac_biet = pd.concat(ds_dac_biet, ignore_index=True) if ds_dac_biet else pd.DataFrame()
                if not df_hd_dac_biet.empty:
                    self.log_msg.emit(
                        f"   ⚠️ Phát hiện {len(df_hd_dac_biet)} hóa đơn có trạng thái đặc biệt (Hủy/Xóa bỏ/Thay thế/Điều chỉnh)."
                    )

                # ===============================================================
                # 3. TRẠM TÍCH HỢP TMS (ĐỒNG BỘ NGÀY NỘP & TRA CỨU TRẠNG THÁI NNT)
                # ===============================================================
                from PySide6.QtCore import QEventLoop
                df_nnt_bat_thuong = pd.DataFrame()

                # A. ĐỒNG BỘ NGÀY NỘP TỜ KHAI
                if self.cau_hinh.get('dong_bo_ngay'):
                    if not config_tms.get("user_tms"):
                        self.log_msg.emit("   ⏭️ [TMS] Bỏ qua tra Ngày nộp do chưa cấu hình Tài khoản SAP.")
                    else:
                        self.progress_ct.emit(f"[{mst}] Đang đồng bộ ngày nộp tờ khai từ TMS...")
                        try:
                            from module_ho_tro.module_dong_bo_tms import BotSAPWorker
                            cac_nam_tms = [
                                int(str(ky)[-4:])
                                for dict_ky in du_lieu_to_khai.values()
                                for ky in dict_ky.keys()
                                if str(ky)[-4:].isdigit()
                            ]
                            nam_nho_nhat = min(cac_nam_tms) if cac_nam_tms else datetime.now().year
                            tu_ngay = config_tms.get("tu_ngay") or f"01.01.{nam_nho_nhat}"
                            den_ngay = config_tms.get("den_ngay") or datetime.now().strftime("%d.%m.%Y")
                            cqt = config_tms.get("cqt", "50500")

                            ngay_tai = datetime.now().strftime("%d%m%Y")
                            ten_file = f"TMS_AUTO_{mst}_{ngay_tai}_{tu_ngay.replace('.', '')}_{den_ngay.replace('.', '')}.xlsx"
                            duong_dan_cache = os.path.join(self.thu_muc_xuat, ten_file)

                            bot_sap = BotSAPWorker(
                                cqt, tu_ngay, den_ngay, mst, du_lieu_to_khai,
                                config_tms["user_tms"], config_tms["pass_tms"], duong_dan_cache
                            )
                            loop = QEventLoop()
                            ket_qua_bot = {}

                            def on_bot_sap_finished(tc, kq, sl):
                                ket_qua_bot['thanh_cong'] = tc
                                ket_qua_bot['ket_qua'] = kq
                                ket_qua_bot['so_luong'] = sl
                                loop.quit()

                            bot_sap.progress.connect(lambda msg: self.log_msg.emit(f"      [BotSAP] {msg}"))
                            bot_sap.finished.connect(on_bot_sap_finished)
                            bot_sap.start()
                            loop.exec()

                            if ket_qua_bot.get('thanh_cong'):
                                self.log_msg.emit(
                                    f"   ✅ [TMS] Đã đồng bộ ngày nộp thực tế cho {ket_qua_bot.get('so_luong')} Tờ khai."
                                )
                            else:
                                self.log_msg.emit(f"   ❌ [TMS] Lỗi BotSAPWorker: {ket_qua_bot.get('ket_qua')}")
                        except Exception as e:
                            self.log_msg.emit(f"   ❌ [TMS] Lỗi đồng bộ ngày nộp: {e}")

                # B. TRA CỨU TRẠNG THÁI HOẠT ĐỘNG NNT TRÊN HÓA ĐƠN
                if self.cau_hinh.get('check_tms'):
                    self.progress_ct.emit(f"[{mst}] Đang tra cứu trạng thái MST đối tác trên TMS...")
                    try:
                        mst_can_tra_cuu = set()
                        c_mv, c_br = "", ""

                        if not df_hdmv.empty:
                            for c_cand in ['MSTNguoiBan', 'Mã số thuế người bán', 'MST người bán', 'MST người bán/MST người xuất hàng']:
                                if c_cand in df_hdmv.columns:
                                    c_mv = c_cand
                                    break
                            if c_mv:
                                mst_can_tra_cuu.update(
                                    df_hdmv[c_mv].dropna().astype(str).str.strip().str.replace(r'\.0$', '', regex=True).unique()
                                )

                        if not df_hdbr.empty:
                            for c_cand in ['MSTNguoiMua', 'MST người mua', 'MST người mua/MST người nhận hàng']:
                                if c_cand in df_hdbr.columns:
                                    c_br = c_cand
                                    break
                            if c_br:
                                mst_can_tra_cuu.update(
                                    df_hdbr[c_br].dropna().astype(str).str.strip().str.replace(r'\.0$', '', regex=True).unique()
                                )

                        mst_can_tra_cuu = [
                            m.strip() for m in mst_can_tra_cuu
                            if m.strip() != "" and m.strip().lower() not in ["nan", "none"]
                        ]

                        if mst_can_tra_cuu:
                            self.log_msg.emit(f"   🔍 [TMS] Đang quét {len(mst_can_tra_cuu)} MST đối tác qua SAP...")
                            from module_ho_tro.sap_worker import SAPWorker
                            sap_robot = SAPWorker(list(mst_can_tra_cuu))
                            sap_robot.progress.connect(lambda step, msg: self.log_msg.emit(f"      [SAPWorker] {msg}"))
                            sap_robot.error.connect(lambda msg: self.log_msg.emit(f"      ❌ [SAPWorker] {msg}"))

                            ket_qua_tms = sap_robot.run_check()
                            if ket_qua_tms:
                                df_tms = pd.DataFrame(
                                    ket_qua_tms,
                                    columns=['MST_SAP', 'TenNNT', 'CQT_TMS', 'MaTrangThai', 'TrangThaiTMS', 'NgayCapNhat']
                                )
                                df_map = df_tms[['MST_SAP', 'TrangThaiTMS', 'CQT_TMS']].drop_duplicates(subset=['MST_SAP'])

                                def _map_tms(df_hd, col_mst):
                                    if df_hd.empty or not col_mst:
                                        return df_hd
                                    for col in ['TrangThaiTMS', 'CQT_TMS']:
                                        if col in df_hd.columns:
                                            df_hd.drop(columns=[col], inplace=True)
                                    df_hd[col_mst] = df_hd[col_mst].fillna('').astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
                                    df_new = pd.merge(df_hd, df_map, how='left', left_on=col_mst, right_on='MST_SAP')
                                    df_new['TrangThaiTMS'] = df_new['TrangThaiTMS'].fillna('Chưa xác định')
                                    df_new['CQT_TMS'] = df_new['CQT_TMS'].fillna('')
                                    df_new.drop(columns=['MST_SAP'], inplace=True, errors='ignore')
                                    return df_new

                                df_hdmv = _map_tms(df_hdmv, c_mv)
                                df_hdbr = _map_tms(df_hdbr, c_br)

                                # Lọc danh sách NNT Bất thường (TMS) giống hệt luồng thủ công
                                trang_thai_an_toan = ['00', '04']
                                tu_khoa_bat_thuong = [
                                    "không hoạt động", "bỏ địa chỉ", "chờ xác minh", "mất tích",
                                    "ngừng hoạt động", "tạm nghỉ", "tạm ngừng", "chấm dứt", "đóng mst"
                                ]

                                def _la_nnt_bat_thuong(row_tms):
                                    ma_tt = str(row_tms.get('MaTrangThai', '')).strip()
                                    ten_tt = str(row_tms.get('TrangThaiTMS', '')).strip().lower()
                                    if any(k in ten_tt for k in tu_khoa_bat_thuong):
                                        return True
                                    if ma_tt and ma_tt not in trang_thai_an_toan:
                                        if "đang hoạt động" in ten_tt and not any(k in ten_tt for k in tu_khoa_bat_thuong):
                                            return False
                                        return True
                                    return False

                                df_nnt_bat_thuong = df_tms[df_tms.apply(_la_nnt_bat_thuong, axis=1)].copy()
                                if not df_nnt_bat_thuong.empty:
                                    df_nnt_bat_thuong.rename(
                                        columns={
                                            'MST_SAP': 'Mã số thuế', 'TenNNT': 'Tên người nộp thuế',
                                            'CQT_TMS': 'Cơ quan thuế quản lý', 'MaTrangThai': 'Mã TT',
                                            'TrangThaiTMS': 'Trạng thái hoạt động', 'NgayCapNhat': 'Ngày cập nhật'
                                        },
                                        inplace=True
                                    )
                                    self.log_msg.emit(
                                        f"   🚨 [TMS] Phát hiện {len(df_nnt_bat_thuong)} NNT có trạng thái hoạt động bất thường!"
                                    )
                                else:
                                    self.log_msg.emit("   ✅ [TMS] 100% NNT đối tác đang hoạt động bình thường.")
                    except Exception as e:
                        self.log_msg.emit(f"   ❌ [TMS] Lỗi tra cứu trạng thái hóa đơn: {e}")

                # ===============================================================
                # 4. CHẠY ĐẦY ĐỦ CÁC ENGINE PHÂN TÍCH (NHƯ THỦ CÔNG 100%)
                # ===============================================================
                self.progress_ct.emit(f"[{mst}] Đang chạy Engine Đối chiếu, Chuỗi cung ứng & Tồn kho...")
                engine = EnginePhanTich(dict_cong_thuc=dict_tong, log_callback=self.log_msg.emit)

                nam_pt = self.cau_hinh['nam_pt']
                if nam_pt != "Tất cả các năm":
                    danh_sach_nam = [nam_pt]
                else:
                    cac_nam = {
                        str(k)[-4:] for d in du_lieu_to_khai.values() for k in d.keys() if str(k)[-4:].isdigit()
                    }
                    if not cac_nam and '_Nam' in df_hdbr.columns:
                        cac_nam.update(df_hdbr['_Nam'].dropna().astype(str).unique())
                    if not cac_nam and '_Nam' in df_hdmv.columns:
                        cac_nam.update(df_hdmv['_Nam'].dropna().astype(str).unique())
                    danh_sach_nam = sorted([n for n in cac_nam if len(n) == 4 and n.isdigit()])
                if not danh_sach_nam:
                    danh_sach_nam = [str(datetime.now().year)]

                # Làm sạch hóa đơn (Bỏ xóa bỏ, hủy, bị thay thế; đảo dấu âm cho điều chỉnh giảm)
                df_br_sach = InvoiceProcessor.loc_hoa_don_tinh_toan(df_hdbr)
                df_mv_sach = InvoiceProcessor.loc_hoa_don_tinh_toan(df_hdmv)

                # Chạy Rà soát Tồn kho cơ bản & Đối chiếu Mặt hàng AI trên tập hóa đơn sạch
                kq_ton_kho = []
                df_ton_kho_ai = pd.DataFrame()
                if (isinstance(df_mv_sach, pd.DataFrame) and not df_mv_sach.empty) or (
                    isinstance(df_br_sach, pd.DataFrame) and not df_br_sach.empty
                ):
                    try:
                        from module_phan_tich.inventory_matcher import InventoryReconcilerLight
                        reconciler_light = InventoryReconcilerLight(df_mv_sach, df_br_sach)
                        kq_ton_kho = reconciler_light.analyze_inventory() or []
                        self.log_msg.emit(f"   📦 Đã phân tích đối chiếu Tồn kho ({len(kq_ton_kho)} nhóm mặt hàng).")
                    except Exception as e:
                        self.log_msg.emit(f"   ⚠️ Lỗi phân tích Tồn kho cơ bản: {e}")

                    try:
                        from module_phan_tich.inventory_matcher1 import InventoryReconciler
                        from module_ho_tro.dictionary_loader import load_synonyms
                        matcher_ai = InventoryReconciler(df_mv_sach, df_br_sach, synonyms=load_synonyms())
                        df_ton_kho_ai = matcher_ai.analyze_discrepancies(threshold=0.4)
                        if not df_ton_kho_ai.empty:
                            self.log_msg.emit(
                                f"   🤖 Đã đối chiếu tên mặt hàng bằng AI ({len(df_ton_kho_ai)} mặt hàng bán ra)."
                            )
                    except Exception as e:
                        self.log_msg.emit(f"   ⚠️ Lỗi đối chiếu mặt hàng AI: {e}")

                kho_ket_qua_nam = {}
                for nam in danh_sach_nam:
                    self.log_msg.emit(f"\n   ⚙️ --- ĐANG TÍNH TOÁN RỦI RO & CHUỖI CUNG ỨNG NĂM {nam} ---")
                    ket_qua_list = engine.thuc_thi_phan_tich(
                        du_lieu_to_khai, df_hdbr, df_hdmv,
                        nam_phan_tich=nam,
                        df_cdsps_tt200=du_lieu_bcdsps_tt200,
                        loai_bctc=self.cau_hinh['loai_bctc'],
                        danh_sach_rule_chon=danh_sach_rule_chon
                    )

                    du_lieu_br = BaoCaoDoanhSoEngine.xuat_du_lieu_bao_cao(
                        df_br_sach, du_lieu_to_khai, nam, False, self.log_msg.emit
                    )
                    du_lieu_mv = BaoCaoDoanhSoEngine.xuat_du_lieu_bao_cao(
                        df_mv_sach, du_lieu_to_khai, nam, True, self.log_msg.emit
                    )

                    # Phân tích Chuỗi cung ứng cho năm hiện tại
                    sc_result = None
                    try:
                        from module_phan_tich.supply_chain_analyzer import SupplyChainAnalyzer
                        df_mv_nam = df_mv_sach[df_mv_sach['_Nam'] == str(nam)] if (
                            isinstance(df_mv_sach, pd.DataFrame) and not df_mv_sach.empty and '_Nam' in df_mv_sach.columns
                        ) else df_mv_sach
                        df_br_nam = df_br_sach[df_br_sach['_Nam'] == str(nam)] if (
                            isinstance(df_br_sach, pd.DataFrame) and not df_br_sach.empty and '_Nam' in df_br_sach.columns
                        ) else df_br_sach
                        if (isinstance(df_mv_nam, pd.DataFrame) and not df_mv_nam.empty) or (
                            isinstance(df_br_nam, pd.DataFrame) and not df_br_nam.empty
                        ):
                            sc_result = SupplyChainAnalyzer(df_mv_nam, df_br_nam).run_analysis()
                    except Exception as e:
                        self.log_msg.emit(f"   ⚠️ Lỗi phân tích Chuỗi cung ứng năm {nam}: {e}")

                    kho_ket_qua_nam[nam] = {
                        "ket_qua_list": ket_qua_list,
                        "du_lieu_ban_ra": du_lieu_br,
                        "du_lieu_mua_vao": du_lieu_mv,
                        "df_br_sach": df_br_sach,
                        "df_mv_sach": df_mv_sach,
                        "sc_result": sc_result
                    }

                    # ===============================================================
                    # 5. KẾT XUẤT FULL EXCEL TỰ ĐỘNG (ĐẦY ĐỦ 100% SHEET NHƯ THỦ CÔNG)
                    # ===============================================================
                    self.progress_ct.emit(f"[{mst}] Đang kết xuất Báo cáo Excel toàn diện năm {nam}...")
                    file_name = f"BaoCaoAuto_{mst}_Nam_{nam}.xlsx"
                    path_save = os.path.join(self.thu_muc_xuat, file_name)

                    self._ghi_full_excel_bot(
                        duong_dan_luu=path_save,
                        ds_kq=ket_qua_list,
                        du_lieu_br=du_lieu_br,
                        du_lieu_mv=du_lieu_mv,
                        dict_to_khai=du_lieu_to_khai,
                        df_hdbr=df_hdbr,
                        df_hdmv=df_hdmv,
                        nam=nam,
                        du_lieu_bcdsps_tt200=du_lieu_bcdsps_tt200,
                        df_hd_rui_ro=df_hd_rui_ro,
                        df_hd_dac_biet=df_hd_dac_biet,
                        df_nnt_bat_thuong=df_nnt_bat_thuong,
                        sc_result=sc_result,
                        kq_ton_kho=kq_ton_kho,
                        df_ton_kho_ai=df_ton_kho_ai
                    )
                    self.log_msg.emit(f"   💾 Đã lưu Báo cáo Excel đầy đủ: {file_name}")

                # Lưu lại trạng thái công ty vừa chạy xong để đồng bộ lên Giao diện chính
                last_company_state = {
                    "mst": mst,
                    "du_lieu_to_khai": du_lieu_to_khai,
                    "ds_file_xml": ds_file_xml,
                    "du_lieu_bcdsps_tt200": du_lieu_bcdsps_tt200,
                    "df_hdbr": df_hdbr,
                    "df_hdmv": df_hdmv,
                    "df_nnt_bat_thuong": df_nnt_bat_thuong,
                    "kho_ket_qua_nam": kho_ket_qua_nam,
                    "danh_sach_nam": danh_sach_nam,
                    "kq_ton_kho": kq_ton_kho
                }

                self.log_msg.emit(f"✅ HOÀN THÀNH HỒ SƠ MST: {mst}")
                thanh_cong += 1

            except Exception as e:
                import traceback
                self.log_msg.emit(f"❌ THẤT BẠI tại MST {mst}.\n   Chi tiết lỗi: {str(e)}\n{traceback.format_exc()}")
                that_bai += 1

        # Đồng bộ kết quả lên Giao diện chính nếu người dùng bật tùy chọn Agent Sync
        if last_company_state and self.cau_hinh.get("sync_ui", True):
            self.progress_ct.emit("Đang đồng bộ toàn bộ bảng biểu lên Giao diện chính...")
            self.sync_main_ui.emit(last_company_state)

        self.progress_ct.emit("Đã chạy xong toàn bộ danh sách!")
        self.finished_batch.emit(thanh_cong, that_bai)

    # =========================================================================
    # BỘ CÔNG CỤ XUẤT EXCEL FULL OPTION (ĐẦY ĐỦ 100% SHEET NHƯ THỦ CÔNG)
    # =========================================================================
    def _dinh_dang_tieu_de(self, sheet, danh_sach_cot):
        fill_header = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
        font_header = Font(color="FFFFFF", bold=True)
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border_thin = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        sheet.append(danh_sach_cot)
        for col_idx in range(1, len(danh_sach_cot) + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_thin

    def _ghi_so_excel(self, sheet, row, col, gia_tri):
        cell = sheet.cell(row=row, column=col)
        cell.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        try:
            if pd.isna(gia_tri) or gia_tri == "":
                val_num = 0.0
            elif isinstance(gia_tri, (int, float)):
                val_num = float(gia_tri)
            else:
                str_val = re.sub(r'[^\d.,-]', '', str(gia_tri).strip())
                if not str_val or str_val == "-":
                    cell.value = str(gia_tri)
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    return
                elif str_val.count(',') > 1:
                    val_num = float(str_val.replace(',', ''))
                elif ',' in str_val and '.' in str_val:
                    val_num = float(str_val.replace('.', '').replace(',', '.')) if str_val.rfind(',') > str_val.rfind('.') else float(str_val.replace(',', ''))
                elif ',' in str_val:
                    val_num = float(str_val.replace(',', '')) if len(str_val.split(',')[-1]) == 3 else float(str_val.replace(',', '.'))
                elif '.' in str_val and str_val.count('.') > 1:
                    val_num = float(str_val.replace('.', ''))
                else:
                    val_num = float(str_val)
            cell.value = val_num
            cell.number_format = '#,##0' if val_num.is_integer() else '#,##0.00'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        except Exception:
            cell.value = str(gia_tri)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _ghi_du_lieu_thong_minh(self, sheet, row, col, ten_cot, gia_tri):
        cell = sheet.cell(row=row, column=col)
        cell.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        val_str = str(gia_tri).strip() if pd.notna(gia_tri) else ""
        ten_cot_lower = str(ten_cot).lower()

        nhom_text = [
            "mst", "mã số thuế", "masothue", "cmnd", "cccd", "hộ chiếu",
            "số giấy tờ", "điện thoại", "tài khoản", "số hđ", "số hóa đơn", "sohoadon", "ký hiệu", "kyhieu"
        ]
        if any(k in ten_cot_lower for k in nhom_text):
            if val_str.endswith(".0"):
                val_str = val_str[:-2]
            cell.value = val_str
            cell.number_format = '@'
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            return

        if "ngày" in ten_cot_lower or "ngay" in ten_cot_lower:
            if re.match(r'^\d{4}-\d{2}-\d{2}', val_str):
                try:
                    cell.value = datetime.strptime(val_str[:10], '%Y-%m-%d')
                    cell.number_format = 'dd/mm/yyyy'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    return
                except Exception:
                    pass
            elif re.match(r'^\d{2}/\d{2}/\d{4}', val_str):
                cell.value = val_str
                cell.number_format = 'dd/mm/yyyy'
                cell.alignment = Alignment(horizontal="center", vertical="center")
                return
        self._ghi_so_excel(sheet, row, col, gia_tri)

    def _can_chinh_kich_thuoc_cot(self, sheet):
        font_chuan = Font(name="Times New Roman", size=11)
        for col_idx in range(1, sheet.max_column + 1):
            max_length = 0
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            for row_idx in range(1, sheet.max_row + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                if type(cell).__name__ == 'MergedCell':
                    continue
                if getattr(cell, 'font', None):
                    cell.font = Font(
                        name="Times New Roman", size=11,
                        bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color
                    )
                else:
                    cell.font = font_chuan
                if col_idx == 1 and getattr(cell, 'font', None) and cell.font.bold and cell.value and len(str(cell.value)) > 35:
                    continue
                try:
                    if cell.value is not None:
                        for line in str(cell.value).split('\n'):
                            if len(line) > max_length:
                                max_length = len(line)
                except Exception:
                    pass
            adjusted_width = max_length + 2.0 if max_length > 0 else 12
            if adjusted_width < 5:
                adjusted_width = 5
            elif adjusted_width > 55:
                adjusted_width = 55
            sheet.column_dimensions[col_letter].width = adjusted_width

    def _ghi_full_excel_bot(
        self, duong_dan_luu, ds_kq, du_lieu_br, du_lieu_mv, dict_to_khai,
        df_hdbr, df_hdmv, nam, du_lieu_bcdsps_tt200=None, df_hd_rui_ro=None,
        df_hd_dac_biet=None, df_nnt_bat_thuong=None, sc_result=None,
        kq_ton_kho=None, df_ton_kho_ai=None
    ):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        border_thin = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_wrap = Alignment(vertical="center", wrap_text=True)
        from openpyxl.comments import Comment

        # =========================================================
        # 1. VẼ KẾT QUẢ PHÂN TÍCH (BCTC, NGANG, ĐỐI CHIẾU CHÉO)
        # =========================================================
        if ds_kq:
            def _dich_ten_chi_tieu(chuoi_goc):
                try:
                    from module_doc_xml.tu_dien_chi_tieu import TU_DIEN_TO_KHAI
                except Exception:
                    TU_DIEN_TO_KHAI = {}
                if not chuoi_goc:
                    return ""
                chuoi_str = str(chuoi_goc)
                if "Chỉ tiêu" in chuoi_str or "=" in chuoi_str:
                    return chuoi_str
                ma_ct = chuoi_str.split("|")[-1].replace("]", "").strip() if "|" in chuoi_str else chuoi_str.replace("[", "").replace("]", "").strip()
                for form_data in TU_DIEN_TO_KHAI.values():
                    if isinstance(form_data, dict):
                        for phan_he, phan_he_data in form_data.items():
                            if isinstance(phan_he_data, dict) and ma_ct in phan_he_data:
                                return f"[{ma_ct}] {phan_he_data[ma_ct]}"
                            elif isinstance(phan_he_data, str) and phan_he == ma_ct:
                                return f"[{ma_ct}] {phan_he_data}"
                return f"Chỉ tiêu {ma_ct}"

            ds_cheo, ds_bctc, ds_ngang = [], [], []
            for r in ds_kq:
                nhom = str(r.get('nhom', '')).lower()
                if "chéo" in nhom:
                    ds_cheo.append(r)
                elif "tài chính" in nhom or "bctc" in nhom:
                    ds_bctc.append(r)
                elif "ngang" in nhom or "tỷ lệ" in nhom or "ty le" in nhom:
                    ds_ngang.append(r)
                else:
                    ds_bctc.append(r)

            def _ve_bang_7_cot(ten_sheet, data_list):
                if not data_list:
                    return
                ws = wb.create_sheet(ten_sheet)
                cols = [
                    "Nhóm Phân Tích", "Tên Phép Kiểm Tra (Logic)", "Số liệu Vế Trái (A)",
                    "Số liệu Vế Phải (B)", "Chênh lệch (A - B)", "Đánh giá", "Khuyến nghị / Thông điệp"
                ]
                self._dinh_dang_tieu_de(ws, cols)

                for r_idx, r in enumerate(data_list, start=2):
                    is_rui_ro = r.get('is_rui_ro', False)
                    mau_nen = PatternFill(
                        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
                    ) if r_idx % 2 == 0 else PatternFill(
                        start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
                    )

                    ws.cell(row=r_idx, column=1, value=r.get('nhom', '')).alignment = align_wrap
                    ws.cell(row=r_idx, column=2, value=r.get('ten_phan_tich', '')).alignment = align_wrap

                    self._ghi_so_excel(ws, r_idx, 3, r.get('gia_tri_a', 0))
                    ws.cell(row=r_idx, column=3).comment = Comment(_dich_ten_chi_tieu(r.get('ve_a_goc', '')), "Chi tiết")
                    ws.cell(row=r_idx, column=3).font = Font(bold=True)

                    if r.get('is_ab'):
                        self._ghi_so_excel(ws, r_idx, 4, r.get('gia_tri_b', 0))
                        ws.cell(row=r_idx, column=4).comment = Comment(_dich_ten_chi_tieu(r.get('ve_b_goc', '')), "Chi tiết")
                    else:
                        ws.cell(row=r_idx, column=4, value="---").alignment = align_center
                    ws.cell(row=r_idx, column=4).font = Font(bold=True)

                    so_chenh_lech = r.get('ket_qua_so', 0)
                    self._ghi_so_excel(ws, r_idx, 5, so_chenh_lech)
                    ws.cell(row=r_idx, column=5).font = Font(color="DC2626", bold=True) if is_rui_ro else Font(bold=True)

                    tt_text = r.get('trang_thai', '')
                    c_tt = ws.cell(row=r_idx, column=6, value=tt_text)
                    c_tt.alignment = align_center
                    if is_rui_ro:
                        c_tt.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
                        c_tt.font = Font(color="FFFFFF", bold=True)
                    elif "Thiếu" in str(so_chenh_lech) or "Không đủ" in tt_text:
                        c_tt.fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
                        c_tt.font = Font(color="FFFFFF", bold=True)
                    else:
                        c_tt.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
                        c_tt.font = Font(color="FFFFFF", bold=True)

                    ws.cell(row=r_idx, column=7, value=r.get('thong_diep', '')).alignment = align_wrap
                    for c_idx in range(1, 8):
                        c = ws.cell(row=r_idx, column=c_idx)
                        if c_idx != 6:
                            c.fill = mau_nen
                        c.border = border_thin

            _ve_bang_7_cot("PT Báo cáo Tài chính", ds_bctc)
            _ve_bang_7_cot("Phân tích Ngang", ds_ngang)

            if ds_cheo:
                ws_cheo = wb.create_sheet("Đối chiếu chéo Tờ khai")

                def _lay_ma_form_excel(chuoi_goc):
                    if not chuoi_goc:
                        return "Dữ liệu"
                    match = re.search(
                        r'(01/GTGT|02/GTGT|03/TNDN|04/TNDN|05/QTT[-_]TNCN|05/KK[-_]TNCN|BCTC[-_]TT133|BCTC[-_]TT200|01/TAIN|02/TAIN|TAIN)',
                        str(chuoi_goc)
                    )
                    if match:
                        return match.group(1).replace('_', '-')
                    if "HD_" in str(chuoi_goc) or "bán ra" in str(chuoi_goc).lower():
                        return "Hóa đơn Bán ra"
                    if "HD_MUA_VAO" in str(chuoi_goc) or "mua vào" in str(chuoi_goc).lower():
                        return "Hóa đơn Mua vào"
                    return "Dữ liệu"

                dict_nhom = {}
                for r in ds_cheo:
                    form_a = _lay_ma_form_excel(r.get('ve_a_goc', ''))
                    form_b = _lay_ma_form_excel(r.get('ve_b_goc', ''))
                    ten_cap = f"{form_a} VÀ {form_b}" if r.get('is_ab') else f"{form_a}"
                    dict_nhom.setdefault(ten_cap, []).append(r)

                r_idx = 1
                for _, ds_loi in dict_nhom.items():
                    ma_a = _lay_ma_form_excel(ds_loi[0].get('ve_a_goc', ''))
                    ma_b = _lay_ma_form_excel(ds_loi[0].get('ve_b_goc', ''))
                    tk_a = ma_a if "hóa đơn" in ma_a.lower() else f"Tờ khai {ma_a}"
                    tk_b = (ma_b if "hóa đơn" in ma_b.lower() else f"Tờ khai {ma_b}") if ds_loi[0].get('is_ab') else "Dữ liệu Vế Phải (B)"
                    cols = ["STT", "Nội dung đối chiếu", tk_a, tk_b, "Chênh lệch (A-B)", "Đánh giá", "Khuyến nghị / Thông điệp"]

                    fill_header = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
                    for c_i, c_name in enumerate(cols, start=1):
                        c = ws_cheo.cell(row=r_idx, column=c_i, value=c_name)
                        c.fill = fill_header
                        c.font = Font(color="0F172A", bold=True)
                        c.alignment = align_center
                        c.border = border_thin
                    r_idx += 1

                    for stt, r in enumerate(ds_loi, start=1):
                        is_rui_ro = r.get('is_rui_ro', False)
                        mau_nen = PatternFill(
                            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
                        ) if stt % 2 != 0 else PatternFill(
                            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
                        )

                        ws_cheo.cell(row=r_idx, column=1, value=stt).alignment = align_center
                        ws_cheo.cell(row=r_idx, column=2, value=r.get('ten_phan_tich', '')).alignment = align_wrap

                        self._ghi_so_excel(ws_cheo, r_idx, 3, r.get('gia_tri_a', 0))
                        ws_cheo.cell(row=r_idx, column=3).comment = Comment(_dich_ten_chi_tieu(r.get('ve_a_goc', '')), "Chi tiết")
                        ws_cheo.cell(row=r_idx, column=3).font = Font(bold=True)

                        if r.get('is_ab'):
                            self._ghi_so_excel(ws_cheo, r_idx, 4, r.get('gia_tri_b', 0))
                            ws_cheo.cell(row=r_idx, column=4).comment = Comment(_dich_ten_chi_tieu(r.get('ve_b_goc', '')), "Chi tiết")
                        else:
                            ws_cheo.cell(row=r_idx, column=4, value="---").alignment = align_center
                        ws_cheo.cell(row=r_idx, column=4).font = Font(bold=True)

                        so_chenh_lech = r.get('ket_qua_so', 0)
                        self._ghi_so_excel(ws_cheo, r_idx, 5, so_chenh_lech)
                        ws_cheo.cell(row=r_idx, column=5).font = Font(color="DC2626", bold=True) if is_rui_ro else Font(bold=True)

                        tt_text = r.get('trang_thai', '')
                        c_tt = ws_cheo.cell(row=r_idx, column=6, value=tt_text)
                        c_tt.alignment = align_center
                        if is_rui_ro:
                            c_tt.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
                            c_tt.font = Font(color="FFFFFF", bold=True)
                        elif "Thiếu" in str(so_chenh_lech) or "Không đủ" in tt_text:
                            c_tt.fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
                            c_tt.font = Font(color="FFFFFF", bold=True)
                        else:
                            c_tt.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
                            c_tt.font = Font(color="FFFFFF", bold=True)

                        ws_cheo.cell(row=r_idx, column=7, value=r.get('thong_diep', '')).alignment = align_wrap
                        for c_idx in range(1, 8):
                            c = ws_cheo.cell(row=r_idx, column=c_idx)
                            if c_idx != 6:
                                c.fill = mau_nen
                            c.border = border_thin
                        r_idx += 1

        # =========================================================
        # 2. VẼ BÁO CÁO ĐỐI CHIẾU DOANH SỐ HÓA ĐƠN & TỜ KHAI
        # =========================================================
        if du_lieu_br or du_lieu_mv:
            ws_bc = wb.create_sheet("Báo cáo Đối chiếu HĐ")
            dong_ht = 1

            def _ve_bang_bao_cao_excel(ws, start_row, tieu_de, data):
                if not data:
                    return start_row
                ws.cell(row=start_row, column=1, value=tieu_de).font = Font(bold=True, size=11, color="000000")
                start_row += 1
                fill_header = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
                for r_offset, h_list in enumerate([
                    ["STT", "Kỳ", "Hóa đơn điện tử", "", "Hồ sơ khai thuế", "", "Chênh lệch (KK-Hóa đơn)", ""],
                    ["", "", "Giá trị chưa thuế", "Thuế GTGT", "Giá trị chưa thuế", "Thuế GTGT", "Giá trị chưa thuế", "Thuế GTGT"]
                ]):
                    for i, text in enumerate(h_list, 1):
                        c = ws.cell(row=start_row + r_offset, column=i, value=text)
                        c.fill = fill_header
                        c.font = Font(bold=True)
                        c.alignment = align_center
                        c.border = border_thin

                ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 1, end_column=1)
                ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row + 1, end_column=2)
                ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row, end_column=4)
                ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row, end_column=6)
                ws.merge_cells(start_row=start_row, start_column=7, end_row=start_row, end_column=8)
                row_idx = start_row + 2
                tong_hd_truoc = tong_hd_thue = tong_tk_truoc = tong_tk_thue = 0

                for stt, r_data in enumerate(data, 1):
                    hd_truoc, hd_thue = r_data.get('hd_truoc', 0), r_data.get('hd_thue', 0)
                    tk_truoc, tk_thue = r_data.get('tk_truoc', 0), r_data.get('tk_thue', 0)
                    tong_hd_truoc += hd_truoc
                    tong_hd_thue += hd_thue
                    tong_tk_truoc += tk_truoc
                    tong_tk_thue += tk_thue

                    ws.cell(row=row_idx, column=1, value=stt).alignment = align_center
                    ws.cell(row=row_idx, column=2, value=r_data.get('ky', '')).alignment = align_center
                    self._ghi_so_excel(ws, row_idx, 3, hd_truoc if hd_truoc else "-")
                    self._ghi_so_excel(ws, row_idx, 4, hd_thue if hd_thue else "-")
                    self._ghi_so_excel(ws, row_idx, 5, tk_truoc if tk_truoc else "-")
                    self._ghi_so_excel(ws, row_idx, 6, tk_thue if tk_thue else "-")
                    self._ghi_so_excel(ws, row_idx, 7, (tk_truoc - hd_truoc) if (tk_truoc - hd_truoc) else "-")
                    self._ghi_so_excel(ws, row_idx, 8, (tk_thue - hd_thue) if (tk_thue - hd_thue) else "-")
                    row_idx += 1

                ws.cell(row=row_idx, column=1, value="Tổng").alignment = align_center
                ws.cell(row=row_idx, column=1).font = Font(bold=True)
                ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
                self._ghi_so_excel(ws, row_idx, 3, tong_hd_truoc if tong_hd_truoc else "-")
                self._ghi_so_excel(ws, row_idx, 4, tong_hd_thue if tong_hd_thue else "-")
                self._ghi_so_excel(ws, row_idx, 5, tong_tk_truoc if tong_tk_truoc else "-")
                self._ghi_so_excel(ws, row_idx, 6, tong_tk_thue if tong_tk_thue else "-")
                self._ghi_so_excel(ws, row_idx, 7, (tong_tk_truoc - tong_hd_truoc) if (tong_tk_truoc - tong_hd_truoc) else "-")
                self._ghi_so_excel(ws, row_idx, 8, (tong_tk_thue - tong_hd_thue) if (tong_tk_thue - tong_hd_thue) else "-")

                for c in range(1, 9):
                    ws.cell(row=row_idx, column=c).border = border_thin
                    ws.cell(row=row_idx, column=c).font = Font(bold=True)
                return row_idx + 3

            dong_ht = _ve_bang_bao_cao_excel(
                ws_bc, dong_ht,
                "1.1.2. Chênh lệch doanh số bán ra trên hồ sơ khai thuế và dữ liệu hóa đơn điện tử:", du_lieu_br
            )
            dong_ht = _ve_bang_bao_cao_excel(
                ws_bc, dong_ht,
                "1.1.3. Chênh lệch giá trị HHDV mua vào, thuế GTGT đầu vào được khấu trừ trên hồ sơ khai thuế và dữ liệu hóa đơn điện tử:", du_lieu_mv
            )

        # =========================================================
        # 2.5. XUẤT CHUỖI CUNG ỨNG, TỒN KHO CƠ BẢN & ĐỐI CHIẾU MẶT HÀNG AI
        # =========================================================
        if sc_result:
            ws_sc = wb.create_sheet("Chuỗi Cung Ứng")
            ws_sc.cell(row=1, column=1, value=f"MÔ HÌNH NHẬN DIỆN: {sc_result.get('mo_hinh', '')}").font = Font(bold=True, color="1E3A8A")
            r_sc = 2
            for cb in sc_result.get("canh_bao", []):
                c_cb = ws_sc.cell(row=r_sc, column=1, value=cb)
                c_cb.font = Font(bold=True, color="059669" if "✅" in cb else "DC2626")
                r_sc += 1
            r_sc += 1

            cols_sc = ["Phân Luồng", "Nhóm Ngành / Mặt Hàng", "Tổng Tiền Chưa Thuế", "Tỷ Trọng (%)"]
            fill_hdr = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
            for c_i, h_txt in enumerate(cols_sc, 1):
                c = ws_sc.cell(row=r_sc, column=c_i, value=h_txt)
                c.fill = fill_hdr
                c.font = Font(color="FFFFFF", bold=True)
                c.alignment = align_center
                c.border = border_thin
            r_sc += 1

            for luong_label, dict_tags, dict_details, tong_t in [
                ("ĐẦU VÀO (MUA VÀO)", sc_result.get("in_tags", {}), sc_result.get("in_details", {}), sc_result.get("in_total", 0)),
                ("ĐẦU RA (BÁN RA)", sc_result.get("out_tags", {}), sc_result.get("out_details", {}), sc_result.get("out_total", 0))
            ]:
                for tag, tien in sorted(dict_tags.items(), key=lambda x: x[1], reverse=True):
                    if tien == 0:
                        continue
                    ty_trong = (tien / tong_t * 100) if tong_t > 0 else 0.0
                    ws_sc.cell(row=r_sc, column=1, value=luong_label).border = border_thin
                    c_tag = ws_sc.cell(row=r_sc, column=2, value=tag)
                    c_tag.font = Font(bold=True)
                    c_tag.border = border_thin
                    self._ghi_so_excel(ws_sc, r_sc, 3, tien)
                    ws_sc.cell(row=r_sc, column=3).font = Font(bold=True)
                    c_tt = ws_sc.cell(row=r_sc, column=4, value=f"{ty_trong:.1f}%")
                    c_tt.alignment = align_center
                    c_tt.font = Font(bold=True)
                    c_tt.border = border_thin
                    r_sc += 1

                    gom_ct = {}
                    for it in dict_details.get(tag, []):
                        gom_ct[it['ten']] = gom_ct.get(it['ten'], 0) + it['tien']
                    for ten_h, tien_h in sorted(gom_ct.items(), key=lambda x: x[1], reverse=True):
                        ws_sc.cell(row=r_sc, column=1, value="").border = border_thin
                        ws_sc.cell(row=r_sc, column=2, value=f"   ↳ {ten_h}").border = border_thin
                        self._ghi_so_excel(ws_sc, r_sc, 3, tien_h)
                        ws_sc.cell(row=r_sc, column=4, value="").border = border_thin
                        r_sc += 1

        if kq_ton_kho:
            ws_tkho = wb.create_sheet("Rà soát Tồn kho")
            cols_tkho = [
                "Nhóm Phân Loại", "Mặt hàng Bán ra (Xuất)", "ĐVT Xuất", "SL Xuất",
                "Mặt hàng Mua vào (Nhập)", "ĐVT Nhập", "SL Nhập", "Chênh lệch Tồn", "Trạng thái", "Đánh giá / Ghi chú"
            ]
            self._dinh_dang_tieu_de(ws_tkho, cols_tkho)
            fill_do_tk = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            fill_vang_tk = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

            for r_i, item in enumerate(kq_ton_kho, start=2):
                ws_tkho.cell(row=r_i, column=1, value=item.get("nhom", "")).border = border_thin
                ws_tkho.cell(row=r_i, column=2, value=item.get("ban_ra", "")).border = border_thin
                ws_tkho.cell(row=r_i, column=3, value=item.get("dvt_ban", "")).border = border_thin
                self._ghi_so_excel(ws_tkho, r_i, 4, item.get("sl_ban", 0))
                ws_tkho.cell(row=r_i, column=5, value=item.get("mua_vao", "")).border = border_thin
                ws_tkho.cell(row=r_i, column=6, value=item.get("dvt_mua", "")).border = border_thin
                self._ghi_so_excel(ws_tkho, r_i, 7, item.get("sl_mua", 0))
                self._ghi_so_excel(ws_tkho, r_i, 8, item.get("ton_kho", 0))
                ws_tkho.cell(row=r_i, column=9, value=item.get("trang_thai", "")).border = border_thin
                ws_tkho.cell(row=r_i, column=10, value=item.get("thong_diep", "")).border = border_thin

                muc_do = item.get("muc_do", "")
                if muc_do == "red":
                    for c_i in range(1, 11):
                        ws_tkho.cell(row=r_i, column=c_i).fill = fill_do_tk
                elif muc_do == "orange":
                    for c_i in range(1, 11):
                        ws_tkho.cell(row=r_i, column=c_i).fill = fill_vang_tk

        if isinstance(df_ton_kho_ai, pd.DataFrame) and not df_ton_kho_ai.empty:
            ws_ai = wb.create_sheet("Đối chiếu Mặt hàng AI")
            cols_ai = list(df_ton_kho_ai.columns)
            self._dinh_dang_tieu_de(ws_ai, cols_ai)
            fill_do_ai = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            fill_vang_ai = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
            for r_i, row in enumerate(df_ton_kho_ai.itertuples(index=False), start=2):
                kq_ai = str(row[-1]).lower()
                for c_i, val in enumerate(row, start=1):
                    cell = ws_ai.cell(row=r_i, column=c_i, value=str(val))
                    cell.border = border_thin
                    cell.alignment = align_center if c_i >= 3 else align_wrap
                    if "rủi ro" in kq_ai:
                        cell.fill = fill_do_ai
                    elif "cảnh báo" in kq_ai:
                        cell.fill = fill_vang_ai

        # =========================================================
        # 3. VẼ TỜ KHAI (XML) MA TRẬN, PHỤ LỤC BẢNG KÊ, BCTC & CĐSPS TT200
        # =========================================================
        def _sort_ky(ky_str):
            try:
                t_q, y = ky_str.split('/')
                m = int(t_q.upper().replace('Q', '')) * 3 - 2 if 'Q' in t_q.upper() else int(t_q)
                return int(y) * 100 + m
            except Exception:
                return 999999

        fill_section = PatternFill(start_color="E2EBF5", end_color="E2EBF5", fill_type="solid")
        font_bold = Font(bold=True)

        for ma_form, dict_ky in dict_to_khai.items():
            danh_sach_ky = [ky for ky in sorted(list(dict_ky.keys()), key=_sort_ky) if str(nam) in str(ky)]
            if not danh_sach_ky:
                continue

            ws_tk = wb.create_sheet(f"TKhai {str(ma_form).replace('/', '_').replace(chr(92), '_')[:25]}")
            ma_tran, ds_bang_ke, ds_bctc_xml = {}, [], []

            for ky in danh_sach_ky:
                data = dict_ky[ky]
                for ma_ct, val_dict in data.get("ChiTieuChinh", {}).items():
                    ma_tran.setdefault("Chỉ tiêu chính", {}).setdefault(
                        ma_ct, {"DienGiai": val_dict.get("DienGiai", ""), "GiaTri": {}}
                    )["GiaTri"][ky] = val_dict.get("GiaTri", 0)
                for ph in ["B01_CDKT", "B02_KQKD", "B03_LCTT", "B04_CDTK"]:
                    if ph in data and data[ph]:
                        ds_bctc_xml.append({"Ky": ky, "PhanHe": ph, "DuLieu": data[ph]})
                for ten_pl, pl_data in data.get("PhuLuc", {}).items():
                    is_bang_ke = pl_data.get("IsBangKe", False)
                    du_lieu_pl = pl_data.get("DuLieu", {})
                    if not is_bang_ke:
                        for ma_ct, val_dict in du_lieu_pl.items():
                            ma_tran.setdefault(ten_pl, {}).setdefault(
                                ma_ct, {"DienGiai": val_dict.get("DienGiai", ""), "GiaTri": {}}
                            )["GiaTri"][ky] = val_dict.get("GiaTri", 0)
                    else:
                        ds_bang_ke.append({"Ky": ky, "TenPL": ten_pl, "DuLieu": du_lieu_pl})

            r_idx = 1
            if ma_tran:
                cols_header = ["Phân Hệ / Phụ Lục", "Mã CT", "Diễn Giải"] + danh_sach_ky + ["TỔNG / CHỐT SỐ"]
                self._dinh_dang_tieu_de(ws_tk, cols_header)
                r_idx += 1
                for phan_he, dict_ct in ma_tran.items():
                    ws_tk.cell(row=r_idx, column=1, value=phan_he).font = font_bold
                    for c in range(1, len(cols_header) + 1):
                        ws_tk.cell(row=r_idx, column=c).fill = fill_section
                        ws_tk.cell(row=r_idx, column=c).border = border_thin
                    r_idx += 1
                    for ma_ct, info in dict_ct.items():
                        ws_tk.cell(row=r_idx, column=1, value="").border = border_thin
                        ws_tk.cell(row=r_idx, column=2, value=ma_ct).border = border_thin
                        ws_tk.cell(row=r_idx, column=3, value=info["DienGiai"]).border = border_thin
                        mang_gia_tri_so = []
                        for c_idx, ky in enumerate(danh_sach_ky):
                            val = info["GiaTri"].get(ky, 0)
                            self._ghi_so_excel(ws_tk, r_idx, 4 + c_idx, val)
                            try:
                                mang_gia_tri_so.append(float(str(val).replace(",", "")))
                            except Exception:
                                mang_gia_tri_so.append(0.0)
                        tong_so = 0.0
                        if mang_gia_tri_so:
                            if str(ma_ct).strip().lower() == "ct22":
                                tong_so = mang_gia_tri_so[0]
                            elif str(ma_ct).strip().lower() == "ct43":
                                tong_so = mang_gia_tri_so[-1]
                            else:
                                tong_so = sum(mang_gia_tri_so)
                        self._ghi_so_excel(ws_tk, r_idx, 4 + len(danh_sach_ky), tong_so)
                        ws_tk.cell(row=r_idx, column=4 + len(danh_sach_ky)).font = font_bold
                        r_idx += 1

            if ds_bang_ke:
                r_idx += 2
                for bk in ds_bang_ke:
                    ws_tk.cell(row=r_idx, column=1, value=f"--- {bk['TenPL']} (Kỳ: {bk['Ky']}) ---").font = Font(bold=True, color="0052D4")
                    r_idx += 1
                    du_lieu_dong = bk['DuLieu']
                    if not du_lieu_dong:
                        continue
                    dong_dau = list(du_lieu_dong.values())[0]
                    danh_sach_den = ["Mã CT", "Chỉ tiêu ct09a_ma", "Chỉ tiêu ct12_ma", "Chỉ tiêu ct14_ma", "Chỉ tiêu ct09_ma"]
                    cols_bk = [c for c in list(dong_dau.keys()) if str(c).strip() not in danh_sach_den]
                    for c_i, c_name in enumerate(cols_bk, 1):
                        c = ws_tk.cell(row=r_idx, column=c_i, value=c_name)
                        c.fill = fill_section
                        c.font = font_bold
                        c.border = border_thin
                    r_idx += 1
                    for row_data in du_lieu_dong.values():
                        for c_i, c_name in enumerate(cols_bk, 1):
                            self._ghi_du_lieu_thong_minh(ws_tk, r_idx, c_i, c_name, row_data.get(c_name, ""))
                        r_idx += 1
                    r_idx += 1

            if ds_bctc_xml:
                r_idx += 2
                for bctc in ds_bctc_xml:
                    ws_tk.cell(row=r_idx, column=1, value=f"--- BÁO CÁO TÀI CHÍNH: {bctc['PhanHe']} (Kỳ: {bctc['Ky']}) ---").font = Font(bold=True, color="D83B01")
                    r_idx += 1
                    du_lieu_bctc = bctc['DuLieu']
                    if not du_lieu_bctc:
                        continue
                    cot_gia_tri = [k for k in du_lieu_bctc[list(du_lieu_bctc.keys())[0]].keys() if k != "DienGiai"]
                    cols_bctc = ["Mã Chỉ Tiêu", "Nội Dung Diễn Giải"] + cot_gia_tri
                    fill_hdr = PatternFill(start_color="0052D4", end_color="0052D4", fill_type="solid")
                    for c_i, c_name in enumerate(cols_bctc, 1):
                        c = ws_tk.cell(row=r_idx, column=c_i, value=c_name)
                        c.fill = fill_hdr
                        c.font = Font(color="FFFFFF", bold=True)
                        c.alignment = align_center
                        c.border = border_thin
                    r_idx += 1
                    for ma_ct, row_data in du_lieu_bctc.items():
                        ws_tk.cell(row=r_idx, column=1, value=ma_ct).border = border_thin
                        ws_tk.cell(row=r_idx, column=2, value=row_data.get("DienGiai", "")).border = border_thin
                        for i, ten_cot in enumerate(cot_gia_tri):
                            self._ghi_so_excel(ws_tk, r_idx, 3 + i, row_data.get(ten_cot, "0"))
                        r_idx += 1
                    r_idx += 1

        # Xuất Bảng CĐSPS TT200 nếu có
        if isinstance(du_lieu_bcdsps_tt200, dict) and str(nam) in du_lieu_bcdsps_tt200:
            df_cd = du_lieu_bcdsps_tt200[str(nam)]
            if isinstance(df_cd, pd.DataFrame) and not df_cd.empty:
                ws_cd = wb.create_sheet("CĐSPS (TT200)")
                cot_cd = ["MaTK", "TenTK", "DuNoDauKy", "DuCoDauKy", "PhatSinhNo", "PhatSinhCo", "DuNoCuoiKy", "DuCoCuoiKy"]
                cot_cd = [c for c in cot_cd if c in df_cd.columns]
                self._dinh_dang_tieu_de(ws_cd, cot_cd)
                for r_i, row in enumerate(df_cd[cot_cd].itertuples(index=False), start=2):
                    for c_i, val in enumerate(row, start=1):
                        if c_i <= 2:
                            self._ghi_du_lieu_thong_minh(ws_cd, r_i, c_i, cot_cd[c_i - 1], val)
                        else:
                            self._ghi_so_excel(ws_cd, r_i, c_i, val)

        # =========================================================
        # 4. VẼ HÓA ĐƠN BÁN RA, MUA VÀO, RỦI RO, TT ĐẶC BIỆT & NNT BẤT THƯỜNG
        # =========================================================
        def _xuat_bang_hoa_don_chuan(ten_sheet, df_hd, loc_theo_nam=True):
            if not isinstance(df_hd, pd.DataFrame) or df_hd.empty:
                return
            df_xuat = df_hd.copy()
            if loc_theo_nam and '_Nam' in df_xuat.columns:
                mask_nam = (
                    (df_xuat['_Nam'].astype(str) == str(nam))
                    | (df_xuat['_Nam'].astype(str).isin(['', 'nan', 'NaT']))
                )
                df_xuat = df_xuat[mask_nam].copy()
            if df_xuat.empty:
                return

            ws_hd = wb.create_sheet(ten_sheet)
            cols = [col for col in df_xuat.columns if not str(col).startswith("_") and str(col) != "Row_Color"]
            self._dinh_dang_tieu_de(ws_hd, cols)

            tu_khoa_so = [
                "tiền", "thuế", "đơn giá", "số lượng", "giá trị", "tỷ giá",
                "tien", "thue", "dongia", "soluong", "giatri", "tygia"
            ]
            tu_khoa_text_bat_buoc = [
                "mst", "mãsốthuế", "masothue", "cccd", "cmnd",
                "sốhđ", "sốhóađơn", "sohoadon", "kýhiệu", "kyhieu", "tàikhoản", "mãtt"
            ]
            fill_do = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            fill_vang = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
            tu_khoa_tms_do = ["không hoạt động tại địa chỉ", "bỏ địa chỉ", "chờ xác minh", "mất tích"]
            tu_khoa_tms_vang = ["ngừng hoạt động", "tạm nghỉ", "tạm ngừng", "chấm dứt"]
            tu_khoa_hd_do = ["hủy", "xóa bỏ", "bị xóa bỏ", "bị hủy", "bị thay thế"]
            tu_khoa_hd_vang = ["thay thế", "điều chỉnh", "bị điều chỉnh", "chiết khấu", "trả lại"]

            idx_tms = cols.index('TrangThaiTMS') if 'TrangThaiTMS' in cols else (
                cols.index('Trạng thái hoạt động') if 'Trạng thái hoạt động' in cols else -1
            )
            idx_rr = cols.index('Đối chiếu HĐ rủi ro') if 'Đối chiếu HĐ rủi ro' in cols else -1
            cot_tt_indices = [
                cols.index(c) for c in [
                    'TrangThaiHoaDon', 'Trạng thái hóa đơn', 'Tinhchat', 'Tính chất',
                    'PhanLoaiTrangThai', 'Nhóm Trạng Thái Đặc Biệt'
                ] if c in cols
            ]

            for r_idx, row in enumerate(df_xuat[cols].itertuples(index=False), start=2):
                tt_tms = str(row[idx_tms]).lower() if idx_tms != -1 and pd.notna(row[idx_tms]) else ""
                tt_rr = str(row[idx_rr]).lower() if idx_rr != -1 and pd.notna(row[idx_rr]) else ""
                tt_hd_goc = " | ".join(str(row[i]).lower() for i in cot_tt_indices if pd.notna(row[i])) if cot_tt_indices else ""

                is_do = (
                    ten_sheet in ["HĐ Rủi Ro", "NNT Bất Thường (TMS)"]
                    or "thuộc ds hóa đơn rủi ro" in tt_rr
                    or any(k in tt_tms for k in tu_khoa_tms_do)
                    or any(k in tt_hd_goc for k in tu_khoa_hd_do)
                )
                is_vang = (not is_do) and (
                    any(k in tt_tms for k in tu_khoa_tms_vang)
                    or any(k in tt_hd_goc for k in tu_khoa_hd_vang)
                )

                for c_idx, val in enumerate(row, start=1):
                    col_name = cols[c_idx - 1]
                    col_lower = str(col_name).lower().replace(" ", "").replace("_", "")
                    is_id_col = any(k in col_lower for k in tu_khoa_text_bat_buoc)
                    is_num_col = (not is_id_col) and any(k in col_lower for k in tu_khoa_so)

                    if is_num_col:
                        self._ghi_so_excel(ws_hd, r_idx, c_idx, val)
                    else:
                        self._ghi_du_lieu_thong_minh(ws_hd, r_idx, c_idx, col_name, val)

                    if is_do:
                        ws_hd.cell(row=r_idx, column=c_idx).fill = fill_do
                    elif is_vang:
                        ws_hd.cell(row=r_idx, column=c_idx).fill = fill_vang

        _xuat_bang_hoa_don_chuan("HĐ Bán Ra", df_hdbr, loc_theo_nam=True)
        _xuat_bang_hoa_don_chuan("HĐ Mua Vào", df_hdmv, loc_theo_nam=True)
        _xuat_bang_hoa_don_chuan("HĐ Rủi Ro", df_hd_rui_ro, loc_theo_nam=True)
        _xuat_bang_hoa_don_chuan("HĐ TT Đặc Biệt", df_hd_dac_biet, loc_theo_nam=True)
        _xuat_bang_hoa_don_chuan("NNT Bất Thường (TMS)", df_nnt_bat_thuong, loc_theo_nam=False)

        # Chốt lưu file
        if not wb.sheetnames:
            wb.create_sheet("Không có dữ liệu")
        for sheet_name in wb.sheetnames:
            self._can_chinh_kich_thuoc_cot(wb[sheet_name])
        wb.save(duong_dan_luu)


# =========================================================================
# GIAO DIỆN CẤU HÌNH AUTO BOT (POPUP AGENT)
# =========================================================================
class DialogAutoBot(QDialog):
    def __init__(self, main_app, parent=None):
        super().__init__(parent)
        self.main_app = main_app
        self.setWindowTitle("🤖 Agent Phân tích & Kết xuất Tự động Toàn diện")
        self.setFixedSize(860, 680)
        self.setWindowFlag(Qt.WindowType.WindowContextHelpButtonHint, False)

        self.setStyleSheet("""
            QDialog { background-color: #f8fafc; }
            QGroupBox { font-weight: bold; color: #1e3a8a; border: 1px solid #cbd5e1; border-radius: 6px; margin-top: 10px; }
            QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 5px 0 5px; }
            QLineEdit { padding: 8px; border: 1px solid #cbd5e1; border-radius: 4px; background: #ffffff; }
            QPushButton { font-weight: bold; border-radius: 5px; padding: 8px; }
        """)

        self._thiet_lap_ui()

    def _thiet_lap_ui(self):
        layout_chinh = QVBoxLayout(self)

        # ---------------------------------------------------------
        # KHỐI 1: CHỌN ĐƯỜNG DẪN & QUÉT TỰ ĐỘNG
        # ---------------------------------------------------------
        grp_duong_dan = QGroupBox("1. Thiết lập Thư mục")
        layout_dd = QVBoxLayout(grp_duong_dan)

        khung_chon_goc = QHBoxLayout()
        self.txt_thu_muc_goc = QLineEdit()
        self.txt_thu_muc_goc.setReadOnly(True)
        self.txt_thu_muc_goc.setPlaceholderText("Chọn thư mục chứa các MST hoặc thư mục chứa trực tiếp file XML/Hóa đơn...")
        btn_chon_goc = QPushButton(" Chọn Thư mục Nguồn")
        btn_chon_goc.setIcon(qta.icon('fa5s.folder-open', color='#3b82f6'))
        btn_chon_goc.setStyleSheet("background-color: #e0f2fe; color: #0284c7; border: 1px solid #7dd3fc;")
        btn_chon_goc.clicked.connect(self.chon_thu_muc_goc)
        khung_chon_goc.addWidget(self.txt_thu_muc_goc)
        khung_chon_goc.addWidget(btn_chon_goc)

        khung_chon_xuat = QHBoxLayout()
        self.txt_thu_muc_xuat = QLineEdit()
        self.txt_thu_muc_xuat.setReadOnly(True)
        self.txt_thu_muc_xuat.setPlaceholderText("Chưa chọn thư mục lưu Báo cáo Excel...")
        btn_chon_xuat = QPushButton(" Chọn Thư mục Đích")
        btn_chon_xuat.setIcon(qta.icon('fa5s.save', color='#10b981'))
        btn_chon_xuat.setStyleSheet("background-color: #d1fae5; color: #059669; border: 1px solid #6ee7b7;")
        btn_chon_xuat.clicked.connect(self.chon_thu_muc_xuat)
        khung_chon_xuat.addWidget(self.txt_thu_muc_xuat)
        khung_chon_xuat.addWidget(btn_chon_xuat)

        layout_dd.addLayout(khung_chon_goc)
        layout_dd.addLayout(khung_chon_xuat)
        layout_chinh.addWidget(grp_duong_dan)

        # ---------------------------------------------------------
        # KHỐI 2: THÂN APP (LIST MST VÀ CẤU HÌNH DÙNG CHUNG)
        # ---------------------------------------------------------
        khung_giua = QHBoxLayout()

        # Cột trái: Danh sách MST quét được
        grp_mst = QGroupBox("2. Danh sách Hồ sơ Doanh nghiệp")
        layout_mst = QVBoxLayout(grp_mst)

        self.chk_chon_tat_ca = QCheckBox("Chọn tất cả")
        self.chk_chon_tat_ca.setChecked(True)
        self.chk_chon_tat_ca.toggled.connect(self.xu_ly_chon_tat_ca_mst)
        layout_mst.addWidget(self.chk_chon_tat_ca)

        self.list_mst = QListWidget()
        self.list_mst.setStyleSheet("background-color: #ffffff; border: 1px solid #cbd5e1; border-radius: 4px;")
        layout_mst.addWidget(self.list_mst)
        khung_giua.addWidget(grp_mst, 5)

        # Cột phải: Cấu hình dùng chung (Đồng bộ với màn hình chính)
        grp_cau_hinh = QGroupBox("3. Cấu hình Agent Phân tích")
        grp_cau_hinh.setStyleSheet("""
            QGroupBox { border: none; margin-top: 5px; }
            QGroupBox::title { color: #1e3a8a; font-size: 15px; font-weight: bold; }
            QLabel { color: #334155; margin-top: 5px; }
            QComboBox { padding: 6px; border: 1px solid #cbd5e1; border-radius: 4px; background: white; min-height: 22px; }
            QComboBox:focus { border: 1px solid #3b82f6; }
            QCheckBox { color: #475569; font-size: 13px; }
        """)
        layout_ch = QVBoxLayout(grp_cau_hinh)
        layout_ch.setSpacing(8)

        # 1. Loại Báo cáo tài chính
        layout_ch.addWidget(QLabel("Loại Báo cáo tài chính:"))
        self.cb_bctc = QComboBox()
        self.cb_bctc.addItems(["Tự động nhận diện", "Theo Thông tư 200", "Theo Thông tư 133"])
        if hasattr(self.main_app, 'giao_dien') and hasattr(self.main_app.giao_dien, 'cb_bctc'):
            self.cb_bctc.setCurrentIndex(self.main_app.giao_dien.cb_bctc.currentIndex())
        layout_ch.addWidget(self.cb_bctc)

        # 2. Năm phân tích
        layout_ch.addWidget(QLabel("Năm phân tích:"))
        self.cb_nam = QComboBox()
        self.cb_nam.addItems(["Tất cả các năm", "2026", "2025", "2024", "2023", "2022"])
        if hasattr(self.main_app, 'giao_dien') and hasattr(self.main_app.giao_dien, 'cb_nam'):
            self.cb_nam.setCurrentIndex(self.main_app.giao_dien.cb_nam.currentIndex())
        layout_ch.addWidget(self.cb_nam)

        # 3. Cụm TMS & Đồng bộ UI
        khung_tms = QHBoxLayout()
        khung_tms.setContentsMargins(0, 8, 0, 4)

        khung_checkbox_tms = QVBoxLayout()
        khung_checkbox_tms.setSpacing(6)
        self.chk_tms = QCheckBox("Tra cứu trạng thái HĐ doanh nghiệp từ TMS")
        self.chk_dong_bo_ngay = QCheckBox("Tra cứu ngày nộp Tờ khai từ TMS")
        self.chk_sync_ui = QCheckBox("Đổ kết quả lên Giao diện chính sau khi chạy xong")
        self.chk_sync_ui.setChecked(True)
        self.chk_sync_ui.setStyleSheet("color: #059669; font-weight: bold;")

        if hasattr(self.main_app, 'giao_dien'):
            if hasattr(self.main_app.giao_dien, 'chk_tms'):
                self.chk_tms.setChecked(self.main_app.giao_dien.chk_tms.isChecked())
            if hasattr(self.main_app.giao_dien, 'chk_dong_bo_ngay'):
                self.chk_dong_bo_ngay.setChecked(self.main_app.giao_dien.chk_dong_bo_ngay.isChecked())

        khung_checkbox_tms.addWidget(self.chk_tms)
        khung_checkbox_tms.addWidget(self.chk_dong_bo_ngay)
        khung_checkbox_tms.addWidget(self.chk_sync_ui)

        btn_cau_hinh_tms = QPushButton(" Cấu hình TMS")
        btn_cau_hinh_tms.setIcon(qta.icon('fa5s.cog', color='#475569'))
        btn_cau_hinh_tms.setStyleSheet("""
            QPushButton { background-color: #e2e8f0; color: #334155; border: 1px solid #cbd5e1; border-radius: 4px; padding: 6px 12px; font-weight: bold; }
            QPushButton:hover { background-color: #cbd5e1; }
        """)
        btn_cau_hinh_tms.clicked.connect(self.main_app.mo_cau_hinh_tms)

        khung_tms.addLayout(khung_checkbox_tms)
        khung_tms.addStretch()
        khung_tms.addWidget(btn_cau_hinh_tms, alignment=Qt.AlignmentFlag.AlignTop)
        layout_ch.addLayout(khung_tms)

        # Nút Chọn Công thức (Đồng bộ với bộ lọc Rule thủ công)
        self.btn_chon_rule_auto = QPushButton(" Bộ lọc Công thức (Đang chạy Tất cả)")
        self.btn_chon_rule_auto.setIcon(qta.icon('fa5s.filter', color='#1e3a8a'))
        self.btn_chon_rule_auto.setStyleSheet(
            "QPushButton { background-color: #dbeafe; color: #1e3a8a; border: 1px solid #93c5fd; border-radius: 4px; padding: 7px; }"
        )
        if getattr(self.main_app, 'danh_sach_rule_duoc_chon', None) is not None:
            self.btn_chon_rule_auto.setText(f" Đã chọn ({len(self.main_app.danh_sach_rule_duoc_chon)} Rule)")
        self.btn_chon_rule_auto.clicked.connect(self._chon_rule_tu_bot)
        layout_ch.addWidget(self.btn_chon_rule_auto)

        layout_ch.addStretch()
        khung_giua.addWidget(grp_cau_hinh, 4)
        layout_chinh.addLayout(khung_giua)

        # ---------------------------------------------------------
        # KHỐI 3: TIẾN ĐỘ & LOG XỬ LÝ
        # ---------------------------------------------------------
        grp_log = QGroupBox("4. Tiến độ xử lý")
        layout_log = QVBoxLayout(grp_log)

        self.lbl_trang_thai = QLabel("Sẵn sàng...")
        self.lbl_trang_thai.setStyleSheet("font-weight: bold; color: #0284c7;")
        layout_log.addWidget(self.lbl_trang_thai)

        self.prg_tong = QProgressBar()
        self.prg_tong.setStyleSheet(
            "QProgressBar { border: 1px solid #cbd5e1; border-radius: 4px; text-align: center; } QProgressBar::chunk { background-color: #10b981; }"
        )
        layout_log.addWidget(self.prg_tong)

        self.txt_log = QTextEdit()
        self.txt_log.setReadOnly(True)
        self.txt_log.setStyleSheet("""
            QTextEdit {
                background-color: #ffffff; 
                color: #334155; 
                font-family: Consolas, 'Courier New', monospace; 
                font-size: 13px; 
                border: 1px solid #cbd5e1;
                border-radius: 4px;
                padding: 5px;
            }
        """)
        layout_log.addWidget(self.txt_log)
        layout_chinh.addWidget(grp_log)

        # ---------------------------------------------------------
        # KHỐI NÚT ĐIỀU KHIỂN
        # ---------------------------------------------------------
        khung_dieu_khien = QHBoxLayout()
        self.btn_dong = QPushButton(" Đóng cửa sổ")
        self.btn_dong.setIcon(qta.icon('fa5s.times', color='#64748b'))
        self.btn_dong.setStyleSheet("background-color: #f1f5f9; border: 1px solid #cbd5e1; color: #475569;")
        self.btn_dong.clicked.connect(self.reject)

        self.btn_chay_auto = QPushButton(" 🚀 CHẠY AGENT PHÂN TÍCH & KẾT XUẤT TOÀN DIỆN ")
        self.btn_chay_auto.setIcon(qta.icon('fa5s.robot', color='white'))
        self.btn_chay_auto.setStyleSheet("background-color: #2563eb; color: white; padding: 12px; font-size: 14px;")
        self.btn_chay_auto.clicked.connect(self.thuc_thi_auto)

        khung_dieu_khien.addWidget(self.btn_dong)
        khung_dieu_khien.addWidget(self.btn_chay_auto)
        layout_chinh.addLayout(khung_dieu_khien)

    def _chon_rule_tu_bot(self):
        if hasattr(self.main_app, 'mo_dialog_chon_rule'):
            self.main_app.mo_dialog_chon_rule()
            if getattr(self.main_app, 'danh_sach_rule_duoc_chon', None) is not None:
                self.btn_chon_rule_auto.setText(f" Đã chọn ({len(self.main_app.danh_sach_rule_duoc_chon)} Rule)")

    def chon_thu_muc_goc(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Chọn Thư mục chứa các MST hoặc Hồ sơ thuế")
        if folder_path:
            self.txt_thu_muc_goc.setText(folder_path)
            if not self.txt_thu_muc_xuat.text().strip():
                self.txt_thu_muc_xuat.setText(folder_path)
            self.quet_thu_muc_mst(folder_path)

    def chon_thu_muc_xuat(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Chọn Thư mục xuất Báo cáo")
        if folder_path:
            self.txt_thu_muc_xuat.setText(folder_path)

    def quet_thu_muc_mst(self, folder_path):
        """
        Quét thông minh:
        1. Tìm các thư mục con có chứa TK/HDMV/HDBR/CDSPS hoặc chứa trực tiếp file .xml/.xlsx/.xls/.csv.
        2. Nếu người dùng chọn thẳng thư mục của 1 công ty (có chứa TK/HDMV/HDBR hoặc file .xml/.xlsx), nhận diện luôn thư mục đó!
        """
        self.list_mst.clear()
        so_luong = 0
        try:
            subdirs_root = [
                d.upper() for d in os.listdir(folder_path)
                if os.path.isdir(os.path.join(folder_path, d))
            ]
            files_root = [
                f.lower() for f in os.listdir(folder_path)
                if os.path.isfile(os.path.join(folder_path, f)) and f.lower().endswith(('.xml', '.xlsx', '.xls', '.csv'))
            ]

            # Kiểm tra nếu bản thân folder_path đã là thư mục của 1 công ty
            if any(k in subdirs_root for k in ["TK", "HDMV", "HDBR", "CDSPS", "TT200"]) or files_root:
                ten_don_vi = os.path.basename(os.path.normpath(folder_path))
                item = QListWidgetItem(ten_don_vi)
                item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                item.setCheckState(Qt.CheckState.Checked)
                self.list_mst.addItem(item)
                so_luong = 1
                self.ghi_log(f"🔍 Đã nhận diện trực tiếp hồ sơ đơn vị: {ten_don_vi}")
                return

            # Ngược lại: Quét danh sách các thư mục công ty con bên trong
            for ten_thu_muc in sorted(os.listdir(folder_path)):
                full_path = os.path.join(folder_path, ten_thu_muc)
                if os.path.isdir(full_path):
                    subdirs = [d.upper() for d in os.listdir(full_path) if os.path.isdir(os.path.join(full_path, d))]
                    subfiles = [
                        f.lower() for f in os.listdir(full_path)
                        if os.path.isfile(os.path.join(full_path, f)) and f.lower().endswith(('.xml', '.xlsx', '.xls', '.csv'))
                    ]
                    if any(k in subdirs for k in ["TK", "HDMV", "HDBR", "CDSPS", "TT200"]) or subfiles:
                        item = QListWidgetItem(ten_thu_muc)
                        item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                        item.setCheckState(Qt.CheckState.Checked)
                        self.list_mst.addItem(item)
                        so_luong += 1

            self.ghi_log(f"🔍 Quét thành công: Phát hiện {so_luong} hồ sơ doanh nghiệp.")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi quét thư mục", str(e))

    def xu_ly_chon_tat_ca_mst(self, checked):
        for i in range(self.list_mst.count()):
            item = self.list_mst.item(i)
            item.setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)

    def ghi_log(self, text):
        self.txt_log.append(text)
        self.txt_log.verticalScrollBar().setValue(self.txt_log.verticalScrollBar().maximum())

    def _dong_bo_ket_qua_len_main_ui(self, state: dict):
        """
        Đổ toàn bộ dữ liệu do Agent phân tích được lên Giao diện chính (giống hệt như người dùng thao tác thủ công).
        """
        try:
            app = self.main_app
            ui = app.giao_dien

            app.du_lieu_he_thong_to_khai = state.get("du_lieu_to_khai", {})
            app.toan_bo_file_xml = set(state.get("ds_file_xml", []))
            app.du_lieu_bcdsps_tt200 = state.get("du_lieu_bcdsps_tt200", {})
            app.dict_hd_ban_ra = state.get("df_hdbr", pd.DataFrame())
            app.dict_hd_mua_vao = state.get("df_hdmv", pd.DataFrame())
            app.kho_ket_qua_nam = state.get("kho_ket_qua_nam", {})
            danh_sach_nam = state.get("danh_sach_nam", [])

            # 1. Vẽ lại các Tab Tờ khai & CĐSPS TT200
            while ui.tab_tong_hop_chi_tiet.count() > 0:
                ui.tab_tong_hop_chi_tiet.removeTab(0)
            app.dict_khbs_widgets.clear()

            for dict_ky in app.du_lieu_he_thong_to_khai.values():
                for ket_qua in dict_ky.values():
                    app.cap_nhat_bang_tong_hop(ket_qua)

            if isinstance(app.du_lieu_bcdsps_tt200, dict) and app.du_lieu_bcdsps_tt200:
                nam_cd_dau = sorted(list(app.du_lieu_bcdsps_tt200.keys()), reverse=True)[0]
                app._ve_bang_cdsps_len_ui(app.du_lieu_bcdsps_tt200[nam_cd_dau], nam=nam_cd_dau)

            # 2. Vẽ lại các Bảng Hóa đơn Mua vào, Bán ra, Rủi ro, Trạng thái đặc biệt
            from module_giao_dien.quan_ly_bang_bieu import QuanLyBangBieu
            if isinstance(app.dict_hd_ban_ra, pd.DataFrame) and not app.dict_hd_ban_ra.empty:
                QuanLyBangBieu.xu_ly_bang_ke(ui.bang_hd_ban_ra, app.dict_hd_ban_ra)
                QuanLyBangBieu.cau_hinh_bang_du_lieu(ui.bang_hd_ban_ra)
                QuanLyBangBieu.can_bang_kich_thuoc_bang(ui.bang_hd_ban_ra)

            if isinstance(app.dict_hd_mua_vao, pd.DataFrame) and not app.dict_hd_mua_vao.empty:
                QuanLyBangBieu.xu_ly_bang_ke(ui.bang_hd_mua_vao, app.dict_hd_mua_vao)
                QuanLyBangBieu.cau_hinh_bang_du_lieu(ui.bang_hd_mua_vao)
                QuanLyBangBieu.can_bang_kich_thuoc_bang(ui.bang_hd_mua_vao)

            if hasattr(app, 'cap_nhat_kpi_chuan_bi'):
                app.cap_nhat_kpi_chuan_bi()

            # 3. Đổ kết quả Tồn kho cơ bản lên Tab Hóa đơn
            kq_ton_kho = state.get("kq_ton_kho", [])
            if kq_ton_kho and hasattr(ui, 'tabs_hoa_don'):
                from module_giao_dien.builder_phan_tich_ton_kho import BuilderPhanTichTonKho
                tab_idx_old = -1
                for i in range(ui.tabs_hoa_don.count()):
                    if ui.tabs_hoa_don.tabText(i) == "Rà soát Tồn kho":
                        tab_idx_old = i
                        break
                if tab_idx_old != -1:
                    ui.tabs_hoa_don.widget(tab_idx_old).nap_du_lieu(kq_ton_kho)
                else:
                    ui.tabs_hoa_don.addTab(
                        BuilderPhanTichTonKho(kq_ton_kho),
                        qta.icon('fa5s.boxes', color='#059669'),
                        "Rà soát Tồn kho"
                    )

            # 4. Cập nhật Combobox Năm kết quả và hiển thị màn hình Báo cáo Phân tích
            if danh_sach_nam and hasattr(ui, 'cb_nam_kq'):
                ui.cb_nam_kq.blockSignals(True)
                ui.cb_nam_kq.clear()
                ui.cb_nam_kq.addItems(danh_sach_nam)
                ui.cb_nam_kq.blockSignals(False)
                app.chuyen_view_ket_qua_nam(danh_sach_nam[0])
                app.chuyen_tab(4)

            self.ghi_log("✨ Đã đồng bộ toàn bộ dữ liệu và bảng biểu lên Giao diện chính!")
        except Exception as e:
            self.ghi_log(f"⚠️ Lỗi đồng bộ lên Giao diện chính: {e}")

    def thuc_thi_auto(self):
        tm_goc = self.txt_thu_muc_goc.text().strip()
        tm_xuat = self.txt_thu_muc_xuat.text().strip()

        if not tm_goc or not tm_xuat:
            QMessageBox.warning(self, "Lỗi", "Vui lòng chọn đầy đủ Thư mục Nguồn và Thư mục Đích!")
            return

        ds_mst_chay = []
        for i in range(self.list_mst.count()):
            item = self.list_mst.item(i)
            if item.checkState() == Qt.CheckState.Checked:
                ds_mst_chay.append(item.text())

        if not ds_mst_chay:
            QMessageBox.warning(self, "Lỗi", "Bạn chưa tích chọn Hồ sơ doanh nghiệp nào để chạy!")
            return

        self.btn_chay_auto.setEnabled(False)
        self.btn_dong.setEnabled(False)
        self.prg_tong.setMaximum(len(ds_mst_chay))
        self.prg_tong.setValue(0)
        self.txt_log.clear()

        config_tms = {}
        try:
            from module_ho_tro.module_cau_hinh_tms import TMSConfig
            config_tms = TMSConfig.load()
        except Exception:
            pass

        cau_hinh = {
            "nam_pt": self.cb_nam.currentText(),
            "loai_bctc": self.cb_bctc.currentText(),
            "check_tms": self.chk_tms.isChecked(),
            "dong_bo_ngay": self.chk_dong_bo_ngay.isChecked(),
            "sync_ui": self.chk_sync_ui.isChecked(),
            "config_tms": config_tms,
            "danh_sach_rule_chon": getattr(self.main_app, 'danh_sach_rule_duoc_chon', None),
            "allowed_cqt": getattr(self.main_app, 'allowed_cqt', ["*"])
        }

        self.worker = AutoBotWorker(tm_goc, tm_xuat, ds_mst_chay, cau_hinh)
        self.worker.log_msg.connect(self.ghi_log)
        self.worker.progress_tong.connect(self.cap_nhat_tien_do)
        self.worker.progress_ct.connect(lambda txt: self.lbl_trang_thai.setText(txt))
        self.worker.sync_main_ui.connect(self._dong_bo_ket_qua_len_main_ui)
        self.worker.finished_batch.connect(self.hoan_thanh_batch)
        self.worker.start()

    def cap_nhat_tien_do(self, hien_tai, tong):
        self.prg_tong.setValue(hien_tai)

    def hoan_thanh_batch(self, thanh_cong, that_bai):
        self.lbl_trang_thai.setText("HOÀN TẤT!")
        self.btn_chay_auto.setEnabled(True)
        self.btn_dong.setEnabled(True)
        QMessageBox.information(
            self, "Hoàn thành",
            f"Agent đã phân tích và kết xuất hoàn tất!\n- Thành công: {thanh_cong} hồ sơ\n- Thất bại: {that_bai} hồ sơ"
        )