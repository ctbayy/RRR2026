import os
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QSplitter,
                               QTableWidget, QTableWidgetItem, QHeaderView,
                               QPushButton, QLabel, QCheckBox, QMessageBox, QFileDialog,
                               QGroupBox, QFormLayout, QLineEdit)
from PySide6.QtCore import Qt
import qtawesome as qta
from module_ho_tro.dinh_dang_so import dinh_dang_so_vn
import json  # <--- THÊM THƯ VIỆN NÀY ĐỂ ĐỌC GHI CẤU HÌNH

class GiaoDienXuatPhieuNhanXet(QWidget):
    def __init__(self, ds_rui_ro, du_lieu_tong_hop=None, parent=None):
        super().__init__(parent)
        # 1. Lọc danh sách rủi ro (Chỉ lấy lỗi hoặc có chênh lệch)
        self.ds_rui_ro = [r for r in ds_rui_ro if r.get('is_rui_ro', False) or r.get('ket_qua_so', 0) != 0]

        # 2. Hứng cục data khổng lồ (Tờ khai GTGT 12 tháng, TNDN, TNCN, Hóa đơn) từ hệ thống truyền vào
        self.du_lieu_tong_hop = du_lieu_tong_hop if du_lieu_tong_hop else {}

        self.thiet_lap_ui()
        self.nap_danh_sach_rui_ro()

    def _doc_cau_hinh_cqt(self):
        """Đọc thông tin Cơ quan thuế đã lưu từ phiên trước"""
        self.file_config_cqt = "Data/config_cqt.json"
        if os.path.exists(self.file_config_cqt):
            try:
                with open(self.file_config_cqt, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except:
                pass
        return {}

    def _luu_cau_hinh_cqt(self):
        """Lưu thông tin hiện tại trên Form vào bộ nhớ"""
        data = {
            "cqt": self.txt_cqt.text(),
            "phong_doi": self.txt_phong_doi.text(),
            "to_truong": self.txt_to_truong.text(),
            "thanh_vien": self.txt_thanh_vien.text(),
            "lanh_dao": self.txt_lanh_dao.text(),
            "chuc_vu": self.txt_chuc_vu.text()
        }
        try:
            os.makedirs("Data", exist_ok=True)
            with open(self.file_config_cqt, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False)
        except Exception as e:
            pass
    def thiet_lap_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(15, 15, 15, 15)

        bo_chia = QSplitter(Qt.Orientation.Horizontal)
        bo_chia.setStyleSheet("QSplitter::handle { background-color: #cbd5e1; width: 4px; }")

        # ==========================================
        # KHUNG TRÁI: CHỈ TICK CHỌN RỦI RO (KHÔNG SỬA TEXT)
        # ==========================================
        khung_trai = QWidget()
        layout_trai = QVBoxLayout(khung_trai)
        layout_trai.setContentsMargins(0, 0, 10, 0)

        lbl_trai = QLabel("1. Lọc Nội dung vi phạm/chênh lệch đưa vào Báo cáo:")
        lbl_trai.setStyleSheet("font-weight: bold; color: #1e3a8a; font-size: 14px;")
        layout_trai.addWidget(lbl_trai)

        self.bang_chon = QTableWidget(0, 3)
        # --- THÊM ĐOẠN CSS NÀY VÀO ĐỂ BẢNG NỔI BẬT KHỎI NỀN ---
        self.bang_chon.setStyleSheet("""
                    QTableWidget {
                        background-color: #ffffff;
                        alternate-background-color: #f8fafc;
                        gridline-color: #cbd5e1;
                        border: 1px solid #94a3b8;
                        color: #1e293b;
                    }
                    QHeaderView::section {
                        background-color: #f1f5f9;
                        font-weight: bold;
                        border: 1px solid #cbd5e1;
                        color: #0f172a;
                    }
                """)
        # -----------------------------------------------------
        self.bang_chon.setHorizontalHeaderLabels(["Chọn", "Nhóm", "Nội dung tóm tắt"])
        header = self.bang_chon.horizontalHeader()
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.bang_chon.setColumnWidth(0, 50)
        self.bang_chon.setColumnWidth(1, 150)
        self.bang_chon.verticalHeader().setVisible(False)
        self.bang_chon.setAlternatingRowColors(True)
        self.bang_chon.setWordWrap(True)
        self.bang_chon.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        layout_trai.addWidget(self.bang_chon)

        bo_chia.addWidget(khung_trai)

        # ==========================================
        # KHUNG PHẢI: CẤU HÌNH THÔNG TIN VÀ XUẤT WORD
        # ==========================================
        khung_phai = QWidget()
        layout_phai = QVBoxLayout(khung_phai)
        layout_phai.setContentsMargins(10, 0, 0, 0)
        layout_phai.setAlignment(Qt.AlignmentFlag.AlignTop)

        lbl_phai = QLabel("2. Thông tin Tổ kiểm tra & Cơ quan ban hành:")
        lbl_phai.setStyleSheet("font-weight: bold; color: #1e3a8a; font-size: 14px;")
        layout_phai.addWidget(lbl_phai)

        box_cau_hinh = QGroupBox()
        form_layout = QFormLayout(box_cau_hinh)
        form_layout.setSpacing(15)

        # Các ô nhập liệu cấu hình
        cau_hinh_cu = self._doc_cau_hinh_cqt()  # <--- KÉO DỮ LIỆU CŨ RA

        self.txt_cqt = QLineEdit(cau_hinh_cu.get("cqt", "Thuế Tỉnh..."))
        self.txt_phong_doi = QLineEdit(cau_hinh_cu.get("phong_doi", "Phòng ..."))
        self.txt_to_truong = QLineEdit(cau_hinh_cu.get("to_truong", ""))
        self.txt_thanh_vien = QLineEdit(cau_hinh_cu.get("thanh_vien", ""))
        self.txt_lanh_dao = QLineEdit(cau_hinh_cu.get("lanh_dao", ""))
        self.txt_chuc_vu = QLineEdit(cau_hinh_cu.get("chuc_vu", "..."))

        # Định dạng css cho input
        css_input = "padding: 8px; border: 1px solid #94a3b8; border-radius: 4px; font-size: 13px;"
        for widget in [self.txt_cqt, self.txt_phong_doi, self.txt_to_truong, self.txt_thanh_vien, self.txt_lanh_dao,
                       self.txt_chuc_vu]:
            widget.setStyleSheet(css_input)

        form_layout.addRow("Cơ quan Thuế:", self.txt_cqt)
        form_layout.addRow("Phòng/Đội:", self.txt_phong_doi)
        form_layout.addRow("Tổ trưởng:", self.txt_to_truong)
        form_layout.addRow("Thành viên:", self.txt_thanh_vien)
        form_layout.addRow("Lãnh đạo ký:", self.txt_lanh_dao)
        form_layout.addRow("Chức vụ Lãnh đạo:", self.txt_chuc_vu)

        layout_phai.addWidget(box_cau_hinh)
        layout_phai.addStretch()

        self.btn_xuat_word = QPushButton(" TẠO FILE WORD (MẪU 02/QTKT)")
        self.btn_xuat_word.setIcon(qta.icon('fa5s.file-word', color='white'))
        self.btn_xuat_word.setStyleSheet(
            "background-color: #2563EB; color: white; font-weight: bold; padding: 15px; border-radius: 5px; font-size: 14px;")
        self.btn_xuat_word.clicked.connect(self.xuat_file_word)
        layout_phai.addWidget(self.btn_xuat_word)

        bo_chia.addWidget(khung_phai)
        bo_chia.setSizes([500, 400])
        layout.addWidget(bo_chia)

    def nap_danh_sach_rui_ro(self):
        self.bang_chon.setRowCount(len(self.ds_rui_ro))
        for row_idx, r in enumerate(self.ds_rui_ro):
            chk_widget = QWidget()
            chk_layout = QHBoxLayout(chk_widget)
            chk_layout.setContentsMargins(0, 0, 0, 0)
            chk_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
            chk = QCheckBox()
            chk.setChecked(True)
            chk_layout.addWidget(chk)
            self.bang_chon.setCellWidget(row_idx, 0, chk_widget)

            self.bang_chon.setItem(row_idx, 1, QTableWidgetItem(r.get('nhom', '')))
            self.bang_chon.setItem(row_idx, 2, QTableWidgetItem(r.get('ten_phan_tich', '')))
            self.bang_chon.resizeRowsToContents()

    def _format_so(self, val):
        try:
            return dinh_dang_so_vn(round(float(val), 0), 0)
        except:
            return str(val)

    def xuat_file_word(self):
        """Xuất Báo cáo trực tiếp bằng python-docx (Khóa cứng độ rộng bảng chống tràn)"""
        self._luu_cau_hinh_cqt()  # <--- THÊM DÒNG NÀY ĐỂ TỰ LƯU NGAY KHI BẤM NÚT XUẤT
        try:
            from docx import Document
            from docx.shared import Pt, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            QMessageBox.critical(self, "Thiếu Thư viện", "Chưa cài đặt python-docx. Gõ lệnh:\npip install python-docx")
            return

        file_path, _ = QFileDialog.getSaveFileName(self, "Lưu Phiếu nhận xét", "Phieu_Nhan_Xet_02QTKT.docx",
                                                   "Word Documents (*.docx)")
        if not file_path: return

        try:
            doc = Document()
            # Căn lề chuẩn A4: Trái 3cm, Phải 2cm -> Vùng vẽ bảng còn đúng 16cm
            for section in doc.sections:
                section.top_margin, section.bottom_margin = Cm(2), Cm(2)
                section.left_margin, section.right_margin = Cm(3), Cm(2)

            doc.styles['Normal'].font.name = 'Times New Roman'
            doc.styles['Normal'].font.size = Pt(13)

            # --- DỮ LIỆU ĐẦU VÀO ---
            cqt = self.txt_cqt.text().upper()
            phong_doi = self.txt_phong_doi.text().upper()
            to_truong = self.txt_to_truong.text()
            thanh_vien = self.txt_thanh_vien.text()
            lanh_dao = self.txt_lanh_dao.text()
            chuc_vu = self.txt_chuc_vu.text().upper()

            ten_nnt = self.du_lieu_tong_hop.get("ten_nnt", "<Chưa có Tên>").upper()
            mst = self.du_lieu_tong_hop.get("mst", "<Chưa có MST>")
            dia_chi = self.du_lieu_tong_hop.get("dia_chi", "")
            nnkd = self.du_lieu_tong_hop.get("nnkd", "")
            if not nnkd: nnkd = "..........................................................."

            # XỬ LÝ CHUỖI NĂM THÔNG MINH
            danh_sach_nam = self.du_lieu_tong_hop.get("danh_sach_nam", [])
            du_lieu_theo_nam = self.du_lieu_tong_hop.get("du_lieu_theo_nam", {})
            ky_pt = self.du_lieu_tong_hop.get("ky_pt", "Năm 2025")

            if not danh_sach_nam:
                chuoi_nam = ky_pt
            elif len(danh_sach_nam) == 1:
                chuoi_nam = f"năm {danh_sach_nam[0]}"
            else:
                chuoi_nam = "các năm " + ", ".join(danh_sach_nam)

            # =======================================================
            # 1. QUỐC HIỆU & TIÊU ĐỀ (CHUẨN FORM NGHỊ ĐỊNH 30 - BẢNG ẨN)
            # =======================================================

            # --- GÓC PHẢI: MẪU SỐ ---
            p_mau = doc.add_paragraph()
            p_mau.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_mau.add_run("Mẫu số: ")
            p_mau.add_run("02/QTKT\n").bold = True
            run_qd = p_mau.add_run("(Ban hành kèm theo Quyết định\nsố 970/QĐ-TCT ngày 14/7/2023\ncủa Tổng cục Thuế)")
            run_qd.italic = True

            # --- BẢNG ẨN BỐ TRÍ QUỐC HIỆU, TIÊU NGỮ VÀ TIÊU ĐỀ ---
            tb_head = doc.add_table(rows=5, cols=2)
            tb_head.autofit = False

            # Cân chỉnh lại tỷ lệ: Nới rộng cột phải để chữ Quốc hiệu nằm trọn 1 dòng
            for r in tb_head.rows:
                r.cells[0].width = Cm(5.5)  # Thu hẹp CQT
                r.cells[1].width = Cm(10.5) # Mở rộng Quốc hiệu (Tổng = 16cm)

            # DÒNG 1: CQT & Quốc hiệu
            p_cqt = tb_head.cell(0, 0).paragraphs[0]
            p_cqt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cqt = p_cqt.add_run(cqt)
            run_cqt.font.size = Pt(12) # Cơ quan chủ quản (size 12-13)

            p_qh = tb_head.cell(0, 1).paragraphs[0]
            p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_qh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
            run_qh.bold = True
            run_qh.font.size = Pt(13) # NĐ 30: Quốc hiệu size 12-13 in đậm

            # DÒNG 2: Phòng/Đội & Tiêu ngữ (Gạch chân)
            p_phong = tb_head.cell(1, 0).paragraphs[0]
            p_phong.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_phong = p_phong.add_run(phong_doi)
            run_phong.bold = True
            run_phong.underline = True
            run_phong.font.size = Pt(12) # Cơ quan ban hành (size 12-13 đậm)

            p_tn = tb_head.cell(1, 1).paragraphs[0]
            p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_tn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
            run_tn.bold = True
            run_tn.underline = True
            run_tn.font.size = Pt(14) # NĐ 30: Tiêu ngữ size 13-14 in đậm

            # DÒNG 3: Khoảng trắng ngắt quãng

            # DÒNG 4: Tên Phiếu Nhận Xét (Trộn 2 cột lại thành 1)
            cell_tieude = tb_head.cell(3, 0)
            cell_tieude.merge(tb_head.cell(3, 1))
            p_tieude = cell_tieude.paragraphs[0]
            p_tieude.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_tieude = p_tieude.add_run("PHIẾU NHẬN XÉT HỒ SƠ KHAI THUẾ")
            run_tieude.bold = True
            run_tieude.font.size = Pt(14) # Tiêu đề to (size 13-14)

            # DÒNG 5: Kỳ thuế (Trộn 2 cột lại thành 1)
            cell_ky = tb_head.cell(4, 0)
            cell_ky.merge(tb_head.cell(4, 1))
            p_ky = cell_ky.paragraphs[0]
            p_ky.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_kyl = p_ky.add_run("Kỳ thuế: ")
            run_kyl.bold = True
            run_kyl.font.size = Pt(13)
            run_kyv = p_ky.add_run(chuoi_nam.capitalize())
            run_kyv.bold = True
            run_kyv.font.size = Pt(13)

            # Đệm thêm 1 dòng trống nhỏ trước khi xuống Phần Thông tin NNT
            doc.add_paragraph()

            # --- THÔNG TIN NNT ---

            # --- THÔNG TIN NNT ---
            # [Phần thông tin NNT giữ nguyên ở đây...]

            doc.add_paragraph(f"- Người nộp thuế: {ten_nnt}.")
            doc.add_paragraph(f"- Mã số thuế: {mst}")
            doc.add_paragraph(f"- Địa chỉ: {dia_chi}")
            doc.add_paragraph(f"- Ngành nghề kinh doanh chính: {nnkd}")
            doc.add_paragraph(f"- Tổ kiểm tra:\n1. {to_truong}: Tổ trưởng.\n2. {thanh_vien}: Thành viên.")

            # =======================================================
            # 2. PHẦN KÊ KHAI
            # =======================================================
            p_i = doc.add_paragraph("I. PHẦN KÊ KHAI CỦA NGƯỜI NỘP THUẾ")
            p_i.runs[0].bold = True

            doc.add_paragraph("Căn cứ hồ sơ khai thuế:")
            doc.add_paragraph(f"- Tờ khai thuế GTGT mẫu số 01/GTGT: {chuoi_nam}.")
            doc.add_paragraph(f"- Tờ khai khấu trừ thuế TNCN mẫu số 05/KK-TNCN: {chuoi_nam}.")
            doc.add_paragraph(f"- Tờ khai quyết toán thuế TNDN mẫu số 03/TNDN {chuoi_nam} và Báo cáo tài chính {chuoi_nam}.")

            import datetime
            now = datetime.datetime.now()
            ngay_pt = f"{now.day:02d}/{now.month:02d}/{now.year}"
            doc.add_paragraph(f"{ten_nnt} đã nộp cho cơ quan thuế đến ngày {ngay_pt} số liệu kê khai thể hiện:")


            # =======================================================
            # HÀM BÚA TẠ: ĐÓNG ĐINH ĐỘ RỘNG
            # =======================================================
            def ep_do_rong_bang(bang, ds_do_rong):
                bang.autofit = False
                for row in bang.rows:
                    for idx, width in enumerate(ds_do_rong):
                        row.cells[idx].width = width

            def ve_bang_to_khai(tieu_de, data_list):
                doc.add_paragraph(tieu_de)
                if not data_list:
                    doc.add_paragraph("(Không có dữ liệu khai báo)").runs[0].italic = True
                    return

                tb = doc.add_table(rows=1, cols=3)
                tb.style = 'Table Grid'

                for i, text in enumerate(["STT", "Tên chỉ tiêu", "Số liệu (đồng)"]):
                    cell = tb.cell(0, i)
                    cell.text = text
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell.paragraphs[0].runs[0].bold = True

                for idx, row in enumerate(data_list):
                    r = tb.add_row()
                    r.cells[0].text = str(idx + 1)
                    r.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r.cells[1].text = str(row.get('ten_chi_tieu', ''))
                    val = row.get('so_lieu', '')
                    r.cells[2].text = self._format_so(val)
                    r.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

                ep_do_rong_bang(tb, [Cm(1.2), Cm(10.8), Cm(4.0)])

            # =======================================================
            # VẼ BẢNG KÊ KHAI (LẶP THEO TỪNG NĂM)
            # =======================================================
            if not danh_sach_nam:
                doc.add_paragraph("(Chưa có dữ liệu kê khai tờ khai)").runs[0].italic = True
            else:
                for nam in danh_sach_nam:
                    # Nếu kiểm tra nhiều năm, tạo Header ngắt trang nhỏ cho rõ ràng
                    if len(danh_sach_nam) > 1:
                        p_nam = doc.add_paragraph(f"\nSỐ LIỆU KÊ KHAI NĂM {nam}:")
                        p_nam.runs[0].bold = True
                        p_nam.runs[0].italic = True

                    data_nam = du_lieu_theo_nam.get(nam, {})
                    ve_bang_to_khai(f"1. Thuế GTGT năm {nam} (ĐVT: đồng):", data_nam.get("gtgt", []))
                    ve_bang_to_khai(f"2. Thuế TNDN năm {nam} (ĐVT: đồng):", data_nam.get("tndn", []))
                    ve_bang_to_khai(f"3. Thuế TNCN năm {nam} (ĐVT: đồng):", data_nam.get("tncn", []))

            # =======================================================
            # CỖ MÁY VẼ BẢNG ĐỐI CHIẾU HÓA ĐƠN (8 CỘT)
            # =======================================================
            p_ii = doc.add_paragraph("\nII. NHẬN XÉT HỒ SƠ KHAI THUẾ")
            p_ii.runs[0].bold = True
            doc.add_paragraph(
                "1. Phần ứng dụng hỗ trợ\nTheo kết quả phân tích từ ứng dụng TPR: Công ty xếp loại rủi ro ...")
            doc.add_paragraph("2. Phần ứng dụng không hỗ trợ:\na. Thuế GTGT")

            def ve_bang_hoa_don(tieu_de, ds_data):
                doc.add_paragraph(tieu_de)
                if not ds_data:
                    doc.add_paragraph("(Không có dữ liệu phát sinh)").runs[0].italic = True
                    return

                tb = doc.add_table(rows=2, cols=8)
                tb.style = 'Table Grid'

                # Trộn ô Header
                tb.cell(0, 0).merge(tb.cell(1, 0))  # STT
                tb.cell(0, 0).text = "STT"
                tb.cell(0, 1).merge(tb.cell(1, 1))  # Kỳ
                tb.cell(0, 1).text = "Kỳ"

                tb.cell(0, 2).merge(tb.cell(0, 3))  # Hóa đơn ĐT
                tb.cell(0, 2).text = "Hóa đơn điện tử"

                tb.cell(0, 4).merge(tb.cell(0, 5))  # Tờ khai
                tb.cell(0, 4).text = "Tờ khai GTGT"

                tb.cell(0, 6).merge(tb.cell(0, 7))  # Chênh lệch
                tb.cell(0, 6).text = "Chênh lệch (Tờ khai-Hóa đơn)"

                # Header phụ
                sub_heads = ["Giá trị chưa thuế", "Thuế GTGT"] * 3
                for i, text in enumerate(sub_heads):
                    tb.cell(1, i + 2).text = text

                for r in range(2):
                    for c in range(8):
                        for p in tb.cell(r, c).paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in p.runs: run.bold = True

                t_hd_truoc = t_hd_thue = t_tk_truoc = t_tk_thue = 0

                for idx, row_data in enumerate(ds_data):
                    r = tb.add_row()
                    r.cells[0].text = str(idx + 1)
                    r.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r.cells[1].text = str(row_data.get("ky", ""))
                    r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                    hd_tr = float(row_data.get("hd_truoc", 0) or 0)
                    hd_th = float(row_data.get("hd_thue", 0) or 0)
                    tk_tr = float(row_data.get("tk_truoc", 0) or 0)
                    tk_th = float(row_data.get("tk_thue", 0) or 0)

                    t_hd_truoc += hd_tr;
                    t_hd_thue += hd_th
                    t_tk_truoc += tk_tr;
                    t_tk_thue += tk_th

                    r.cells[2].text = self._format_so(hd_tr) if hd_tr else "-"
                    r.cells[3].text = self._format_so(hd_th) if hd_th else "-"
                    r.cells[4].text = self._format_so(tk_tr) if tk_tr else "-"
                    r.cells[5].text = self._format_so(tk_th) if tk_th else "-"
                    r.cells[6].text = self._format_so(tk_tr - hd_tr) if (tk_tr - hd_tr) else "-"
                    r.cells[7].text = self._format_so(tk_th - hd_th) if (tk_th - hd_th) else "-"

                    for c in range(2, 8): r.cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

                r_tong = tb.add_row()
                r_tong.cells[0].merge(r_tong.cells[1])
                r_tong.cells[0].text = "Tổng"
                r_tong.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_tong.cells[0].paragraphs[0].runs[0].bold = True

                vals = [t_hd_truoc, t_hd_thue, t_tk_truoc, t_tk_thue, t_tk_truoc - t_hd_truoc, t_tk_thue - t_hd_thue]
                for i, v in enumerate(vals):
                    c = r_tong.cells[i + 2]
                    c.text = self._format_so(v) if v else "-"
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    c.paragraphs[0].runs[0].bold = True

                # Khóa cứng độ rộng (Tổng = 16cm)
                ep_do_rong_bang(tb, [Cm(1.0), Cm(1.2), Cm(2.3), Cm(2.3), Cm(2.3), Cm(2.3), Cm(2.3), Cm(2.3)])

            # Kéo gọi vẽ 2 bảng Hóa đơn
            ve_bang_hoa_don("- Chênh lệch HHDV bán ra, thuế GTGT đầu ra:",
                            self.du_lieu_tong_hop.get("bang_hoa_don_br", []))
            ve_bang_hoa_don("\n- Chênh lệch HHDV mua vào, thuế GTGT đầu vào được khấu trừ:",
                            self.du_lieu_tong_hop.get("bang_hoa_don_mv", []))

            # =======================================================
            # VĂN PHONG DIỄN GIẢI RỦI RO (KHÔNG DÙNG BẢNG)
            # =======================================================
            doc.add_paragraph("\nb. Kết quả Phân tích Tự động (TNDN, BCTC, Rủi ro khác):")
            co_loi = False
            for row_idx in range(self.bang_chon.rowCount()):
                chk = self.bang_chon.cellWidget(row_idx, 0).findChild(QCheckBox)
                if chk.isChecked():
                    co_loi = True
                    loi = self.ds_rui_ro[row_idx]

                    val_a = self._format_so(loi.get('gia_tri_a', 0))
                    val_b = self._format_so(loi.get('gia_tri_b', 0))
                    try:
                        so_kq = float(loi.get('ket_qua_so', 0))
                    except:
                        so_kq = 0
                    chenh_lech = self._format_so(abs(so_kq))

                    nhom_loi = loi.get('nhom', 'Khác')
                    ten_loi = loi.get('ten_phan_tich', '')
                    kien_nghi = loi.get('thong_diep', '')

                    if loi.get('is_ab', False):
                        cau_van = f"- Qua kiểm tra nhóm chỉ tiêu {nhom_loi}, Tổ kiểm tra nhận thấy: {ten_loi}. Cụ thể, số liệu Vế A là {val_a} đồng đối chiếu với Vế B là {val_b} đồng, dẫn đến chênh lệch {chenh_lech} đồng. Yêu cầu đơn vị giải trình rõ nguyên nhân: {kien_nghi}"
                    else:
                        cau_van = f"- Qua rà soát nhóm {nhom_loi}, phát hiện rủi ro bất thường: {ten_loi}. Cụ thể, hệ thống ghi nhận giá trị là {val_a} đồng. Đề nghị đơn vị kiểm tra lại: {kien_nghi}"

                    p_loi = doc.add_paragraph(cau_van)
                    p_loi.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            if not co_loi:
                doc.add_paragraph("(Qua rà soát, Tổ kiểm tra chưa phát hiện thêm rủi ro bất thường nào khác)").runs[
                    0].italic = True

            p_iii = doc.add_paragraph("\nIII. KIẾN NGHỊ")
            p_iii.runs[0].bold = True
            doc.add_paragraph(
                f"Kiến nghị ban hành Thông báo giải trình để yêu cầu {ten_nnt} làm rõ các nội dung trên./.")

            import datetime
            now = datetime.datetime.now()
            p_date = doc.add_paragraph(f"......, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}")
            p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_date.runs[0].italic = True

            tb_f = doc.add_table(rows=2, cols=2)
            tb_f.autofit = False
            for r in tb_f.rows:
                r.cells[0].width = Cm(8)
                r.cells[1].width = Cm(8)

            c1_f = tb_f.cell(0, 0)
            c1_f.text = f"{phong_doi}\n{chuc_vu}\n\n\n\n{lanh_dao}"
            c1_f.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c1_f.paragraphs[0].runs[0].bold = True

            c2_f = tb_f.cell(0, 1)
            c2_f.text = f"TỔ KIỂM TRA\nTổ trưởng\n\n\n\n{to_truong}\n\nThành viên\n\n\n\n{thanh_vien}"
            c2_f.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c2_f.paragraphs[0].runs[0].bold = True

            doc.save(file_path)
            QMessageBox.information(self, "Thành công", f"Đã xuất Phiếu Nhận Xét ra file Word tại:\n{file_path}")

        except Exception as e:
            QMessageBox.critical(self, "Lỗi Render Word", f"Có lỗi khi tạo file Word: {e}")